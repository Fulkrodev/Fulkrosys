"""Lectura IA de uploads del cliente · extracción de contenido + clasificación.

El sistema debe LEER el contenido de los documentos que sube el cliente
(PDF/Word/Excel/texto) y clasificarlos a la medida ENS por CONTENIDO, no sólo
por el nombre de fichero. OCR (imágenes/PDF escaneado) degrada con gracia si el
binario tesseract no está disponible.
"""
from __future__ import annotations

import io

from backend.app.motors.m07_evidence.ai_classifier_service import (
    suggest_classification,
)
from backend.app.motors.m07_evidence.content_extraction_service import extract_text


def _docx(texto: str) -> bytes:
    import docx
    d = docx.Document()
    for line in texto.split("\n"):
        d.add_paragraph(line)
    bio = io.BytesIO()
    d.save(bio)
    return bio.getvalue()


def _xlsx(rows: list[list]) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _pdf(texto: str) -> bytes:
    from reportlab.pdfgen import canvas
    bio = io.BytesIO()
    c = canvas.Canvas(bio)
    y = 800
    for line in texto.split("\n"):
        c.drawString(50, y, line)
        y -= 20
    c.save()
    return bio.getvalue()


def test_extrae_txt():
    r = extract_text(b"Politica de control de acceso. MFA obligatorio.",
                     "text/plain", "pol.txt")
    assert r["method"] == "text"
    assert "MFA" in r["text"]


def test_extrae_docx():
    r = extract_text(_docx("Procedimiento de copias de seguridad backup restore."),
                     "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                     "proc.docx")
    assert r["method"] == "docx"
    assert "backup" in r["text"].lower()


def test_extrae_xlsx():
    r = extract_text(_xlsx([["fecha", "evento"], ["2026-06-11", "registro_acceso log"]]),
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     "reg.xlsx")
    assert r["method"] == "xlsx"
    assert "log" in r["text"].lower()


def test_extrae_pdf_texto():
    r = extract_text(_pdf("Informe de analisis de riesgos MAGERIT amenaza."),
                     "application/pdf", "riesgos.pdf")
    assert r["method"] == "pdf-text"
    assert "MAGERIT" in r["text"]


def test_clasifica_por_contenido_no_solo_nombre():
    # nombre genérico, pero el CONTENIDO dice backup → debe clasificar por contenido
    data = _docx("Plan de copias de seguridad y restauracion. Backup diario.")
    r = extract_text(data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                     "documento.docx")
    sug = suggest_classification("documento.docx", content_preview=r["text"])
    assert "mp.info.6" in sug.suggested_measure_codes
    assert sug.confidence > 0.0


def test_ocr_imagen_degrada_con_gracia_sin_tesseract():
    # Imagen falsa · sin tesseract debe devolver vacío + ocr_available=False, NO crashear
    r = extract_text(b"\x89PNG\r\n\x1a\n notreal", "image/png", "scan.png")
    assert r["text"] == ""
    assert r["method"].startswith("image-ocr")
    # nunca lanza · siempre devuelve dict
    assert "ocr_available" in r


def test_nunca_lanza_con_binario_desconocido():
    r = extract_text(b"\x00\x01\x02\x03\xff\xfe", "application/octet-stream", "x.bin")
    assert r["text"] == ""
    assert r["method"] in ("unsupported-binary", "fallback-decode")
