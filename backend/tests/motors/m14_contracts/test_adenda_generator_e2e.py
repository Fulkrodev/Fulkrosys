"""Sub-atom 1.B.7.1.3 · Test E2E adenda_generator service.

Flujo end-to-end real (DB + MinIO bucket fulkro-documents):
1. setup_test_project + crear Provider sintetico via real ORM columns
   (name + type + scope + criticality · NO razon_social/nif/normativas
   ver LECCION-OPS-029 adaptaciones documented en adenda_generator.py)
2. invoke AdendaGenerator.generate(...) con normativas_aplicables param
3. Verify row provider_addendums + DOCX en MinIO + signed URL
4. Cleanup: transactional rollback + DELETE MinIO objects al final

LECCION-OPS-027 sostenida: anchor en realidad ORM (no Provider.razon_social).
LECCION-OPS-029 sostenida: realidad codigo tiene precedencia · service adaptado.
"""
from __future__ import annotations

import uuid

import pytest
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from backend.app.core.storage.minio_client import (
    BUCKET_DOCUMENTS,
    get_minio_client,
)
from backend.app.models.m14_providers import Provider, ProviderAddendum
from backend.app.motors.m14_contracts.adenda_generator import (
    AdendaGenerator,
    TEMPLATE_CODE,
)
from backend.tests.conftest import _admin_setup, setup_test_project


def _cleanup_minio_object(object_key: str) -> None:
    """Best-effort cleanup MinIO objeto post-test."""
    try:
        get_minio_client().remove_object(BUCKET_DOCUMENTS, object_key)
    except Exception:
        pass


async def _create_provider(
    db: AsyncSession,
    project_uuid: uuid.UUID,
    name: str = "CloudHost Iberia S.L.",
    p_type: str = "cloud",
    criticality: str = "CRITICO",
    scope: str = "Servicios hosting cloud para sistemas en ambito ENS",
) -> Provider:
    """Inserta Provider via real ORM columns (LECCION-OPS-027)."""
    async with _admin_setup(db):
        provider = Provider(
            project_id=project_uuid,
            name=name,
            type=p_type,
            scope=scope,
            criticality=criticality,
        )
        db.add(provider)
        await db.flush()
    return provider


@pytest.mark.asyncio
async def test_adenda_generator_generates_full_addendum_with_three_normativas(
    db: AsyncSession,
):
    """Provider con normativas=['ENS','RGPD','NIS2'] genera adenda con 3 bloques."""
    _, project_id = await setup_test_project(db)
    project_uuid = uuid.UUID(project_id)

    provider = await _create_provider(db, project_uuid)

    gen = AdendaGenerator(db=db)
    result = await gen.generate(
        project_id=project_uuid,
        provider_id=provider.id,
        normativas_aplicables=["ENS", "RGPD", "NIS2"],
        contract_ref="CONTRATO-CLOUDHOST-2026-001",
        provider_extras={
            "nif": "B12345678",
            "domicilio": "Calle Cloud 1, Madrid",
            "representante": {"nombre": "CEO Cloud", "cargo": "CEO"},
            "categoria_servicio_ens": "ALTA",
        },
        client_extras={
            "domicilio": "Calle Cliente 1, Madrid",
            "representante": {"nombre": "Director Cliente", "cargo": "Director General"},
            "contacto_compliance": {"email": "compliance@cliente.example"},
        },
    )

    try:
        # Verificaciones core
        assert result.addendum_code.startswith("ADENDA-ENS-")
        assert result.minio_object_key == (
            f"corpus/addendums/{project_uuid}/{result.addendum_code}.docx"
        )
        assert result.normativas_cubiertas == ["ENS", "RGPD", "NIS2"]
        assert result.docx_size_bytes >= 10000, (
            f"DOCX E-604 esperado >10KB con 3 bloques · obtenido {result.docx_size_bytes}B"
        )
        assert result.template_code == TEMPLATE_CODE
        assert result.signed_url.startswith("http")

        # Row en provider_addendums
        addendum_res = await db.execute(
            select(ProviderAddendum).where(ProviderAddendum.id == result.addendum_id)
        )
        addendum = addendum_res.scalar_one()
        assert addendum.firmado_cliente is False
        assert addendum.firmado_proveedor is False
        assert addendum.generated_from_template_code == TEMPLATE_CODE
        assert addendum.normativas_cubiertas == ["ENS", "RGPD", "NIS2"]
        assert addendum.minio_object_key == result.minio_object_key
        # metadata_ stored con generator_version + docx_size_bytes
        meta = addendum.metadata_ or {}
        assert meta.get("generator_version") == "1.B.7.1.3"
        assert meta.get("docx_size_bytes") == result.docx_size_bytes

        # Verify DOCX bytes en MinIO (download + size match)
        obj = get_minio_client().get_object(BUCKET_DOCUMENTS, result.minio_object_key)
        downloaded = obj.read()
        obj.close()
        obj.release_conn()
        assert len(downloaded) == result.docx_size_bytes
    finally:
        _cleanup_minio_object(result.minio_object_key)


@pytest.mark.asyncio
async def test_adenda_generator_idempotent_same_code_returns_existing(
    db: AsyncSession,
):
    """Si addendum_code ya existe, generator devuelve existing sin re-render."""
    _, project_id = await setup_test_project(db)
    project_uuid = uuid.UUID(project_id)

    provider = await _create_provider(
        db, project_uuid, name="Test Provider Idempotent", p_type="saas",
        criticality="MEDIO",
    )

    gen = AdendaGenerator(db=db)
    result1 = await gen.generate(
        project_id=project_uuid,
        provider_id=provider.id,
        normativas_aplicables=["ENS"],
        addendum_code="ADENDA-TEST-IDEM-001",
        provider_extras={"nif": "B99999999"},
    )

    try:
        # Segunda llamada con mismo code -> debe devolver existing
        result2 = await gen.generate(
            project_id=project_uuid,
            provider_id=provider.id,
            normativas_aplicables=["ENS"],
            addendum_code="ADENDA-TEST-IDEM-001",
        )

        assert result1.addendum_id == result2.addendum_id
        assert result1.minio_object_key == result2.minio_object_key
        # Solo UNA fila en BD
        rows = await db.execute(
            select(ProviderAddendum).where(
                ProviderAddendum.addendum_code == "ADENDA-TEST-IDEM-001",
            )
        )
        assert len(rows.scalars().all()) == 1
    finally:
        _cleanup_minio_object(result1.minio_object_key)


@pytest.mark.asyncio
async def test_adenda_generator_only_ens_no_rgpd_nis2_blocks(db: AsyncSession):
    """Provider con normativas=['ENS'] solo. Adenda generada NO incluye bloques
    RGPD/NIS2/DORA · resultado mas pequeño (sin secciones 6/7/8)."""
    _, project_id = await setup_test_project(db)
    project_uuid = uuid.UUID(project_id)

    provider = await _create_provider(
        db, project_uuid, name="Solo ENS Provider", p_type="consultoria",
        criticality="BAJO", scope="Servicio minimo ENS",
    )

    gen = AdendaGenerator(db=db)
    result = await gen.generate(
        project_id=project_uuid,
        provider_id=provider.id,
        normativas_aplicables=["ENS"],
        provider_extras={"nif": "B88888888"},
    )

    try:
        assert result.normativas_cubiertas == ["ENS"]
        # Heurística de tamaño (la verificación semántica real es normativas_
        # cubiertas + ausencia del bloque RGPD abajo). Ceiling subido a 18KB
        # tras R11: el normalizador de plantillas añade header/footer canónico
        # (~1-2KB) a E-604, que antes no lo tenía.
        assert 5000 < result.docx_size_bytes < 18000, (
            f"E-604 solo-ENS esperado 5-18KB · obtenido {result.docx_size_bytes}B"
        )
        # Verify DOCX en MinIO + descomprime y busca texto bloque RGPD ausente
        obj = get_minio_client().get_object(BUCKET_DOCUMENTS, result.minio_object_key)
        downloaded = obj.read()
        obj.close()
        obj.release_conn()
        # Parse DOCX y verifica que NO contiene "RGPD Art. 28" en cuerpo
        from io import BytesIO
        from docx import Document as DocxDocument

        doc = DocxDocument(BytesIO(downloaded))
        text = "\n".join(p.text for p in doc.paragraphs)
        for t in doc.tables:
            for r in t.rows:
                for c in r.cells:
                    text += "\n" + c.text

        # Bloque ENS siempre presente (clausula 5)
        assert "ENS" in text
        assert "op.ext.1" in text
        # Bloques condicionales AUSENTES
        assert "RGPD Art. 28" not in text, (
            "E-604 con normativas=['ENS'] NO debe incluir bloque RGPD seccion 6"
        )
        assert "encargado de tratamiento" not in text, (
            "E-604 solo-ENS NO debe incluir clausula encargado tratamiento"
        )
        assert "NIS2" not in text or "OBLIGACIONES DE CIBERSEGURIDAD NIS2" not in text, (
            "E-604 solo-ENS NO debe incluir bloque NIS2 seccion 7"
        )
        assert "ACUERDOS TIC BAJO DORA" not in text, (
            "E-604 solo-ENS NO debe incluir bloque DORA seccion 8"
        )
    finally:
        _cleanup_minio_object(result.minio_object_key)
