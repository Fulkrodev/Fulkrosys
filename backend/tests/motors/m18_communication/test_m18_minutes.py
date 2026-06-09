"""Tests M18 — Acta de Comite (E-005) — Sesion 6.

Cubre:
- Creacion de acta + auto-codigo E-005-NNN secuencial
- Validaciones (tipo_comite, asistentes minimos)
- Generacion DOCX con header/footer/sigblock + hash SHA-256 + Ed25519
  (best-effort) + PDF (best-effort via LibreOffice)
- 0 leaks internos en el DOCX
- Magic link APROBACION_ACTA generado por asistente con email
- Email rico HTML/texto con placeholders correctos
- Registro de firma + transicion automatica fully_signed
- API endpoints completos
"""
from __future__ import annotations

import base64
import uuid
import zipfile
from datetime import date
from io import BytesIO

import pytest

from backend.app.database import set_tenant_context
from backend.app.motors.m18_communication.minutes_service import (
    MinutesService,
    MinutesValidationError,
)
from backend.tests.conftest import setup_test_project


BASE = "/api/v1/communication"


# ─────────── Helpers ───────────

async def _setup_tenant(db):
    client_id, project_id = await setup_test_project(db)
    await set_tenant_context(
        db,
        client_id=uuid.UUID(client_id),
        project_id=uuid.UUID(project_id),
    )
    return client_id, project_id


def _dataforma_attendees() -> list[dict]:
    """Asistentes tipicos del comite trimestral DataForma."""
    return [
        {
            "nombre": "Maria Perez Nunez",
            "cargo": "Consejera Delegada",
            "organizacion": "DataForma Galicia SL",
            "email": "ceo@dataforma.es",
        },
        {
            "nombre": "Jorge Fernandez Rodriguez",
            "cargo": "CISO / Responsable de Seguridad",
            "organizacion": "DataForma Galicia SL",
            "email": "rseg@dataforma.es",
        },
        {
            "nombre": "Marcos Mata Garcia",
            "cargo": "Consultor independiente en ENS",
            "organizacion": "Externo",
            "email": "marcos@matagarciamarcos.es",
        },
    ]


def _dataforma_orden_dia() -> list[dict]:
    return [
        {"punto": 1, "titulo": "Aprobacion del acta de la sesion anterior",
         "ponente": "Maria Perez Nunez", "tiempo_min": 5},
        {"punto": 2, "titulo": "Revision de avance del proyecto ENS",
         "ponente": "Marcos Mata Garcia", "tiempo_min": 15},
        {"punto": 3, "titulo": "Estado de los riesgos materializados",
         "ponente": "Jorge Fernandez Rodriguez", "tiempo_min": 10},
        {"punto": 4, "titulo": "Decision sobre alcance de la auditoria interna",
         "ponente": "Maria Perez Nunez", "tiempo_min": 15},
        {"punto": 5, "titulo": "Ruegos y preguntas",
         "ponente": None, "tiempo_min": 5},
    ]


def _dataforma_acuerdos() -> list[dict]:
    return [
        {"numero": 1,
         "descripcion": (
             "Aprobar el plan de auditoria interna del SGSI con alcance "
             "completo a los sistemas clinicos digitales."
         ),
         "owner": "Jorge Fernandez Rodriguez",
         "fecha_limite": "2026-05-31"},
        {"numero": 2,
         "descripcion": (
             "Designar a Raquel Cabrera Gomez como DPO externa formal "
             "y formalizar el contrato de prestacion de servicios."
         ),
         "owner": "Maria Perez Nunez",
         "fecha_limite": "2026-05-15"},
    ]


# ─────────── Creacion + validaciones ───────────

@pytest.mark.asyncio
async def test_create_minutes_assigns_codigo(db):
    _, project_id = await _setup_tenant(db)
    svc = MinutesService(db)
    m = await svc.create(
        project_id=uuid.UUID(project_id),
        tipo_comite="kickoff",
        titulo="Comite kickoff",
        fecha=date(2026, 4, 21),
        presidente="Maria Perez Nunez",
        secretario="Jorge Fernandez Rodriguez",
        asistentes=_dataforma_attendees(),
    )
    assert m.codigo == "E-005-001"
    assert m.estado == "draft"
    assert m.firmas == []


@pytest.mark.asyncio
async def test_create_minutes_codigo_secuencial(db):
    _, project_id = await _setup_tenant(db)
    svc = MinutesService(db)
    a1 = await svc.create(
        project_id=uuid.UUID(project_id),
        tipo_comite="kickoff", titulo="Comite 1", fecha=date(2026, 1, 15),
        presidente="P", secretario="S", asistentes=_dataforma_attendees(),
    )
    a2 = await svc.create(
        project_id=uuid.UUID(project_id),
        tipo_comite="seguimiento_trimestral", titulo="Comite 2",
        fecha=date(2026, 4, 21),
        presidente="P", secretario="S", asistentes=_dataforma_attendees(),
    )
    a3 = await svc.create(
        project_id=uuid.UUID(project_id),
        tipo_comite="extraordinario", titulo="Comite 3",
        fecha=date(2026, 5, 5),
        presidente="P", secretario="S", asistentes=_dataforma_attendees(),
    )
    assert a1.codigo == "E-005-001"
    assert a2.codigo == "E-005-002"
    assert a3.codigo == "E-005-003"


@pytest.mark.asyncio
async def test_create_minutes_invalid_tipo(db):
    _, project_id = await _setup_tenant(db)
    with pytest.raises(MinutesValidationError, match="tipo_comite"):
        await MinutesService(db).create(
            project_id=uuid.UUID(project_id),
            tipo_comite="reunion_de_pasillo",
            titulo="X", fecha=date(2026, 4, 21),
            presidente="P", secretario="S",
            asistentes=_dataforma_attendees(),
        )


@pytest.mark.asyncio
async def test_create_minutes_no_asistentes(db):
    _, project_id = await _setup_tenant(db)
    with pytest.raises(MinutesValidationError, match="al menos un asistente"):
        await MinutesService(db).create(
            project_id=uuid.UUID(project_id),
            tipo_comite="kickoff", titulo="X", fecha=date(2026, 4, 21),
            presidente="P", secretario="S", asistentes=[],
        )


# ─────────── DOCX ───────────

@pytest.mark.asyncio
async def test_generate_docx_creates_file_with_hash(db, tmp_path):
    _, project_id = await _setup_tenant(db)
    svc = MinutesService(db)
    m = await svc.create(
        project_id=uuid.UUID(project_id),
        tipo_comite="seguimiento_trimestral",
        titulo="Comite Q1 2026",
        fecha=date(2026, 4, 21),
        presidente="Maria Perez Nunez",
        secretario="Jorge Fernandez Rodriguez",
        asistentes=_dataforma_attendees(),
        orden_del_dia=_dataforma_orden_dia(),
        acuerdos=_dataforma_acuerdos(),
        proximos_pasos=[
            {"descripcion": "Ejecutar auditoria interna",
             "owner": "Jorge Fernandez Rodriguez",
             "fecha_limite": "2026-05-31"},
        ],
        lugar="Telematica (videoconferencia)",
    )
    m2 = await svc.generate_docx(
        m.id, cliente_razon="DataForma Galicia SL",
    )
    assert m2.estado == "generated"
    assert m2.docx_path is not None
    assert m2.hash_sha256 is not None
    assert len(m2.hash_sha256) == 64
    # docx valido (zip con [Content_Types].xml)
    from pathlib import Path as _P
    docx_bytes = _P(m2.docx_path).read_bytes()
    assert docx_bytes[:2] == b"PK"  # zip header
    z = zipfile.ZipFile(BytesIO(docx_bytes))
    assert "[Content_Types].xml" in z.namelist()


@pytest.mark.asyncio
async def test_generate_docx_no_internal_leaks(db):
    """El DOCX generado NO debe contener referencias internas FULKRO/Motor N."""
    _, project_id = await _setup_tenant(db)
    svc = MinutesService(db)
    m = await svc.create(
        project_id=uuid.UUID(project_id),
        tipo_comite="kickoff", titulo="Comite kickoff",
        fecha=date(2026, 4, 21),
        presidente="Maria Perez Nunez",
        secretario="Jorge Fernandez Rodriguez",
        asistentes=_dataforma_attendees(),
        orden_del_dia=_dataforma_orden_dia(),
        acuerdos=_dataforma_acuerdos(),
    )
    m2 = await svc.generate_docx(m.id, cliente_razon="DataForma Galicia SL")
    from pathlib import Path as _P
    from docx import Document as _Doc
    doc = _Doc(_P(m2.docx_path))

    # Concatenar todo el texto del documento
    all_text = []
    for p in doc.paragraphs:
        all_text.append(p.text)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    all_text.append(p.text)
    for section in doc.sections:
        for p in section.header.paragraphs:
            all_text.append(p.text)
        for tbl in section.header.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        all_text.append(p.text)
        for p in section.footer.paragraphs:
            all_text.append(p.text)
        for tbl in section.footer.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        all_text.append(p.text)

    full = "\n".join(all_text)
    forbidden = [
        "FULKRO", "Motor 18", "Motor 12", "Motor 6",
        "Document Factory", "Copiloto", "Agente ", "M3-G", "M5-G",
        "via Motor",
    ]
    for needle in forbidden:
        assert needle not in full, f"Leak interno encontrado: {needle!r}"

    # Y verificar contenido positivo: aparece codigo, titulo, asistentes
    assert "E-005-001" in full
    assert "Comite kickoff" in full
    assert "Maria Perez Nunez" in full
    assert "Jorge Fernandez Rodriguez" in full


@pytest.mark.asyncio
async def test_generate_docx_includes_acuerdos(db):
    _, project_id = await _setup_tenant(db)
    svc = MinutesService(db)
    m = await svc.create(
        project_id=uuid.UUID(project_id),
        tipo_comite="seguimiento_trimestral", titulo="X",
        fecha=date(2026, 4, 21),
        presidente="P", secretario="S",
        asistentes=_dataforma_attendees(),
        acuerdos=_dataforma_acuerdos(),
    )
    m2 = await svc.generate_docx(m.id, cliente_razon="DataForma Galicia SL")
    from pathlib import Path as _P
    from docx import Document as _Doc
    doc = _Doc(_P(m2.docx_path))
    full = "\n".join(p.text for p in doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    full += "\n" + p.text
    assert "auditoria interna" in full.lower()
    assert "DPO" in full or "dpo" in full.lower()


# ─────────── Magic link ───────────

@pytest.mark.asyncio
async def test_send_for_signature_creates_one_link_per_attendee(db):
    _, project_id = await _setup_tenant(db)
    svc = MinutesService(db)
    m = await svc.create(
        project_id=uuid.UUID(project_id),
        tipo_comite="kickoff", titulo="Comite",
        fecha=date(2026, 4, 21),
        presidente="P", secretario="S",
        asistentes=_dataforma_attendees(),
    )
    await svc.generate_docx(m.id, cliente_razon="DataForma Galicia SL")
    result = await svc.send_for_signature(
        m.id,
        cliente_razon="DataForma Galicia SL",
        base_url="https://app.fulkro.es",
    )
    assert result["total_envios"] == 3
    # Cada envio tiene magic_link_id, url, otp (purpose APROBACION_ACTA exige OTP)
    for envio in result["envios"]:
        assert envio["magic_link_id"]
        assert envio["url"].startswith("https://app.fulkro.es/ml/consume?token=")
        assert envio["otp"] is not None
        assert len(envio["otp"]) == 6  # 6-digit TOTP
        assert envio["subject"].startswith("Aprobaci")  # subject inicia con Aprobacion


@pytest.mark.asyncio
async def test_send_for_signature_email_html_has_acta_codigo(db):
    _, project_id = await _setup_tenant(db)
    svc = MinutesService(db)
    m = await svc.create(
        project_id=uuid.UUID(project_id),
        tipo_comite="kickoff", titulo="Comite",
        fecha=date(2026, 4, 21),
        presidente="P", secretario="S",
        asistentes=_dataforma_attendees()[:1],  # solo 1
    )
    await svc.generate_docx(m.id, cliente_razon="DataForma Galicia SL")
    result = await svc.send_for_signature(
        m.id, cliente_razon="DataForma Galicia SL", base_url="https://x",
    )
    envio = result["envios"][0]
    html = base64.b64decode(envio["html_b64"]).decode("utf-8")
    # El email debe mencionar el codigo del acta (placeholder resuelto)
    assert "E-005-001" in html
    # Y la razon social del cliente
    assert "DataForma Galicia SL" in html
    # Y NO debe haber leaks internos
    for needle in ["Motor 18", "Motor 12", "FULKRO_", "Document Factory"]:
        assert needle not in html


@pytest.mark.asyncio
async def test_send_for_signature_changes_state(db):
    _, project_id = await _setup_tenant(db)
    svc = MinutesService(db)
    m = await svc.create(
        project_id=uuid.UUID(project_id),
        tipo_comite="kickoff", titulo="X", fecha=date(2026, 4, 21),
        presidente="P", secretario="S",
        asistentes=_dataforma_attendees(),
    )
    await svc.generate_docx(m.id, cliente_razon="DataForma")
    await svc.send_for_signature(m.id, cliente_razon="DataForma", base_url="https://x")
    # Re-leer
    m2 = await svc.get(m.id)
    assert m2.estado == "sent_for_signature"
    assert m2.enviada_at is not None


# ─────────── Firma ───────────

@pytest.mark.asyncio
async def test_register_signature_marks_firma_and_state(db):
    _, project_id = await _setup_tenant(db)
    svc = MinutesService(db)
    m = await svc.create(
        project_id=uuid.UUID(project_id),
        tipo_comite="kickoff", titulo="X", fecha=date(2026, 4, 21),
        presidente="P", secretario="S",
        asistentes=_dataforma_attendees(),
    )
    await svc.generate_docx(m.id, cliente_razon="DataForma")
    sent = await svc.send_for_signature(m.id, cliente_razon="DataForma", base_url="https://x")

    # Primer asistente firma
    ml_id = uuid.UUID(sent["envios"][0]["magic_link_id"])
    m2 = await svc.register_signature(m.id, magic_link_id=ml_id, ip="1.2.3.4")
    assert m2.estado == "partially_signed"
    assert len(m2.firmas) == 1
    assert m2.firmas[0]["asistente_idx"] == 0
    assert m2.firmas[0]["ip"] == "1.2.3.4"

    # Segundo y tercer asistente firman → fully_signed
    ml_id2 = uuid.UUID(sent["envios"][1]["magic_link_id"])
    await svc.register_signature(m.id, magic_link_id=ml_id2)
    ml_id3 = uuid.UUID(sent["envios"][2]["magic_link_id"])
    m3 = await svc.register_signature(m.id, magic_link_id=ml_id3)
    assert m3.estado == "fully_signed"
    assert m3.fully_signed_at is not None
    assert len(m3.firmas) == 3


@pytest.mark.asyncio
async def test_register_signature_idempotent_per_asistente(db):
    """Si el mismo magic_link se registra dos veces, no duplica firma."""
    _, project_id = await _setup_tenant(db)
    svc = MinutesService(db)
    m = await svc.create(
        project_id=uuid.UUID(project_id),
        tipo_comite="kickoff", titulo="X", fecha=date(2026, 4, 21),
        presidente="P", secretario="S",
        asistentes=_dataforma_attendees(),
    )
    await svc.generate_docx(m.id, cliente_razon="X")
    sent = await svc.send_for_signature(m.id, cliente_razon="X", base_url="https://x")
    ml_id = uuid.UUID(sent["envios"][0]["magic_link_id"])
    await svc.register_signature(m.id, magic_link_id=ml_id)
    await svc.register_signature(m.id, magic_link_id=ml_id)
    m2 = await svc.get(m.id)
    assert len(m2.firmas) == 1


@pytest.mark.asyncio
async def test_register_signature_wrong_purpose_rejected(db):
    _, project_id = await _setup_tenant(db)
    svc = MinutesService(db)
    m = await svc.create(
        project_id=uuid.UUID(project_id),
        tipo_comite="kickoff", titulo="X", fecha=date(2026, 4, 21),
        presidente="P", secretario="S",
        asistentes=_dataforma_attendees(),
    )
    await svc.generate_docx(m.id, cliente_razon="X")
    await svc.send_for_signature(m.id, cliente_razon="X", base_url="https://x")

    # Crear un magic link de OTRO purpose → debe rechazar
    from backend.app.motors.m12_magic_link.service import MagicLinkService
    from backend.app.motors.m12_magic_link.schemas import MagicLinkGenerateRequest
    from backend.app.motors.m12_magic_link.purposes import MagicLinkPurpose
    # Post-MB-4.bis3: usa DESCARGA_DOSSIER_FINAL (legitimate · NO deprecated)
    # como purpose distinto a APROBACION_ACTA · test rejection sigue válido.
    other = await MagicLinkService(db).generate_magic_link(
        MagicLinkGenerateRequest(
            project_id=uuid.UUID(project_id),
            purpose=MagicLinkPurpose.DESCARGA_DOSSIER_FINAL,
            recipient_email="x@y.es",
        ),
        base_url="https://x",
    )
    with pytest.raises(MinutesValidationError, match="APROBACION_ACTA"):
        await svc.register_signature(m.id, magic_link_id=other.magic_link_id)


@pytest.mark.asyncio
async def test_signing_status_tracks_progress(db):
    _, project_id = await _setup_tenant(db)
    svc = MinutesService(db)
    m = await svc.create(
        project_id=uuid.UUID(project_id),
        tipo_comite="kickoff", titulo="X", fecha=date(2026, 4, 21),
        presidente="P", secretario="S",
        asistentes=_dataforma_attendees(),
    )
    await svc.generate_docx(m.id, cliente_razon="X")
    sent = await svc.send_for_signature(m.id, cliente_razon="X", base_url="https://x")

    status_initial = await svc.get_signing_status(m.id)
    assert status_initial["firmas_recibidas"] == 0
    assert status_initial["porcentaje_firmado"] == 0.0
    assert len(status_initial["pendientes"]) == 3

    ml = uuid.UUID(sent["envios"][0]["magic_link_id"])
    await svc.register_signature(m.id, magic_link_id=ml)
    status_partial = await svc.get_signing_status(m.id)
    assert status_partial["firmas_recibidas"] == 1
    assert status_partial["porcentaje_firmado"] == pytest.approx(33.3, abs=0.5)
    assert len(status_partial["pendientes"]) == 2


@pytest.mark.asyncio
async def test_regenerate_docx_with_signatures_includes_stamps(db):
    """Despues de fully_signed, el DOCX final incluye sello de firma por asistente."""
    _, project_id = await _setup_tenant(db)
    svc = MinutesService(db)
    m = await svc.create(
        project_id=uuid.UUID(project_id),
        tipo_comite="kickoff", titulo="Comite X", fecha=date(2026, 4, 21),
        presidente="P", secretario="S",
        asistentes=_dataforma_attendees(),
    )
    await svc.generate_docx(m.id, cliente_razon="DataForma")
    sent = await svc.send_for_signature(m.id, cliente_razon="DataForma", base_url="https://x")
    for envio in sent["envios"]:
        await svc.register_signature(
            m.id,
            magic_link_id=uuid.UUID(envio["magic_link_id"]),
            ip="10.0.0.1",
        )
    final = await svc.regenerate_docx_with_signatures(m.id, cliente_razon="DataForma")
    assert final.estado == "fully_signed"

    from pathlib import Path as _P
    from docx import Document as _Doc
    doc = _Doc(_P(final.docx_path))
    full = "\n".join(p.text for p in doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    full += "\n" + p.text
    assert "FIRMADO ELECTRONICAMENTE" in full
    assert "10.0.0.1" in full


# ─────────── API endpoints ───────────

@pytest.mark.asyncio
async def test_api_create_and_list_minutes(async_client, db):
    _, project_id = await setup_test_project(db)
    body = {
        "tipo_comite": "kickoff",
        "titulo": "Comite kickoff",
        "fecha": "2026-04-21",
        "presidente": "Maria Perez Nunez",
        "secretario": "Jorge Fernandez Rodriguez",
        "asistentes": _dataforma_attendees(),
    }
    r = await async_client.post(
        f"{BASE}/projects/{project_id}/minutes", json=body,
    )
    assert r.status_code == 201, r.text
    assert r.json()["codigo"] == "E-005-001"

    r2 = await async_client.get(f"{BASE}/projects/{project_id}/minutes")
    assert r2.status_code == 200
    assert len(r2.json()["minutes"]) == 1


@pytest.mark.asyncio
async def test_api_full_signing_flow(async_client, db):
    """E2E api: crear → docx → enviar → firmar 3 → fully_signed."""
    _, project_id = await setup_test_project(db)
    create_body = {
        "tipo_comite": "seguimiento_trimestral",
        "titulo": "Comite Q1 2026",
        "fecha": "2026-04-21",
        "presidente": "Maria Perez Nunez",
        "secretario": "Jorge Fernandez Rodriguez",
        "asistentes": _dataforma_attendees(),
        "orden_del_dia": _dataforma_orden_dia(),
        "acuerdos": _dataforma_acuerdos(),
    }
    r = await async_client.post(
        f"{BASE}/projects/{project_id}/minutes", json=create_body,
    )
    assert r.status_code == 201
    minutes_id = r.json()["id"]

    r2 = await async_client.post(
        f"{BASE}/minutes/{minutes_id}/generate-docx",
        json={"cliente_razon": "DataForma Galicia SL"},
    )
    assert r2.status_code == 200
    assert r2.json()["estado"] == "generated"
    assert r2.json()["hash_sha256"] is not None

    r3 = await async_client.post(
        f"{BASE}/minutes/{minutes_id}/send-for-signature",
        json={
            "cliente_razon": "DataForma Galicia SL",
            "base_url": "https://app.fulkro.es",
        },
    )
    assert r3.status_code == 200
    envios = r3.json()["envios"]
    assert len(envios) == 3

    # Firmar todos
    for envio in envios:
        rs = await async_client.post(
            f"{BASE}/minutes/{minutes_id}/sign",
            json={"magic_link_id": envio["magic_link_id"]},
        )
        assert rs.status_code == 200, rs.text

    rstatus = await async_client.get(f"{BASE}/minutes/{minutes_id}/signing-status")
    assert rstatus.status_code == 200
    assert rstatus.json()["estado"] == "fully_signed"
    assert rstatus.json()["porcentaje_firmado"] == 100.0
