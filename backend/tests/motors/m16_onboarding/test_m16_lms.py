"""Tests M16 — Mini-LMS — Sesion 6.

Cubre:
- Catalogo LMS carga correctamente con 3 cursos
- Asignacion de curso a lista de empleados (idempotente)
- Registro de asistencia + generacion E-502 con hash
- Envio de cuestionario + correccion + E-503 con detalle de respuestas
- Score calculation + pass threshold
- Progress aggregate por curso
- 0 leaks internos en E-502/E-503
- API endpoints (catalog, assign, attendance, submit-quiz, progress)
"""
from __future__ import annotations

import hashlib
import uuid
from pathlib import Path

import pytest

from backend.app.database import set_tenant_context
from backend.app.motors.m16_onboarding.lms_service import (
    LmsNotFoundError,
    LmsService,
    LmsStateError,
    LmsValidationError,
    get_course,
    list_courses,
)
from backend.tests.conftest import setup_test_project


BASE = "/api/v1/onboarding"


# ─────────── Helpers ───────────

async def _setup(db):
    client_id, project_id = await setup_test_project(db)
    await set_tenant_context(
        db,
        client_id=uuid.UUID(client_id),
        project_id=uuid.UUID(project_id),
    )
    return client_id, project_id


def _empleados_dataforma() -> list[dict]:
    return [
        {"nombre": "Laura Mendez", "email": "laura.mendez@dataforma.es",
         "cargo": "Medico/a especialista", "organizacion": "DataForma Galicia SL"},
        {"nombre": "David Otero", "email": "david.otero@dataforma.es",
         "cargo": "Administrativo", "organizacion": "DataForma Galicia SL"},
    ]


def _respuestas_correctas_lms001() -> dict[str, str]:
    """Respuestas correctas al cuestionario de LMS-001."""
    return {"q1": "b", "q2": "b", "q3": "b", "q4": "c", "q5": "b"}


def _respuestas_correctas_lms002() -> dict[str, str]:
    return {"q1": "b", "q2": "b", "q3": "c", "q4": "b"}


def _docx_full_text(doc) -> str:
    """Extract all text from a python-docx Document, including header/footer
    tables (where the document codigo + cliente live).
    """
    parts: list[str] = []
    parts.extend(p.text for p in doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    parts.append(p.text)
    for section in doc.sections:
        for p in section.header.paragraphs:
            parts.append(p.text)
        for tbl in section.header.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        parts.append(p.text)
        for p in section.footer.paragraphs:
            parts.append(p.text)
        for tbl in section.footer.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        parts.append(p.text)
    return "\n".join(parts)


# ─────────── Catalogo ───────────

def test_catalog_has_three_courses():
    courses = list_courses()
    codigos = {c["codigo"] for c in courses}
    assert codigos == {"LMS-001", "LMS-002", "LMS-003"}


def test_catalog_course_durations_match_spec():
    """LMS-001=30min, LMS-002=15min, LMS-003=20min (spec Sesion 6)."""
    c1 = get_course("LMS-001")
    c2 = get_course("LMS-002")
    c3 = get_course("LMS-003")
    assert c1["duracion_minutos"] == 30
    assert c2["duracion_minutos"] == 15
    assert c3["duracion_minutos"] == 20


def test_catalog_all_courses_have_quiz_with_questions():
    for c in list_courses():
        full = get_course(c["codigo"])
        preguntas = full.get("quiz", {}).get("preguntas", [])
        assert len(preguntas) >= 4
        for q in preguntas:
            assert q.get("id")
            assert q.get("enunciado")
            assert len(q.get("opciones", [])) >= 2
            assert q.get("respuesta_correcta")
            assert q.get("explicacion")


def test_catalog_all_courses_emit_e502_and_e503():
    for c in list_courses():
        full = get_course(c["codigo"])
        assert "E-502" in full.get("evidencias_generadas", [])
        assert "E-503" in full.get("evidencias_generadas", [])


# ─────────── Asignacion ───────────

@pytest.mark.asyncio
async def test_assign_course_to_employees(db):
    _, project_id = await _setup(db)
    svc = LmsService(db)
    assignments = await svc.assign_to_employees(
        uuid.UUID(project_id),
        "LMS-001",
        _empleados_dataforma(),
    )
    assert len(assignments) == 2
    for a in assignments:
        assert a.course_codigo == "LMS-001"
        assert a.course_titulo
        assert a.course_duracion_minutos == 30
        assert a.estado == "assigned"


@pytest.mark.asyncio
async def test_assign_is_idempotent_per_employee(db):
    _, project_id = await _setup(db)
    svc = LmsService(db)
    a1 = await svc.assign_to_employees(
        uuid.UUID(project_id), "LMS-001", _empleados_dataforma(),
    )
    a2 = await svc.assign_to_employees(
        uuid.UUID(project_id), "LMS-001", _empleados_dataforma(),
    )
    assert {str(x.id) for x in a1} == {str(x.id) for x in a2}

    all_ = await svc.list_by_project(uuid.UUID(project_id))
    assert len(all_) == 2  # no se duplicaron


@pytest.mark.asyncio
async def test_assign_invalid_course_codigo(db):
    _, project_id = await _setup(db)
    with pytest.raises(LmsNotFoundError):
        await LmsService(db).assign_to_employees(
            uuid.UUID(project_id), "LMS-999",
            _empleados_dataforma(),
        )


@pytest.mark.asyncio
async def test_assign_empleado_sin_email_rechaza(db):
    _, project_id = await _setup(db)
    with pytest.raises(LmsValidationError, match="email"):
        await LmsService(db).assign_to_employees(
            uuid.UUID(project_id), "LMS-001",
            [{"nombre": "Sin email"}],
        )


# ─────────── E-502 asistencia ───────────

@pytest.mark.asyncio
async def test_record_attendance_generates_e502(db):
    _, project_id = await _setup(db)
    svc = LmsService(db)
    assignments = await svc.assign_to_employees(
        uuid.UUID(project_id), "LMS-001",
        _empleados_dataforma()[:1],
    )
    a = assignments[0]
    updated = await svc.record_attendance(
        a.id, cliente_razon="DataForma Galicia SL",
    )
    assert updated.e502_path is not None
    assert updated.e502_hash is not None
    assert len(updated.e502_hash) == 64
    assert updated.estado == "in_progress"
    assert updated.iniciado_at is not None
    # E-502 es DOCX valido
    docx_bytes = Path(updated.e502_path).read_bytes()
    assert docx_bytes[:2] == b"PK"
    # hash match
    assert hashlib.sha256(docx_bytes).hexdigest() == updated.e502_hash


@pytest.mark.asyncio
async def test_e502_contains_asistente_and_no_leaks(db):
    _, project_id = await _setup(db)
    svc = LmsService(db)
    a = (await svc.assign_to_employees(
        uuid.UUID(project_id), "LMS-001",
        [_empleados_dataforma()[0]],
    ))[0]
    updated = await svc.record_attendance(
        a.id, cliente_razon="DataForma Galicia SL",
    )

    from docx import Document
    doc = Document(updated.e502_path)
    full = _docx_full_text(doc)
    # Content check
    assert "Laura Mendez" in full
    assert "LMS-001" in full
    assert "DataForma Galicia SL" in full
    assert "E-502" in full
    # No leaks
    for needle in ["FULKRO", "Motor 16", "Motor 18", "Document Factory",
                   "Copiloto", "Agente ", "M5-G"]:
        assert needle not in full, f"Leak interno en E-502: {needle!r}"


# ─────────── E-503 cuestionario ───────────

@pytest.mark.asyncio
async def test_submit_quiz_all_correct_marks_completed(db):
    _, project_id = await _setup(db)
    svc = LmsService(db)
    a = (await svc.assign_to_employees(
        uuid.UUID(project_id), "LMS-001",
        [_empleados_dataforma()[0]],
    ))[0]
    await svc.record_attendance(a.id, cliente_razon="DataForma")
    updated = await svc.submit_quiz(
        a.id,
        respuestas=_respuestas_correctas_lms001(),
        cliente_razon="DataForma Galicia SL",
    )
    assert updated.estado == "completed"
    assert updated.quiz_pass is True
    assert updated.quiz_score == 100.0
    assert updated.e503_path is not None
    assert updated.e503_hash is not None


@pytest.mark.asyncio
async def test_submit_quiz_failed_below_threshold(db):
    _, project_id = await _setup(db)
    svc = LmsService(db)
    a = (await svc.assign_to_employees(
        uuid.UUID(project_id), "LMS-001",
        [_empleados_dataforma()[0]],
    ))[0]
    await svc.record_attendance(a.id, cliente_razon="DataForma")
    # 1/5 correctas = 20%
    updated = await svc.submit_quiz(
        a.id,
        respuestas={"q1": "b", "q2": "a", "q3": "a", "q4": "a", "q5": "a"},
        cliente_razon="DataForma",
    )
    assert updated.estado == "failed"
    assert updated.quiz_pass is False
    assert updated.quiz_score == 20.0


@pytest.mark.asyncio
async def test_submit_quiz_partial_at_threshold_passes(db):
    """4/5 = 80% >= 70 → passes."""
    _, project_id = await _setup(db)
    svc = LmsService(db)
    a = (await svc.assign_to_employees(
        uuid.UUID(project_id), "LMS-001",
        [_empleados_dataforma()[0]],
    ))[0]
    await svc.record_attendance(a.id, cliente_razon="DataForma")
    updated = await svc.submit_quiz(
        a.id,
        respuestas={"q1": "b", "q2": "b", "q3": "b", "q4": "c", "q5": "a"},  # q5 mal
        cliente_razon="DataForma",
    )
    assert updated.quiz_score == 80.0
    assert updated.quiz_pass is True
    assert updated.estado == "completed"


@pytest.mark.asyncio
async def test_e503_contains_detail_and_no_leaks(db):
    _, project_id = await _setup(db)
    svc = LmsService(db)
    a = (await svc.assign_to_employees(
        uuid.UUID(project_id), "LMS-001",
        [_empleados_dataforma()[0]],
    ))[0]
    await svc.record_attendance(a.id, cliente_razon="DataForma")
    updated = await svc.submit_quiz(
        a.id,
        respuestas=_respuestas_correctas_lms001(),
        cliente_razon="DataForma Galicia SL",
    )

    from docx import Document
    doc = Document(updated.e503_path)
    full = _docx_full_text(doc)
    # Content check
    assert "E-503" in full
    assert "APTO" in full  # resultado
    assert "Laura Mendez" in full
    assert "100.0" in full  # score
    # Explicaciones incluidas
    assert "RD 311/2022" in full
    # No leaks
    for needle in ["FULKRO", "Motor 16", "Document Factory", "M5-G",
                   "Agente "]:
        assert needle not in full, f"Leak interno en E-503: {needle!r}"


# ─────────── Progreso agregado ───────────

@pytest.mark.asyncio
async def test_progress_aggregates_by_course(db):
    _, project_id = await _setup(db)
    svc = LmsService(db)
    # Asignar LMS-001 a 2 empleados, uno completa, otro no
    asigs = await svc.assign_to_employees(
        uuid.UUID(project_id), "LMS-001", _empleados_dataforma(),
    )
    await svc.record_attendance(asigs[0].id, cliente_razon="X")
    await svc.submit_quiz(
        asigs[0].id,
        respuestas=_respuestas_correctas_lms001(),
        cliente_razon="X",
    )
    # asigs[1] queda en 'assigned'

    # Asignar LMS-002 a 1 empleado, lo completa
    asigs2 = await svc.assign_to_employees(
        uuid.UUID(project_id), "LMS-002",
        [_empleados_dataforma()[0]],
    )
    await svc.record_attendance(asigs2[0].id, cliente_razon="X")
    await svc.submit_quiz(
        asigs2[0].id,
        respuestas=_respuestas_correctas_lms002(),
        cliente_razon="X",
    )

    progress = await svc.get_progress(uuid.UUID(project_id))
    assert progress["total_assignments"] == 3
    by_course = {c["course_codigo"]: c for c in progress["by_course"]}
    assert by_course["LMS-001"]["total"] == 2
    assert by_course["LMS-001"]["completed"] == 1
    assert by_course["LMS-001"]["assigned"] == 1
    assert by_course["LMS-001"]["pct_completados"] == 50.0
    assert by_course["LMS-002"]["completed"] == 1
    assert by_course["LMS-002"]["pct_completados"] == 100.0


# ─────────── API ───────────

@pytest.mark.asyncio
async def test_api_list_courses(async_client):
    r = await async_client.get(f"{BASE}/lms/courses")
    assert r.status_code == 200
    codigos = {c["codigo"] for c in r.json()["courses"]}
    assert codigos == {"LMS-001", "LMS-002", "LMS-003"}


@pytest.mark.asyncio
async def test_api_get_course_does_not_expose_correct_answers(async_client):
    r = await async_client.get(f"{BASE}/lms/courses/LMS-001")
    assert r.status_code == 200
    body = r.json()
    # El endpoint publico no debe exponer respuesta_correcta / explicacion
    for q in body.get("quiz", {}).get("preguntas", []):
        assert "respuesta_correcta" not in q
        assert "explicacion" not in q


@pytest.mark.asyncio
async def test_api_e2e_lms001_dataforma(async_client, db):
    """E2E API: asignar → asistencia → quiz 100% → evidencias generadas."""
    _, project_id = await setup_test_project(db)
    r1 = await async_client.post(
        f"{BASE}/projects/{project_id}/lms/assign",
        json={
            "course_codigo": "LMS-001",
            "empleados": _empleados_dataforma()[:1],
        },
    )
    assert r1.status_code == 200, r1.text
    assignment_id = r1.json()["assignments"][0]["id"]

    r2 = await async_client.post(
        f"{BASE}/lms/assignments/{assignment_id}/attendance",
        json={"cliente_razon": "DataForma Galicia SL"},
    )
    assert r2.status_code == 200
    assert r2.json()["e502_path"]

    r3 = await async_client.post(
        f"{BASE}/lms/assignments/{assignment_id}/submit-quiz",
        json={
            "respuestas": _respuestas_correctas_lms001(),
            "cliente_razon": "DataForma Galicia SL",
        },
    )
    assert r3.status_code == 200, r3.text
    body = r3.json()
    assert body["estado"] == "completed"
    assert body["quiz_pass"] is True
    assert body["quiz_score"] == 100.0
    assert body["e503_path"]

    r4 = await async_client.get(f"{BASE}/projects/{project_id}/lms/progress")
    assert r4.status_code == 200
    courses = {c["course_codigo"]: c for c in r4.json()["by_course"]}
    assert courses["LMS-001"]["completed"] == 1


# ─────────── Límite de intentos (§5.5 audit C6) ───────────

_RESPUESTAS_FALLIDAS_LMS001 = {"q1": "a", "q2": "a", "q3": "a", "q4": "a", "q5": "a"}


@pytest.mark.asyncio
async def test_quiz_attempts_increment_each_submit(db):
    """Cada envío del cuestionario incrementa el contador de intentos."""
    _, project_id = await _setup(db)
    svc = LmsService(db)
    a = (await svc.assign_to_employees(
        uuid.UUID(project_id), "LMS-001",
        [_empleados_dataforma()[0]],
    ))[0]
    await svc.record_attendance(a.id, cliente_razon="X")
    assert a.intentos == 0

    u1 = await svc.submit_quiz(
        a.id, respuestas=_RESPUESTAS_FALLIDAS_LMS001, cliente_razon="X",
    )
    assert u1.intentos == 1
    assert u1.estado == "failed"

    u2 = await svc.submit_quiz(
        a.id, respuestas=_RESPUESTAS_FALLIDAS_LMS001, cliente_razon="X",
    )
    assert u2.intentos == 2


@pytest.mark.asyncio
async def test_quiz_rejected_after_max_attempts(db):
    """Tras agotar max_attempts (3 en el catálogo) el envío se rechaza."""
    _, project_id = await _setup(db)
    svc = LmsService(db)
    a = (await svc.assign_to_employees(
        uuid.UUID(project_id), "LMS-001",
        [_empleados_dataforma()[0]],
    ))[0]
    await svc.record_attendance(a.id, cliente_razon="X")

    # 3 intentos fallidos consumen el límite
    for _ in range(3):
        await svc.submit_quiz(
            a.id, respuestas=_RESPUESTAS_FALLIDAS_LMS001, cliente_razon="X",
        )
    assert a.intentos == 3

    # El 4º envío se rechaza con mensaje claro
    with pytest.raises(LmsStateError) as exc:
        await svc.submit_quiz(
            a.id, respuestas=_respuestas_correctas_lms001(), cliente_razon="X",
        )
    assert "intentos" in str(exc.value).lower()
    # El contador no se mueve tras el rechazo
    assert a.intentos == 3


@pytest.mark.asyncio
async def test_quiz_passed_cannot_be_resubmitted(db):
    """Un cuestionario ya APROBADO no se puede reenviar (no re-aprobar)."""
    _, project_id = await _setup(db)
    svc = LmsService(db)
    a = (await svc.assign_to_employees(
        uuid.UUID(project_id), "LMS-001",
        [_empleados_dataforma()[0]],
    ))[0]
    await svc.record_attendance(a.id, cliente_razon="X")
    u = await svc.submit_quiz(
        a.id, respuestas=_respuestas_correctas_lms001(), cliente_razon="X",
    )
    assert u.estado == "completed"
    assert u.intentos == 1

    with pytest.raises(LmsStateError) as exc:
        await svc.submit_quiz(
            a.id, respuestas=_respuestas_correctas_lms001(), cliente_razon="X",
        )
    assert "aprobad" in str(exc.value).lower()


@pytest.mark.asyncio
async def test_catalog_quizzes_define_max_attempts():
    """Todos los cursos definen max_attempts en el catálogo (§5.5 audit C6)."""
    for c in list_courses():
        full = get_course(c["codigo"])
        assert int(full["quiz"]["max_attempts"]) >= 1
