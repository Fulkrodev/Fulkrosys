"""Integration tests LMS templates E500-E504 · sub-lote 1.B.3 architect-curated.

5 plantillas LMS en parametrize · pattern identico a test_bcp_templates_render.py.
Replica context simplificado architect-required.
"""
from __future__ import annotations

import re
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from backend.app.motors.m06_document_factory.rendering import render_docx

VAR_TEMPLATES = Path(__file__).resolve().parents[4] / "var" / "templates_docx"
PLACEHOLDER_PATTERN = re.compile(r"\{\{|\{%|%\}|\}\}")

LMS_TEMPLATES = ["E-500", "E-501", "E-502", "E-503", "E-504"]


@pytest.fixture
def lms_context() -> dict:
    return {
        "cliente": {
            "razon_social": "Test LMS SL", "nif": "B88888888",
            "organo_aprobador_politicas": "Comite Test",
        },
        "proyecto": {
            "codigo_documento_base": "TEST-ENS-2026", "version_actual": "1.0",
            "fecha_aprobacion_inicial": "2026-01-01",
            "sistema_principal": "Sistema Test",
        },
        "responsables": {
            "consultor": {"nombre": "Consultor Test", "cargo": "Consultor ENS"},
            "responsable_seguridad": {"nombre": "RSEG Test", "cargo": "CISO"},
        },
        "lms": {"plataforma": "Moodle Test", "phishing_plataforma": "GoPhish Test"},
        "periodo": {
            "trimestre": "Q1", "anyo": "2026", "fecha_cierre": "2026-03-31",
            "fecha_inicio": "2026-01-01", "fecha_fin": "2026-03-31",
        },
        "resumen": {"num_sesiones": 5, "num_asistencias": 30, "num_empleados_distintos": 20,
                    "tasa_cumplimiento_pct": 90},
        "sesiones": [
            {"fecha": "2026-01-15", "modulo": "M-G1-001", "grupo": "G1",
             "convocados": 5, "asistentes": 5, "tasa": 100, "facilitador": "F"},
        ],
        "empleados": [
            {"nombre": "Emp1", "rol": "Dev", "grupo": "G3", "horas_plan": 16,
             "horas_completadas": 14, "cumplimiento_pct": 88, "estado": "Parcial"},
        ],
        "nuevas_incorporaciones": [
            {"nombre": "N1", "fecha_alta": "2026-01-01", "fecha_acogida": "2026-01-20",
             "en_plazo": "Si"},
        ],
        "kpi_acogida_pct": "100",
        "evaluaciones": [
            {"modulo": "M1", "num_participantes": 5, "puntuacion_media": "80",
             "tasa_aprobado": 80},
        ],
        "brechas": [
            {"empleado_o_grupo": "Emp1", "grupo": "G3",
             "descripcion_brecha": "B", "accion_correctiva": "A", "plazo": "P"},
        ],
        "acciones_correctivas": [
            {"id": "AC1", "descripcion": "D", "afectado": "A",
             "responsable": "R", "plazo": "P", "estado": "E"},
        ],
        "observaciones_responsable": "OK",
        "campania": {
            "id": "PH-T", "nombre": "Test", "trimestre": "Q1", "anyo": "2026",
            "plataforma": "GoPhish Test", "fecha_envio": "2026-02-01",
            "duracion": "14d", "tipo_escenario": "Test", "dificultad": "Media",
            "plantillas": "1", "total_destinatarios": 50,
            "objetivos": ["O1"],
        },
        "segmentos": [
            {"nombre": "Todos", "num_destinatarios": 50, "justificacion": "J"},
        ],
        "resultados": {
            "entregados": 50, "abiertos": 38, "tasa_abierto": 76,
            "clicks": 4, "tasa_click": 8, "credenciales": 1, "tasa_credenciales": 2,
            "reportes": 28, "tasa_reporte": 56, "ignorados": 18, "tasa_ignorados": 36,
        },
        "segmentos_analisis": [
            {"nombre": "Todos", "tasa_click": 8, "tasa_reporte": 56,
             "delta_anterior": "Mejora"},
        ],
        "patrones": {"recurrentes": "0", "areas_alta": "X",
                     "areas_mejora": "Y", "early_reporters": "Z"},
        "acciones_colectivas": [
            {"id": "AC1", "descripcion": "D", "audiencia": "A",
             "responsable": "R", "plazo": "P"},
        ],
        "valoracion_global": "OK",
        "conclusion_texto": "Conclusion test",
        "recomendaciones": ["R1"],
        "kpis": {
            "cobertura_plan_pct": 92, "acogida_pct": 100, "g3g4_horas_pct": 88,
            "phishing_click_pct": 8, "phishing_reporte_pct": 56, "quiz_aprobado_pct": 85,
            "estado_global": "OK",
            "acogida": {"nuevos": 1, "completadas": 1, "fuera_plazo": 0, "estado": "OK"},
            "phishing_tendencia": "Mejora",
        },
        "cobertura_por_grupo": [
            {"grupo": "G1", "asignados": 20, "completados": 18,
             "cobertura_pct": 90, "estado": "OK"},
        ],
        "g3g4_detalle": [
            {"nombre": "Emp1", "grupo": "G3", "plan": 16, "real": 14,
             "pct": 88, "estado": "Parcial"},
        ],
        "phishing_campanias": [
            {"nombre": "PH-T", "fecha": "2026-02-01", "destinatarios": 50,
             "tasa_click": 8, "tasa_reporte": 56},
        ],
        "quizzes": [
            {"modulo": "M1", "participantes": 5, "media": "80", "tasa_aprobado": 80},
        ],
        "conclusiones_texto": "OK",
        "areas_mejora": [
            {"area": "A", "descripcion": "D", "accion": "AC"},
        ],
        "acciones_siguiente": [
            {"id": "AS1", "descripcion": "D", "responsable": "R",
             "plazo": "P", "estado": "E"},
        ],
        "elevaciones_comite": ["E1"],
    }


@pytest.mark.parametrize("codigo", LMS_TEMPLATES)
def test_lms_template_renders_without_placeholder_leak(codigo, tmp_path, lms_context):
    template_path = VAR_TEMPLATES / f"{codigo}.docx"
    assert template_path.exists(), f"Missing precompiled template: {template_path}"

    output_path = tmp_path / f"{codigo}_test.docx"
    render_docx(template_path, lms_context, output_path)

    assert output_path.exists()
    assert output_path.stat().st_size > 5000, (
        f"{codigo} DOCX too small ({output_path.stat().st_size} bytes)"
    )

    doc = DocxDocument(str(output_path))
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += "\n" + cell.text

    leaks = PLACEHOLDER_PATTERN.findall(text)
    assert len(leaks) == 0, f"{codigo} has placeholder leaks: {leaks[:5]}"
    assert "Test LMS SL" in text, f"{codigo} missing cliente.razon_social"
