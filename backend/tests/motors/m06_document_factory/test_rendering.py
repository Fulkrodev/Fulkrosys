"""Tests for Motor 6 rendering module."""
import pytest
from pathlib import Path

from docx import Document as DocxDocument

from backend.app.motors.m06_document_factory.rendering import (
    render_docx,
    convert_docx_to_pdf,
)
from backend.app.motors.m06_document_factory.exceptions import (
    MissingPlaceholderError,
    TemplateFileMissingError,
    PDFConversionError,
    RenderError,
)


def _create_template_docx(path: Path, placeholders: list[str] | None = None):
    """Create a minimal .docx template with Jinja2 placeholders."""
    doc = DocxDocument()
    doc.add_heading("Test Template", level=1)
    if placeholders:
        for ph in placeholders:
            doc.add_paragraph("{{ " + ph + " }}")
    else:
        doc.add_paragraph("{{ cliente.razon_social }}")
        doc.add_paragraph("{{ cliente.nif }}")
    doc.save(str(path))


class TestRenderDocx:

    def test_render_saves_to_output_path(self, tmp_path):
        tpl = tmp_path / "tpl.docx"
        out = tmp_path / "out.docx"
        _create_template_docx(tpl)
        result = render_docx(tpl, {"cliente": {"razon_social": "ACME", "nif": "B123"}}, out)
        assert result == out
        assert out.exists()

    def test_render_replaces_placeholders(self, tmp_path):
        tpl = tmp_path / "tpl.docx"
        out = tmp_path / "out.docx"
        _create_template_docx(tpl, ["nombre"])
        render_docx(tpl, {"nombre": "Test Value"}, out)
        # Read output and verify replacement
        doc = DocxDocument(str(out))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Test Value" in text

    def test_render_template_not_found_raises(self, tmp_path):
        with pytest.raises(TemplateFileMissingError):
            render_docx(tmp_path / "missing.docx", {}, tmp_path / "out.docx")

    def test_render_missing_required_placeholder_raises(self, tmp_path):
        tpl = tmp_path / "tpl.docx"
        out = tmp_path / "out.docx"
        _create_template_docx(tpl)
        with pytest.raises(MissingPlaceholderError, match="Missing required"):
            render_docx(tpl, {}, out, required_vars=["cliente.razon_social"])

    def test_render_with_nested_context(self, tmp_path):
        tpl = tmp_path / "tpl.docx"
        out = tmp_path / "out.docx"
        _create_template_docx(tpl, ["cliente.razon_social"])
        render_docx(
            tpl,
            {"cliente": {"razon_social": "Nested Corp"}},
            out,
            required_vars=["cliente.razon_social"],
        )
        assert out.exists()

    def test_render_docx_generic_error_raises_RenderError(self, tmp_path):
        """Corrupt template raises RenderError."""
        tpl = tmp_path / "corrupt.docx"
        tpl.write_bytes(b"not a real docx file")
        with pytest.raises(RenderError):
            render_docx(tpl, {}, tmp_path / "out.docx")


class TestConvertDocxToPdf:

    def test_convert_produces_pdf(self, tmp_path):
        tpl = tmp_path / "test.docx"
        _create_template_docx(tpl)
        pdf = convert_docx_to_pdf(tpl, tmp_path)
        assert pdf is not None
        assert pdf.exists()
        assert pdf.suffix == ".pdf"

    def test_convert_nonexistent_raises(self, tmp_path):
        with pytest.raises(PDFConversionError, match="DOCX not found"):
            convert_docx_to_pdf(tmp_path / "missing.docx", tmp_path)

    def test_convert_sets_output_dir(self, tmp_path):
        tpl = tmp_path / "test.docx"
        _create_template_docx(tpl)
        out_dir = tmp_path / "pdfs"
        pdf = convert_docx_to_pdf(tpl, out_dir)
        assert pdf is not None
        assert str(out_dir) in str(pdf.parent)


class TestFulkroIdentityInjection:
    """F-14-01 · Ejecutable 8 Pasada 16: identidad Fulkro centralizada propagada
    a las plantillas (footer/tel/web/email + autor de firmas) desde
    backend.app.fulkro_identity (single source of truth · Ejecutable 7.6)."""

    def test_default_firmas_uses_fulkro_identity(self):
        from backend.app.motors.m06_document_factory.rendering import _default_firmas
        from backend.app.fulkro_identity import (
            FULKRO_AUTHOR_NAME, FULKRO_AUTHOR_ROLE,
        )
        ctx: dict = {}
        _default_firmas(ctx)
        assert ctx["firmas"]["elaborado"]["nombre"] == FULKRO_AUTHOR_NAME
        assert ctx["firmas"]["elaborado"]["cargo"] == FULKRO_AUTHOR_ROLE

    def test_inject_brand_propagates_consultor_identity(self):
        from backend.app.motors.m06_document_factory.rendering import _inject_brand
        from backend.app.fulkro_identity import (
            FULKRO_FOOTER_TEXT, FULKRO_PHONE, FULKRO_EMAIL, FULKRO_WEB,
        )
        ctx: dict = {}
        # Sin logos → rama fallback (tpl no se usa) + inyección de identidad.
        _inject_brand(None, ctx, None, None)
        c = ctx["consultor"]
        assert c["nombre_comercial"] == "Fulkro"
        assert c["footer_text"] == FULKRO_FOOTER_TEXT
        assert c["telefono"] == FULKRO_PHONE
        assert c["email"] == FULKRO_EMAIL
        assert c["web"] == FULKRO_WEB
        # Fallback de marca = Fulkro (no "Marcos")
        assert "Fulkro" in c["header_brand"]
        assert "Marcos" not in c["header_brand"]
