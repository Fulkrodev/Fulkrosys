"""Hardening tests for Motor 6 rendering module.

Tests _find_missing_required deep dot-notation checking
and render_docx with required_vars parameter.
"""
import pytest
from pathlib import Path

from docx import Document as DocxDocument

from backend.app.motors.m06_document_factory.rendering import (
    _find_missing_required,
    render_docx,
)
from backend.app.motors.m06_document_factory.exceptions import (
    MissingPlaceholderError,
)


def _create_template_docx(path: Path, placeholders: list[str] | None = None):
    """Create a minimal .docx template with Jinja2 placeholders."""
    doc = DocxDocument()
    doc.add_heading("Hardening Test Template", level=1)
    if placeholders:
        for ph in placeholders:
            doc.add_paragraph("{{ " + ph + " }}")
    else:
        doc.add_paragraph("{{ cliente.razon_social }}")
    doc.save(str(path))


# ================================================================
# _find_missing_required
# ================================================================


class TestFindMissingRequired:

    def test_empty_required_returns_empty(self):
        result = _find_missing_required({"a": 1}, [])
        assert result == []

    def test_top_level_present(self):
        result = _find_missing_required({"name": "X"}, ["name"])
        assert result == []

    def test_top_level_missing(self):
        result = _find_missing_required({}, ["name"])
        assert result == ["name"]

    def test_dot_notation_present(self):
        ctx = {"cliente": {"razon_social": "ACME"}}
        result = _find_missing_required(ctx, ["cliente.razon_social"])
        assert result == []

    def test_dot_notation_missing_leaf(self):
        ctx = {"cliente": {"nif": "B123"}}
        result = _find_missing_required(ctx, ["cliente.razon_social"])
        assert result == ["cliente.razon_social"]

    def test_dot_notation_missing_parent(self):
        result = _find_missing_required({}, ["cliente.razon_social"])
        assert result == ["cliente.razon_social"]

    def test_none_value_treated_as_missing(self):
        ctx = {"cliente": {"razon_social": None}}
        result = _find_missing_required(ctx, ["cliente.razon_social"])
        assert result == ["cliente.razon_social"]

    def test_empty_string_treated_as_missing(self):
        ctx = {"cliente": {"razon_social": "   "}}
        result = _find_missing_required(ctx, ["cliente.razon_social"])
        assert result == ["cliente.razon_social"]

    def test_triple_dot_notation(self):
        ctx = {
            "responsables": {
                "responsable_seguridad": {
                    "nombre": "Ana",
                },
            },
        }
        result = _find_missing_required(
            ctx, ["responsables.responsable_seguridad.nombre"]
        )
        assert result == []

    def test_triple_dot_missing_deep_leaf(self):
        ctx = {
            "responsables": {
                "responsable_seguridad": {},
            },
        }
        result = _find_missing_required(
            ctx, ["responsables.responsable_seguridad.nombre"]
        )
        assert result == ["responsables.responsable_seguridad.nombre"]

    def test_mixed_present_and_missing(self):
        ctx = {"a": {"b": "ok"}}
        result = _find_missing_required(ctx, ["a.b", "a.c", "x.y"])
        assert sorted(result) == ["a.c", "x.y"]

    def test_numeric_value_not_missing(self):
        ctx = {"count": 0}
        result = _find_missing_required(ctx, ["count"])
        assert result == []

    def test_non_dict_intermediate_treated_as_missing(self):
        ctx = {"cliente": "just a string"}
        result = _find_missing_required(ctx, ["cliente.razon_social"])
        assert result == ["cliente.razon_social"]


# ================================================================
# render_docx with required_vars
# ================================================================


class TestRenderDocxWithRequiredVars:

    def test_raises_on_missing_required_vars(self, tmp_path):
        tpl = tmp_path / "tpl.docx"
        out = tmp_path / "out.docx"
        _create_template_docx(tpl)
        with pytest.raises(MissingPlaceholderError, match="Missing required"):
            render_docx(
                tpl,
                {"cliente": {"nif": "B123"}},
                out,
                required_vars=["cliente.razon_social"],
            )

    def test_succeeds_when_required_vars_present(self, tmp_path):
        tpl = tmp_path / "tpl.docx"
        out = tmp_path / "out.docx"
        _create_template_docx(tpl, ["cliente.razon_social"])
        result = render_docx(
            tpl,
            {"cliente": {"razon_social": "ACME Corp"}},
            out,
            required_vars=["cliente.razon_social"],
        )
        assert result == out
        assert out.exists()

    def test_permissive_when_required_vars_is_none(self, tmp_path):
        """When required_vars is None, render succeeds even with empty context."""
        tpl = tmp_path / "tpl.docx"
        out = tmp_path / "out.docx"
        _create_template_docx(tpl, ["nombre"])
        result = render_docx(tpl, {}, out, required_vars=None)
        assert result == out
        assert out.exists()
