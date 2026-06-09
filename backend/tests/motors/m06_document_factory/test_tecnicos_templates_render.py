"""Integration tests TECNICOS templates E700/E701/E705-E709 · sub-lote 1.B.4 architect-curated.

7 plantillas tecnicas en parametrize · pattern identico a test_lms_templates_render.py y
test_bcp_templates_render.py. Replica context simplificado architect-required cubriendo
auditoria interna (E-700/E-701), phishing formal (E-705), tabletop incidente (E-706),
restauracion backup (E-707), auditoria externa (E-708) y INES snapshot (E-709).
"""
from __future__ import annotations

import re
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from backend.app.motors.m06_document_factory.rendering import render_docx

VAR_TEMPLATES = Path(__file__).resolve().parents[4] / "var" / "templates_docx"
PLACEHOLDER_PATTERN = re.compile(r"\{\{|\{%|%\}|\}\}")

TECNICOS_TEMPLATES = ["E-700", "E-701", "E-705", "E-706", "E-707", "E-708", "E-709"]


@pytest.fixture
def tecnicos_context() -> dict:
    return {
        "cliente": {
            "razon_social": "Test Tecnicos SL", "nif": "B77777777",
            "sector": "servicios_TIC",
            "direccion": "Calle Test 1, 28001 Madrid",
            "representante_legal": "Director General Test",
            "organo_aprobador_politicas": "Comite Test",
            "ambito_ens": "Plataforma Test",
            "num_empleados_alcance": 50,
        },
        "proyecto": {
            "codigo_documento_base": "TEST-ENS-2026", "version_actual": "1.0",
            "fecha_aprobacion_inicial": "2026-01-01",
            "alcance": "Plataforma Test",
            "categoria_ens": "BASICA",
            "sistema_principal": "Plataforma Test",
        },
        "responsables": {
            "consultor": {"nombre": "Consultor Test", "cargo": "Consultor ENS"},
            "responsable_seguridad": {"nombre": "RSEG Test", "cargo": "CISO"},
        },
        # E-700 / E-701 auditoria interna
        "auditoria": {
            "periodo_inicio": "2025-06", "periodo_fin": "2026-05",
            "fecha_inicio": "2026-05-10", "fecha_fin": "2026-05-15",
            "equipo_auditor": "Auditor Test", "auditor_jefe": "Auditor Jefe Test",
            "e700_ref": "TEST-ENS-2026-700",
            "externa_fecha": "2026-06-30",
        },
        "resultados_por_familia": [
            {"nombre": "org", "aplicables": 4, "conformes": 4,
             "nc_menores": 0, "nc_mayores": 0, "observaciones": 0, "no_aplica": 0},
        ],
        "resumen": {
            "total_auditadas": 73, "conformidad_pct": 90, "total": 73,
            "ncs_iniciales": 3, "ncs_cerradas": 3, "ncs_pendientes": 0, "ncs_nuevas": 0,
        },
        "no_conformidades": [
            {"id": "NC-1", "medida": "op.exp.7", "tipo": "menor",
             "descripcion": "Test NC", "severidad": "Media", "plazo": "2026-06-30"},
        ],
        "observaciones": [
            {"id": "OBS-1", "medida": "mp.per.3",
             "descripcion": "Test obs", "recomendacion": "R test"},
        ],
        "plan_acciones": [
            {"id": "AC-1", "nc_ref": "NC-1", "descripcion": "Cerrar NC",
             "responsable": "CISO", "plazo": "2026-06-30", "estado": "En curso"},
        ],
        "conclusion": {
            "estado_global": "Conforme con NCs menores",
            "texto": "Conclusion test",
            "recomendacion_externa": "Avanzar",
            "estado": "Listo", "recomendacion": "Cerrar NCs",
        },
        "cierre_ncs": [
            {"id": "NC-1", "tipo": "menor", "descripcion": "T",
             "accion": "A", "evidencia": "E", "estado": "Cerrada"},
        ],
        "nuevas_ncs": [
            {"id": "NC-N1", "medida": "mp.com.4", "tipo": "menor",
             "descripcion": "T", "severidad": "Baja", "plazo": "2026-06-30"},
        ],
        "documentos_revisados": [
            {"nombre": "Politica", "version": "1.0",
             "aprobado_por": "Comite", "fecha": "2026-01-01", "coherencia": "OK"},
        ],
        "puntos_riesgo": [
            {"area": "Inventario", "descripcion": "T",
             "mitigacion": "M", "plazo": "2026-06-30"},
        ],
        "recomendaciones": ["R1", "R2"],
        # E-705 phishing formal
        "campania": {
            "id": "PH-T-FORMAL", "fecha_inicio": "2026-02-15", "fecha_fin": "2026-02-28",
            "duracion": "14d", "plataforma": "GoPhish",
            "tipo": "Spearphishing", "sector_simulado": "Helpdesk",
            "modalidad": "Ad-hoc", "autorizacion": "Comite", "total_destinatarios": 50,
        },
        "escenario": {"descripcion": "Test escenario"},
        "infraestructura": {
            "dominio_remitente": "test.example", "plataforma_envio": "GoPhish",
            "landing": "https://landing.test.example", "captura": "Token UUID",
            "tracking": "GoPhish events",
        },
        "segmentacion": [
            {"nombre": "Todos", "justificacion": "J", "numero": 50},
        ],
        "resultados": {
            "entregados": 50, "abiertos": 38, "tasa_abierto": 76,
            "clicks": 4, "tasa_click": 8, "credenciales": 1, "tasa_credenciales": 2,
            "reportes": 28, "tasa_reporte": 56,
            "benchmark_abierto": "65%", "benchmark_click": "12%",
            "benchmark_credenciales": "3%", "benchmark_reporte": "40%",
            "benchmark_tiempo": "30min",
            "tiempo_primer_reporte": "8min",
        },
        "campanyas_previas": [
            {"id": "PH-Q4", "fecha": "2025-11-15",
             "tasa_click": 12, "tasa_reporte": 48, "variacion": "-"},
        ],
        "tendencia_texto": "Mejora sostenida",
        "analisis_segmentos": [
            {"segmento": "Todos", "descripcion": "Homogeneo"},
        ],
        "señales_detectadas": ["Remitente externo", "Urgencia"],
        "señales_no_detectadas": ["Firma alterada"],
        "acciones_colectivas": [
            {"id": "AC1", "descripcion": "D", "audiencia": "A",
             "responsable": "R", "plazo": "P"},
        ],
        "valoracion": "Positivo · objetivo cumplido",
        "conclusion_texto": "Conclusion test",
        # E-706 tabletop incidente
        "ejercicio": {
            "fecha": "2026-03-15", "hora_inicio": "10:00", "hora_fin": "12:30",
            "duracion": "2h 30min", "modalidad": "On-site",
            "facilitador": "CISO Test", "observadores": "DPO",
            "escenario_resumen": "Ransomware BD",
            "escenario_completo": "Ransomware afecta BD primaria",
        },
        "inyecciones": [
            {"tiempo": "T+0", "descripcion": "Alerta SIEM",
             "reaccion_esperada": "Activar Comite"},
        ],
        "participantes": [
            {"nombre": "P1", "rol_ejercicio": "Presidente",
             "funcion_habitual": "CEO", "asistencia": "Presente"},
        ],
        "cronologia": [
            {"timestamp": "10:00", "descripcion": "Apertura"},
            {"timestamp": "10:30", "descripcion": "Restauracion completa"},
        ],
        "plazos": {
            "aepd_real": "60min", "aepd_estado": "OK",
            "nis2_inicial": "20min", "nis2_inicial_estado": "OK",
            "nis2_detallada": "65min", "nis2_detallada_estado": "OK",
            "interesados": "120min", "interesados_estado": "OK",
            "lucia": "Registrado", "lucia_estado": "OK",
            "dora": "N/A", "dora_estado": "No aplica",
        },
        "eficacia_e204": [
            {"apartado": "Deteccion", "aplicado": "Si", "eficaz": "Si",
             "comentario": "Inmediata"},
        ],
        "gaps": [
            {"descripcion": "Demora decision", "severidad": "Media",
             "procedimiento": "E-204 3.2", "accion": "Pre-aprobar",
             "responsable": "Legal", "plazo": "2026-06-30",
             "afectado": "Comunicacion", "causa_raiz": "Sin plantilla",
             "recomendacion": "Pre-aprobar"},
        ],
        # E-707 restauracion backup
        "prueba": {
            "id": "RT-T", "sistema": "BD Test", "tipo": "Completa",
            "modalidad": "Aislada", "fecha": "2026-05-12",
            "hora_inicio": "10:00", "hora_fin": "11:05",
            "duracion": "1h 05min", "operador": "DBA Test", "verificador": "CISO",
        },
        "backup": {
            "id": "BCK-T", "tipo": "Completa", "fecha": "2026-05-11 02:00",
            "tamaño": "45GB", "ubicacion": "S3 eu-west-1",
            "cifrado": "Si", "algoritmo": "AES-256-GCM",
            "hash_verificado": "OK SHA256",
        },
        "entorno": {
            "tipo": "EC2 t3.large", "aislamiento": "VPC dedicada",
            "recursos": "4 vCPU 16GB", "red": "deny egress",
            "no_interferencia": "Confirmada",
        },
        "metricas": {
            "rto_objetivo": "30min", "rto_real": "27min", "rto_cumplimiento": "OK",
            "rpo_objetivo": "5min", "rpo_real": "4min", "rpo_cumplimiento": "OK",
            "tiempo_descifrado": "3min", "tiempo_validacion": "8min",
            "tiempo_funcional": "12min",
        },
        "validaciones_integridad": [
            {"prueba": "Hash", "esperado": "Match", "obtenido": "Match", "estado": "OK"},
        ],
        "validaciones_funcionales": [
            {"funcionalidad": "Login", "esperado": "OK",
             "obtenido": "OK", "estado": "OK"},
        ],
        "incidencias": [
            {"id": "INC-1", "descripcion": "Throughput KMS",
             "causa": "Regional", "impacto": "Sin", "resolucion": "Documentar"},
        ],
        "hallazgos": [
            {"descripcion": "Baseline KMS", "severidad": "Baja",
             "accion": "Anexar runbook", "responsable": "DBA",
             "plazo": "2026-06-15", "estado": "Pendiente"},
        ],
        # E-708 auditoria externa ENS
        "auditoria_externa": {
            "entidad_auditora": "Entidad Test ENAC",
            "acreditacion": "ENAC 999/C-AC999",
            "auditor_jefe": "Auditor Externo Test",
            "equipo": "Equipo Test",
            "fecha_inicio": "2026-06-30", "fecha_fin": "2026-07-02",
            "ciclo": "Inicial", "modalidad": "On-site + remoto",
        },
        "resultados_externa_familia": [
            {"nombre": "org", "aplicables": 4, "conformes": 4,
             "nc_menores": 0, "nc_mayores": 0, "no_aplica": 0},
        ],
        "ncs_externas": [
            {"id": "NCE-1", "medida": "op.exp.7", "tipo": "menor",
             "descripcion": "T", "plazo": "2026-09-30"},
        ],
        "observaciones_externas": [
            {"id": "OBS-E-1", "medida": "mp.per.3",
             "descripcion": "Cobertura", "recomendacion": "Pildora"},
        ],
        "evidencias_revisadas": [
            {"nombre": "Politica", "tipo": "Politica",
             "periodo": "2025-2026", "estado": "OK"},
        ],
        "dictamen": {
            "resultado": "FAVORABLE CON OBSERVACIONES",
            "fundamentacion": "Test fundamentacion",
        },
        "propuesta_certificacion": "Procede certificacion",
        "vigencia_certificacion": "24 meses",
        "proxima_auditoria": "Seguimiento 12m",
        "acciones_correctivas_exigidas": [
            {"id": "ACE-1", "nc_ref": "NCE-1", "descripcion": "Revision anual",
             "plazo": "2026-09-30", "verificacion": "Acta firmada"},
        ],
        # E-709 INES snapshot
        "ines": {
            "año": "2026", "fecha_snapshot": "2026-12-31",
            "fecha_inicio": "2026-01-01", "fecha_fin": "2026-12-31",
            "fecha_envio": "2027-01-31",
        },
        "sistemas_declarados": [
            {"nombre": "Plataforma Test", "alcance": "Produccion",
             "categoria": "BASICA", "estado_certificacion": "Certificado",
             "proxima_auditoria": "2027-07"},
        ],
        "madurez": {
            "confidencialidad": {"objetivo": "E3", "real": "E3", "estado": "OK"},
            "integridad": {"objetivo": "E3", "real": "E3", "estado": "OK"},
            "disponibilidad": {"objetivo": "E3", "real": "E2", "estado": "WARN"},
            "autenticidad": {"objetivo": "E3", "real": "E3", "estado": "OK"},
            "trazabilidad": {"objetivo": "E3", "real": "E3", "estado": "OK"},
        },
        "cumplimiento_familia": [
            {"nombre": "org", "aplicables": 4, "cumplidas": 4,
             "parciales": 0, "no_cumplidas": 0, "pct": 100},
        ],
        "cumplimiento_global": 96,
        "incidentes_periodo": [
            {"id": "INC-1", "fecha": "2026-03-10", "tipo": "Phishing menor",
             "severidad": "Baja", "reportado": "No", "cerrado": "Si"},
        ],
        "resumen_incidentes": {"total": 3, "notificados": 0},
        "auditorias_periodo": [
            {"tipo": "Interna", "fecha": "2026-05-15",
             "resultado": "Conforme", "documento": "E-700"},
        ],
        "formacion": {
            "cobertura_pct": 94, "acogida_pct": 100,
            "click_pct": 8, "reporte_pct": 56,
        },
        "riesgos": {"total": 42, "alto_residual": 0, "tratados": 18, "seguimiento": 24},
        "inversion": {
            "infraestructura": "EUR 12.500", "servicios": "EUR 18.000",
            "formacion": "EUR 3.500", "personal": "0,5 FTE", "total": "EUR 34.000",
        },
        "evolucion": [
            {"indicador": "Cumplimiento %", "año_anterior": "88%",
             "año_actual": "96%", "variacion": "+8pp"},
        ],
        "hitos_año": ["Hito 1", "Hito 2"],
        "objetivos_proximo": ["Obj 1", "Obj 2"],
    }


@pytest.mark.parametrize("codigo", TECNICOS_TEMPLATES)
def test_tecnicos_template_renders_without_placeholder_leak(codigo, tmp_path, tecnicos_context):
    template_path = VAR_TEMPLATES / f"{codigo}.docx"
    assert template_path.exists(), f"Missing precompiled template: {template_path}"

    output_path = tmp_path / f"{codigo}_test.docx"
    render_docx(template_path, tecnicos_context, output_path)

    assert output_path.exists()
    assert output_path.stat().st_size > 5000, (
        f"{codigo} DOCX too small ({output_path.stat().st_size} bytes)"
    )

    doc = DocxDocument(str(output_path))
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += "\n" + cell.text

    leaks = PLACEHOLDER_PATTERN.findall(text)
    assert len(leaks) == 0, f"{codigo} has placeholder leaks: {leaks[:5]}"
    assert "Test Tecnicos SL" in text, f"{codigo} missing cliente.razon_social"
