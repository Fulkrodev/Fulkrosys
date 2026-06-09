"""Integration tests BCP templates E401-E406 · sub-lote 1.B.2 architect-curated.

Context BCP completo (procesos_criticos · estrategias · sistemas_criticos · etc.)
con structures architect-required. 6 plantillas en parametrize.
"""
from __future__ import annotations

import re
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from backend.app.motors.m06_document_factory.rendering import render_docx

VAR_TEMPLATES = Path(__file__).resolve().parents[4] / "var" / "templates_docx"
PLACEHOLDER_PATTERN = re.compile(r"\{\{|\{%|%\}|\}\}")

BCP_TEMPLATES = ["E-401", "E-402", "E-403", "E-404", "E-405", "E-406"]


@pytest.fixture
def bcp_context() -> dict:
    return {
        "cliente": {
            "razon_social": "Test Corp SL", "nif": "B99999999",
            "organo_aprobador_politicas": "Comite Test",
            "representante_legal": "Director Test",
        },
        "proyecto": {
            "codigo_documento_base": "TEST-ENS-2026", "version_actual": "1.0",
            "fecha_aprobacion_inicial": "2026-01-01",
            "sistema_principal": "Sistema Test", "alcance": "alcance test",
        },
        "responsables": {
            "consultor": {"nombre": "Consultor Test", "cargo": "Consultor ENS"},
            "responsable_seguridad": {"nombre": "RSEG Test", "cargo": "CISO"},
            "responsable_sistema": {"nombre": "CTO Test", "cargo": "CTO"},
            "responsable_servicio": {"nombre": "OPS Test", "cargo": "Head of Ops"},
            "comite_seguridad": {
                "presidente": "Presidente Test", "secretario": "Secretario Test",
                "miembros": ["Miembro A", "Miembro B", "Miembro C"],
            },
        },
        "procesos_criticos": [
            {"nombre": "P1", "rto": "4h", "rpo": "1h",
             "criticidad": "Alta", "descripcion": "desc"},
        ],
        "estrategias": [
            {"proceso": "P1", "estrategia": "E1", "coste_estimado": "C",
             "responsable": "R", "plazo": "P"},
        ],
        "ubicaciones_alternas": [
            {"nombre": "U1", "direccion": "D", "capacidad": "10",
             "tiempo_activacion": "2h"},
        ],
        "proveedores_criticos": [
            {"nombre": "PR1", "tipo_servicio": "TS", "sla": "99",
             "contacto_emergencia": "+34"},
        ],
        "servicios_criticos": [
            {"nombre": "SC1", "rto": "4h", "responsable": "R",
             "procedimiento": "P", "recursos": "R"},
        ],
        "sistemas_criticos": [
            {"nombre": "S1", "funcion": "F", "rto": "1h", "rpo": "15m",
             "mecanismo": "M"},
        ],
        "ubicaciones": {
            "primario": {"nombre": "P", "direccion": "D"},
            "secundario": {"nombre": "S", "direccion": "D",
                           "modalidad": "warm", "tiempo_activacion": "2h"},
        },
        "stakeholders": [
            {"tipo": "T", "prioridad": "P1", "canal": "C",
             "responsable_interno": "R", "plantilla_ref": "Sec",
             "plazo": "I"},
        ],
        "portavoces": [
            {"funcion": "F", "titular": "T", "suplente": "S", "audiencia": "A"},
        ],
        "autoridades_aplicables": [
            {"nombre": "AEPD", "regulacion": "RGPD",
             "cuando": "C", "canal": "S", "plazo": "72h"},
        ],
        "ejercicios_planificados": [
            {"tipo": "T", "fecha": "2026", "responsable": "R", "alcance": "A"},
        ],
        "ejercicio": {
            "tipo": "T", "fecha": "2026", "hora_inicio": "10", "hora_fin": "12",
            "duracion": "2h", "escenario": "E", "alcance": "A",
            "facilitador": "F", "observadores": "O", "sistemas_involucrados": "S",
            "objetivos": ["O1"],
            "participantes": [
                {"nombre": "N1", "rol_ejercicio": "R", "funcion_habitual": "F",
                 "asistencia": "Pres"},
            ],
            "tasa_participacion": "100",
            "cronologia": [{"timestamp": "T", "descripcion": "D"}],
        },
        "metricas": {
            "rto_objetivo": "4h", "rto_real": "3h", "rto_cumplimiento": "OK",
            "rpo_objetivo": "1h", "rpo_real": "30m", "rpo_cumplimiento": "OK",
            "checklist_pct": 90, "checklist_estado": "OK",
            "participacion_pct": 100, "participacion_estado": "OK",
            "integridad_pct": 100, "integridad_estado": "OK",
        },
        "observaciones": {
            "positivas": ["P1"], "negativas": ["N1"], "sorpresas": ["S1"],
        },
        "gaps": [
            {"descripcion": "G", "severidad": "Media", "afectado": "A",
             "causa_raiz": "C", "recomendacion": "R"},
        ],
        "acciones_correctivas": [
            {"id": "AC1", "descripcion": "D", "gap_ref": "G1",
             "responsable": "R", "plazo": "P", "estado": "E"},
        ],
        "lecciones": ["L1"],
        "recomendaciones_siguiente": ["R1"],
        "valoracion_global": "OK",
        "conclusion_texto": "Conclusion test",
    }


@pytest.mark.parametrize("codigo", BCP_TEMPLATES)
def test_bcp_template_renders_without_placeholder_leak(codigo, tmp_path, bcp_context):
    template_path = VAR_TEMPLATES / f"{codigo}.docx"
    assert template_path.exists(), f"Missing precompiled template: {template_path}"

    output_path = tmp_path / f"{codigo}_test.docx"
    render_docx(template_path, bcp_context, output_path)

    assert output_path.exists()
    assert output_path.stat().st_size > 5000, (
        f"{codigo} DOCX too small ({output_path.stat().st_size} bytes)"
    )

    doc = DocxDocument(str(output_path))
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += "\n" + cell.text

    leaks = PLACEHOLDER_PATTERN.findall(text)
    assert len(leaks) == 0, f"{codigo} has placeholder leaks: {leaks[:5]}"
    assert "Test Corp SL" in text, f"{codigo} missing cliente.razon_social"
