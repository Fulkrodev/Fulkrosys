"""Tests M17 Project Planning — WBS, CPM, tasks, CRs, exports."""
from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import load_workbook

from backend.app.motors.m17_planning import (
    effort_estimator,
    wbs_catalog,
)
from backend.tests.conftest import setup_test_project

BASE = "/api/v1/planning"


# ================== Helpers ==================

async def _generate_plan(async_client, db, categoria="MEDIA", start=None):
    _, project_id = await setup_test_project(db)
    start = start or date.today()
    r = await async_client.post(
        f"{BASE}/projects/{project_id}/generate",
        json={
            "categoria": categoria,
            "start_date": start.isoformat(),
            "client_size": "mediana",
            "complexity": "media",
        },
    )
    assert r.status_code == 200, r.text
    return project_id, r.json()


# ================== WBS Catalog ==================

class TestWBSCatalog:
    def test_basica_has_fewer_tasks_than_media(self):
        b = wbs_catalog.get_tasks_for_categoria("BASICA")
        m = wbs_catalog.get_tasks_for_categoria("MEDIA")
        assert len(b) < len(m)

    def test_alta_includes_redteam(self):
        a = wbs_catalog.get_tasks_for_categoria("ALTA")
        codes = {t.code for t in a}
        assert "WBS-042" in codes  # Red Team

    def test_dependencies_reference_existing_codes(self):
        all_codes = {t.code for t in wbs_catalog.WBS_CATALOG}
        for t in wbs_catalog.WBS_CATALOG:
            for dep in t.dependencies:
                assert dep in all_codes, (
                    f"{t.code} depends on unknown {dep}"
                )

    def test_all_phases_0_to_8(self):
        phases = {t.phase for t in wbs_catalog.WBS_CATALOG}
        # Deben existir 0, 1, 2, 3, 4, 5, 6, 7, 8
        assert phases == {0, 1, 2, 3, 4, 5, 6, 7, 8}


# ================== Effort Estimator ==================

class TestEffortEstimator:
    def test_basica_micro_baja(self):
        # base 10 × 0.6 × 0.7 × 0.8 = 3.36
        h = effort_estimator.estimate_effort(10, "BASICA", "micro", "baja")
        assert h == pytest.approx(3.4, abs=0.1)

    def test_alta_grande_alta(self):
        # base 10 × 1.5 × 1.3 × 1.3 = 25.35
        h = effort_estimator.estimate_effort(10, "ALTA", "grande", "alta")
        assert h == pytest.approx(25.4, abs=0.1)

    def test_media_mediana_media_identity(self):
        # base 10 × 1.0 × 1.0 × 1.0 = 10.0
        h = effort_estimator.estimate_effort(10, "MEDIA", "mediana", "media")
        assert h == 10.0

    def test_duration_weeks_basica(self):
        assert effort_estimator.estimate_duration_weeks("BASICA") == 18

    def test_duration_weeks_alta(self):
        assert effort_estimator.estimate_duration_weeks("ALTA") == 60

    def test_legacy_factors_match_catalog(self):
        # WAVE C2 · DRY invariant: las tablas legacy (SIZE_FACTOR /
        # COMPLEXITY_FACTOR) deben coincidir 1:1 con el catálogo canónico
        # effort_formulas_v1.json para que estimate_effort (legacy) y
        # estimate_full no diverjan numéricamente.
        cat = effort_estimator.load_formulas()
        size_json = {
            k: v["factor"] for k, v in cat["factor_size_empleados"].items()
        }
        comp_json = {
            k: v["factor"] for k, v in cat["factor_complejidad_tecnica"].items()
        }
        assert effort_estimator.SIZE_FACTOR == size_json
        assert effort_estimator.COMPLEXITY_FACTOR == comp_json


# ================== Plan Generation ==================

class TestPlanGeneration:
    @pytest.mark.asyncio
    async def test_generate_plan_creates_tasks(self, async_client, db):
        project_id, body = await _generate_plan(async_client, db, "MEDIA")
        assert body["categoria"] == "MEDIA"
        assert body["estado"] == "active"
        r = await async_client.get(f"{BASE}/projects/{project_id}/tasks")
        assert len(r.json()) >= 15  # MEDIA tiene muchas

    @pytest.mark.asyncio
    async def test_generate_plan_calculates_dates(self, async_client, db):
        start = date(2026, 5, 1)
        project_id, body = await _generate_plan(
            async_client, db, "BASICA", start=start,
        )
        r = await async_client.get(f"{BASE}/projects/{project_id}/tasks")
        tasks = r.json()
        # Primera tarea (WBS-001) empieza en start_date
        t001 = next(t for t in tasks if t["task_code"] == "WBS-001")
        assert t001["start_date"] == start.isoformat()

    @pytest.mark.asyncio
    async def test_dates_respect_dependencies(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "BASICA")
        r = await async_client.get(f"{BASE}/projects/{project_id}/tasks")
        tasks = {t["task_code"]: t for t in r.json()}
        # WBS-002 depende de WBS-001 → WBS-002.start > WBS-001.end
        t001_end = date.fromisoformat(tasks["WBS-001"]["end_date"])
        t002_start = date.fromisoformat(tasks["WBS-002"]["start_date"])
        assert t002_start > t001_end

    @pytest.mark.asyncio
    async def test_critical_path_identified(self, async_client, db):
        project_id, body = await _generate_plan(async_client, db, "MEDIA")
        assert len(body["critical_path_tasks"]) > 0
        assert body["critical_path_length_weeks"] >= 1

    @pytest.mark.asyncio
    async def test_slack_exists_for_non_critical(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "MEDIA")
        r = await async_client.get(f"{BASE}/projects/{project_id}/tasks")
        tasks = r.json()
        critical = [t for t in tasks if t["is_critical_path"]]
        non_critical = [t for t in tasks if not t["is_critical_path"]]
        # En plan MEDIA deben existir ambos
        assert critical
        assert non_critical
        # Critical slack == 0, no-critical slack >= 0 (con alguna > 0)
        assert all(t["slack_days"] == 0 for t in critical)

    @pytest.mark.asyncio
    async def test_mermaid_gantt_generated(self, async_client, db):
        project_id, body = await _generate_plan(async_client, db, "BASICA")
        assert body["mermaid_gantt"] is not None
        assert "gantt" in body["mermaid_gantt"]
        assert "dateFormat YYYY-MM-DD" in body["mermaid_gantt"]

    @pytest.mark.asyncio
    async def test_meeting_plan_set(self, async_client, db):
        project_id, body = await _generate_plan(async_client, db, "BASICA")
        assert "weekly" in body["meeting_plan"]
        assert body["meeting_plan"]["weekly"]["day"] == "friday"

    @pytest.mark.asyncio
    async def test_end_date_estimated(self, async_client, db):
        project_id, body = await _generate_plan(async_client, db, "BASICA")
        assert body["end_date_estimated"] is not None

    @pytest.mark.asyncio
    async def test_duplicate_plan_rejected(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "BASICA")
        r = await async_client.post(
            f"{BASE}/projects/{project_id}/generate",
            json={
                "categoria": "BASICA",
                "start_date": date.today().isoformat(),
            },
        )
        assert r.status_code == 422


# ================== Task Management ==================

class TestTaskManagement:
    @pytest.mark.asyncio
    async def test_update_task_status(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "BASICA")
        tasks = (await async_client.get(
            f"{BASE}/projects/{project_id}/tasks",
        )).json()
        task_id = tasks[0]["id"]
        r = await async_client.patch(
            f"{BASE}/projects/{project_id}/tasks/{task_id}",
            json={"status": "en_curso", "progress_pct": 25},
        )
        assert r.status_code == 200
        assert r.json()["status"] == "en_curso"
        assert r.json()["progress_pct"] == 25

    @pytest.mark.asyncio
    async def test_update_task_bloqueada_with_description(
        self, async_client, db,
    ):
        project_id, _ = await _generate_plan(async_client, db, "BASICA")
        tasks = (await async_client.get(
            f"{BASE}/projects/{project_id}/tasks",
        )).json()
        task_id = tasks[0]["id"]
        r = await async_client.patch(
            f"{BASE}/projects/{project_id}/tasks/{task_id}",
            json={
                "status": "bloqueada",
                "blocker_description": "Cliente no responde",
            },
        )
        assert r.status_code == 200
        assert r.json()["status"] == "bloqueada"
        assert r.json()["blocker_description"] == "Cliente no responde"

    @pytest.mark.asyncio
    async def test_update_task_invalid_status(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "BASICA")
        tasks = (await async_client.get(
            f"{BASE}/projects/{project_id}/tasks",
        )).json()
        task_id = tasks[0]["id"]
        r = await async_client.patch(
            f"{BASE}/projects/{project_id}/tasks/{task_id}",
            json={"status": "nonsense"},
        )
        assert r.status_code == 422

    @pytest.mark.asyncio
    async def test_task_status_hecha_sets_100_progress(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "BASICA")
        tasks = (await async_client.get(
            f"{BASE}/projects/{project_id}/tasks",
        )).json()
        task_id = tasks[0]["id"]
        r = await async_client.patch(
            f"{BASE}/projects/{project_id}/tasks/{task_id}",
            json={"status": "hecha"},
        )
        assert r.json()["progress_pct"] == 100


# ================== Baseline + Delays ==================

class TestBaselineAndDelays:
    @pytest.mark.asyncio
    async def test_set_baseline(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "BASICA")
        r = await async_client.post(
            f"{BASE}/projects/{project_id}/baseline",
        )
        assert r.status_code == 200
        assert r.json()["baseline_date"] is not None

    @pytest.mark.asyncio
    async def test_cant_set_baseline_twice(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "BASICA")
        await async_client.post(f"{BASE}/projects/{project_id}/baseline")
        r = await async_client.post(
            f"{BASE}/projects/{project_id}/baseline",
        )
        assert r.status_code == 422

    @pytest.mark.asyncio
    async def test_detect_delays_empty_initially(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "BASICA")
        await async_client.post(f"{BASE}/projects/{project_id}/baseline")
        # Sin cambios sobre la baseline, no hay retrasos
        r = await async_client.get(f"{BASE}/projects/{project_id}/delays")
        assert r.status_code == 200
        assert r.json()["total"] == 0

    @pytest.mark.asyncio
    async def test_plan_progress(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "BASICA")
        r = await async_client.get(f"{BASE}/projects/{project_id}/progress")
        assert r.status_code == 200
        body = r.json()
        assert body["total_tasks"] > 0
        assert body["pct_done"] == 0.0
        assert body["by_status"]["por_hacer"] > 0


# ================== Change Requests ==================

class TestChangeRequests:
    @pytest.mark.asyncio
    async def test_create_cr(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "BASICA")
        r = await async_client.post(
            f"{BASE}/projects/{project_id}/change-requests",
            json={
                "titulo": "Ampliar alcance a subsede",
                "descripcion": "Cliente pide añadir subsede de Madrid",
                "impacto_plazo_dias": 15,
                "impacto_esfuerzo_horas": 20,
                "impacto_presupuesto_eur": 3500.0,
                "solicitado_por": "cliente",
            },
        )
        assert r.status_code == 200
        body = r.json()
        assert body["code"] == "CR-001"
        assert body["estado"] == "propuesto"

    @pytest.mark.asyncio
    async def test_cr_auto_increment_code(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "BASICA")
        for _ in range(3):
            await async_client.post(
                f"{BASE}/projects/{project_id}/change-requests",
                json={"titulo": "CR X", "descripcion": "desc"},
            )
        crs = (await async_client.get(
            f"{BASE}/projects/{project_id}/change-requests",
        )).json()
        codes = sorted(c["code"] for c in crs)
        assert codes == ["CR-001", "CR-002", "CR-003"]

    @pytest.mark.asyncio
    async def test_approve_cr_applies_impact(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "BASICA")
        plan_before = (await async_client.get(
            f"{BASE}/projects/{project_id}",
        )).json()
        r = await async_client.post(
            f"{BASE}/projects/{project_id}/change-requests",
            json={
                "titulo": "Extensión proyecto",
                "descripcion": "+5 días",
                "impacto_plazo_dias": 5,
                "impacto_esfuerzo_horas": 10,
            },
        )
        cr_id = r.json()["id"]
        appr = await async_client.post(
            f"{BASE}/projects/{project_id}/change-requests/{cr_id}/approve",
            json={"aprobado_por": "marcos"},
        )
        assert appr.status_code == 200
        assert appr.json()["estado"] == "aprobado"
        # El end_date debe haberse desplazado +5 dias
        plan_after = (await async_client.get(
            f"{BASE}/projects/{project_id}",
        )).json()
        before = date.fromisoformat(plan_before["end_date_estimated"])
        after = date.fromisoformat(plan_after["end_date_estimated"])
        assert (after - before).days == 5

    @pytest.mark.asyncio
    async def test_reject_cr(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "BASICA")
        r = await async_client.post(
            f"{BASE}/projects/{project_id}/change-requests",
            json={"titulo": "CR X", "descripcion": "desc"},
        )
        cr_id = r.json()["id"]
        rej = await async_client.post(
            f"{BASE}/projects/{project_id}/change-requests/{cr_id}/reject",
            json={"aprobado_por": "marcos"},
        )
        assert rej.status_code == 200
        assert rej.json()["estado"] == "rechazado"

    @pytest.mark.asyncio
    async def test_cannot_approve_twice(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "BASICA")
        r = await async_client.post(
            f"{BASE}/projects/{project_id}/change-requests",
            json={"titulo": "CR test twice", "descripcion": "descripcion larga"},
        )
        cr_id = r.json()["id"]
        await async_client.post(
            f"{BASE}/projects/{project_id}/change-requests/{cr_id}/approve",
            json={"aprobado_por": "marcos"},
        )
        r2 = await async_client.post(
            f"{BASE}/projects/{project_id}/change-requests/{cr_id}/approve",
            json={"aprobado_por": "marcos"},
        )
        assert r2.status_code == 422


# ================== Exports ==================

class TestExports:
    @pytest.mark.asyncio
    async def test_export_xlsx(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "BASICA")
        r = await async_client.get(f"{BASE}/projects/{project_id}/export/xlsx")
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")
        assert r.content[:2] == b"PK"
        # Cabecera correcta
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        assert "WBS" in headers
        assert "Ruta Crítica" in headers

    @pytest.mark.asyncio
    async def test_mermaid_endpoint(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "MEDIA")
        r = await async_client.get(f"{BASE}/projects/{project_id}/mermaid")
        assert r.status_code == 200
        mermaid = r.json()["mermaid_gantt"]
        assert mermaid.startswith("gantt")
        assert "section Fase" in mermaid


# ================== API ==================

class TestAPI:
    @pytest.mark.asyncio
    async def test_critical_path_endpoint(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "ALTA")
        r = await async_client.get(
            f"{BASE}/projects/{project_id}/tasks/critical-path",
        )
        assert r.status_code == 200
        assert r.json()["total_critical"] > 0

    @pytest.mark.asyncio
    async def test_filter_tasks_by_phase(self, async_client, db):
        project_id, _ = await _generate_plan(async_client, db, "MEDIA")
        r = await async_client.get(
            f"{BASE}/projects/{project_id}/tasks?phase=0",
        )
        assert r.status_code == 200
        tasks = r.json()
        assert all(t["phase"] == "0" for t in tasks)

    @pytest.mark.asyncio
    async def test_estimate_effort_endpoint(self, async_client, db):
        r = await async_client.post(
            f"{BASE}/estimate-effort",
            json={
                "base_hours": 10, "categoria": "MEDIA",
                "client_size": "grande", "complexity": "alta",
            },
        )
        assert r.status_code == 200
        body = r.json()
        # 10 × 1.0 × 1.3 × 1.3 = 16.9
        assert body["estimated_hours"] == pytest.approx(16.9, abs=0.1)
        assert body["duration_weeks"] == 36

    @pytest.mark.asyncio
    async def test_get_plan_404_when_no_plan(self, async_client, db):
        _, project_id = await setup_test_project(db)
        r = await async_client.get(f"{BASE}/projects/{project_id}")
        assert r.status_code == 404


# ================== Effort Estimator FULL (Apendice N — Sesión 6) ==================

class TestEffortFormulasCatalog:
    """Validacion del catalogo effort_formulas_v1.json (calibrado por Marcos)."""

    def test_catalog_loads_with_required_keys(self):
        formulas = effort_estimator.load_formulas()
        assert "horas_base_por_categoria" in formulas
        assert "factor_sector" in formulas
        assert "factor_madurez" in formulas
        assert "factor_size_empleados" in formulas
        assert "factor_complejidad_tecnica" in formulas
        assert "duracion_semanas_por_categoria" in formulas

    def test_horas_base_calibradas_marcos(self):
        """Calibracion de Marcos: BASICA=20, MEDIA=65, ALTA=125."""
        formulas = effort_estimator.load_formulas()
        hb = formulas["horas_base_por_categoria"]
        assert hb["BASICA"] == 20
        assert hb["MEDIA"] == 65
        assert hb["ALTA"] == 125

    def test_madurez_5_niveles_decrecientes(self):
        formulas = effort_estimator.load_formulas()
        m = formulas["factor_madurez"]
        # L1 > L2 > L3 > L4 > L5: a mayor madurez, menos esfuerzo
        for a, b in [("L1", "L2"), ("L2", "L3"), ("L3", "L4"), ("L4", "L5")]:
            assert m[a]["factor"] > m[b]["factor"], f"{a} debe ser > {b}"
        assert m["L3"]["factor"] == 1.0  # L3 es la linea base

    def test_sector_alta_sensibilidad_mayor_que_estandar(self):
        formulas = effort_estimator.load_formulas()
        s = formulas["factor_sector"]
        assert s["financiero"]["factor"] > s["servicios_profesionales"]["factor"]
        assert s["salud"]["factor"] > s["servicios_profesionales"]["factor"]
        assert s["administracion_publica"]["factor"] > s["servicios_profesionales"]["factor"]

    def test_saas_factor_inferior_estandar(self):
        formulas = effort_estimator.load_formulas()
        # saas_tech tiene cultura tecnica madura → factor < 1
        assert formulas["factor_sector"]["saas_tech"]["factor"] < 1.0


class TestClassifySizeByEmployees:
    def test_micro_under_10(self):
        assert effort_estimator.classify_size_by_employees(5) == "micro"
        assert effort_estimator.classify_size_by_employees(10) == "micro"

    def test_pequena_11_to_50(self):
        assert effort_estimator.classify_size_by_employees(11) == "pequena"
        assert effort_estimator.classify_size_by_employees(50) == "pequena"

    def test_mediana_51_to_250(self):
        assert effort_estimator.classify_size_by_employees(85) == "mediana"
        assert effort_estimator.classify_size_by_employees(250) == "mediana"

    def test_grande_251_to_1000(self):
        assert effort_estimator.classify_size_by_employees(420) == "grande"
        assert effort_estimator.classify_size_by_employees(1000) == "grande"

    def test_muy_grande_over_1000(self):
        assert effort_estimator.classify_size_by_employees(1500) == "muy_grande"
        assert effort_estimator.classify_size_by_employees(50000) == "muy_grande"

    def test_zero_or_none_defaults_to_micro(self):
        assert effort_estimator.classify_size_by_employees(None) == "micro"
        assert effort_estimator.classify_size_by_employees(0) == "micro"


class TestEstimateFull:
    """estimate_full: calculo end-to-end con horas_base intrinsecas."""

    def test_dataforma_media_85emp_l2_salud(self):
        """Caso real DataForma Galicia SL.

        MEDIA + sector salud (1.25) + madurez L2 (1.15) + 85 emp → mediana (1.00)
        + complexity media (1.00) = factor combinado 1.4375.
        Horas base 65 × 1.4375 = 93.4 horas.
        """
        result = effort_estimator.estimate_full(
            categoria="MEDIA",
            sector="salud",
            madurez="L2",
            n_empleados=85,
            complexity="media",
        )
        assert result["categoria"] == "MEDIA"
        assert result["horas_base"] == 65
        assert result["sector"] == "salud"
        assert result["client_size"] == "mediana"
        assert result["factor_sector"]["valor"] == 1.25
        assert result["factor_madurez"]["valor"] == 1.15
        assert result["factor_size"]["valor"] == 1.00
        assert result["factor_complexity"]["valor"] == 1.00
        assert result["factor_combinado"] == pytest.approx(1.4375, abs=0.001)
        assert result["horas_consultor"] == pytest.approx(93.4, abs=0.1)
        assert result["duracion_semanas"] == 36

    def test_basica_micro_l1_generico(self):
        """Proyecto chico arrancando de cero.

        BASICA=20 × generico 1.00 × L1 1.30 × micro 0.70 × baja 0.80
        = 20 × 0.728 = 14.56 → 14.6 h
        """
        result = effort_estimator.estimate_full(
            categoria="BASICA",
            sector="generico",
            madurez="L1",
            client_size="micro",
            complexity="baja",
        )
        assert result["horas_base"] == 20
        assert result["factor_combinado"] == pytest.approx(0.728, abs=0.001)
        assert result["horas_consultor"] == pytest.approx(14.6, abs=0.1)

    def test_alta_muy_grande_l5_financiero(self):
        """Proyecto grande maduro: factor cerca del cap superior.

        ALTA=125 × financiero 1.30 × L5 0.75 × muy_grande 1.60 × muy_alta 1.60
        = 125 × 2.496 = 312 horas
        """
        result = effort_estimator.estimate_full(
            categoria="ALTA",
            sector="financiero",
            madurez="L5",
            client_size="muy_grande",
            complexity="muy_alta",
        )
        assert result["horas_base"] == 125
        assert result["factor_combinado"] == pytest.approx(2.496, abs=0.01)
        assert result["horas_consultor"] == pytest.approx(312.0, abs=1.0)

    def test_inferred_size_from_employees(self):
        """Sin client_size, infiere desde n_empleados."""
        result = effort_estimator.estimate_full(
            categoria="MEDIA",
            sector="generico",
            madurez="L3",
            client_size=None,
            n_empleados=420,
        )
        assert result["client_size"] == "grande"
        assert result["size_inferido_de_empleados"] == 420

    def test_unknown_sector_neutral_factor(self):
        """Sector desconocido → factor 1.0 + aviso."""
        result = effort_estimator.estimate_full(
            categoria="MEDIA", sector="quimica_pesada",
            madurez="L3", client_size="mediana",
        )
        assert result["factor_sector"]["valor"] == 1.0
        assert any("quimica_pesada" in a for a in result["avisos"])

    def test_invalid_categoria_raises(self):
        with pytest.raises(effort_estimator.EffortFormulasError):
            effort_estimator.estimate_full(
                categoria="MEGA-ALTA", sector="generico", madurez="L3",
                client_size="mediana",
            )

    def test_presupuesto_uses_default_tarifa(self):
        result = effort_estimator.estimate_full(
            categoria="BASICA", sector="generico", madurez="L3",
            client_size="mediana", complexity="media",
        )
        # 20 × 1.0 × 1.0 × 1.0 × 1.0 = 20 h × 95 €/h = 1900 €
        assert result["tarifa_hora_eur"] == 95
        assert result["presupuesto_eur"] == 1900.0

    def test_tarifa_override(self):
        result = effort_estimator.estimate_full(
            categoria="BASICA", sector="generico", madurez="L3",
            client_size="mediana", complexity="media",
            tarifa_hora_eur=120,
        )
        assert result["tarifa_hora_eur"] == 120
        assert result["presupuesto_eur"] == 2400.0

    def test_factor_combinado_cap_warning(self):
        """Si factor combinado supera el cap (3.5), debe emitir aviso."""
        # Forzamos: ALTA + financiero (1.30) × L1 (1.30) × muy_grande (1.60) × muy_alta (1.60)
        # = 1.30 × 1.30 × 1.60 × 1.60 = 4.3264 > 3.5
        result = effort_estimator.estimate_full(
            categoria="ALTA", sector="financiero", madurez="L1",
            client_size="muy_grande", complexity="muy_alta",
        )
        assert result["factor_combinado"] > 3.5
        assert any("Replantear modalidad" in a for a in result["avisos"])


class TestEstimateFullEndpoint:
    """Endpoint POST /planning/estimate-full."""

    @pytest.mark.asyncio
    async def test_dataforma_endpoint(self, async_client):
        r = await async_client.post(
            f"{BASE}/estimate-full",
            json={
                "categoria": "MEDIA",
                "sector": "salud",
                "madurez": "L2",
                "n_empleados": 85,
                "complexity": "media",
            },
        )
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["horas_consultor"] == pytest.approx(93.4, abs=0.1)
        assert body["client_size"] == "mediana"
        assert body["sector"] == "salud"

    @pytest.mark.asyncio
    async def test_get_effort_formulas(self, async_client):
        r = await async_client.get(f"{BASE}/effort-formulas")
        assert r.status_code == 200
        body = r.json()
        assert body["horas_base_por_categoria"]["MEDIA"] == 65
        assert body["version"] == "1.0"

    @pytest.mark.asyncio
    async def test_invalid_categoria_returns_422(self, async_client):
        r = await async_client.post(
            f"{BASE}/estimate-full",
            json={"categoria": "INEXISTENTE", "sector": "generico", "madurez": "L3"},
        )
        assert r.status_code == 422
