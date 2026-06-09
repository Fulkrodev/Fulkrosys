"""Integration tests m_live_records service · sub-lote 1.C.B checkpoint.

DB-backed tests sobre tabla ``live_records``. Usa fixtures conftest:
- db: transactional session as fulkro_app (NOSUPERUSER)
- _admin_setup: bypass RLS para data setup
- setup_test_project: crea client + project + setea app.current_project_id

Cobertura:
- RLS isolation cross-project
- create_record validation
- list pagination
- update partial
- archive
- dashboard counts
- export csv / xlsx
"""
from __future__ import annotations

import io
import uuid

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import text

from backend.app.motors.m_live_records.service import LiveRecordsService


_HAPPY_E300 = {
    "codigo_activo": "ACT-001",
    "nombre": "Servidor BBDD principal",
    "categoria_activo": "hardware",
    "propietario": "IT Dept",
    "ubicacion": "CPD Madrid",
    "criticidad": "alta",
    "fecha_alta": "2026-01-15",
}

_HAPPY_E305 = {
    "codigo_incidente": "INC-2026-001",
    "fecha_deteccion": "2026-02-10T14:30:00Z",
    "nivel_criticidad": "alto",
    "descripcion_breve": "Acceso no autorizado detectado",
    "descripcion_detallada": "Detalle del incidente con contexto completo.",
    "fase_nist": "contencion",
}


async def _setup_project(db) -> tuple[uuid.UUID, uuid.UUID]:
    """Create a project + return (project_id, user_id)."""
    from backend.tests.conftest import setup_test_project  # type: ignore

    _client_id, project_id = await setup_test_project(db)
    return uuid.UUID(project_id), uuid.uuid4()


@pytest.mark.asyncio
async def test_create_record_persists_and_validates(db):
    project_id, user_id = await _setup_project(db)
    svc = LiveRecordsService(db)

    record = await svc.create_record(
        project_id=project_id,
        register_type="E-300",
        entry_data=_HAPPY_E300,
        created_by=user_id,
    )

    assert record.id is not None
    assert record.register_type == "E-300"
    assert record.status == "active"
    assert record.entry_data["codigo_activo"] == "ACT-001"


@pytest.mark.asyncio
async def test_create_record_rejects_invalid_entry_data(db):
    project_id, user_id = await _setup_project(db)
    svc = LiveRecordsService(db)

    with pytest.raises(ValidationError):
        await svc.create_record(
            project_id=project_id,
            register_type="E-300",
            entry_data={"codigo_activo": "X"},  # missing required fields
            created_by=user_id,
        )


@pytest.mark.asyncio
async def test_create_record_rejects_unknown_register_type(db):
    project_id, user_id = await _setup_project(db)
    svc = LiveRecordsService(db)

    with pytest.raises(ValueError, match="Unknown register_type"):
        await svc.create_record(
            project_id=project_id,
            register_type="E-999",
            entry_data={},
            created_by=user_id,
        )


@pytest.mark.asyncio
async def test_list_records_paginates(db):
    project_id, user_id = await _setup_project(db)
    svc = LiveRecordsService(db)

    for i in range(5):
        payload = dict(_HAPPY_E300, codigo_activo=f"ACT-{i:03d}")
        await svc.create_record(
            project_id=project_id,
            register_type="E-300",
            entry_data=payload,
            created_by=user_id,
        )

    rows, total = await svc.list_records(
        project_id=project_id, register_type="E-300", limit=2, offset=0
    )
    assert total == 5
    assert len(rows) == 2

    rows_page2, _ = await svc.list_records(
        project_id=project_id, register_type="E-300", limit=2, offset=2
    )
    assert len(rows_page2) == 2
    assert rows[0].id != rows_page2[0].id


@pytest.mark.asyncio
async def test_update_record_partial_status(db):
    project_id, user_id = await _setup_project(db)
    svc = LiveRecordsService(db)

    record = await svc.create_record(
        project_id=project_id,
        register_type="E-300",
        entry_data=_HAPPY_E300,
        created_by=user_id,
    )
    other_user = uuid.uuid4()

    updated = await svc.update_record(
        project_id=project_id,
        record_id=record.id,
        updated_by=other_user,
        status="archived",
    )
    assert updated is not None
    assert updated.status == "archived"
    assert updated.entry_data["codigo_activo"] == "ACT-001"
    assert updated.updated_by == other_user


@pytest.mark.asyncio
async def test_archive_record_changes_status(db):
    project_id, user_id = await _setup_project(db)
    svc = LiveRecordsService(db)

    record = await svc.create_record(
        project_id=project_id,
        register_type="E-300",
        entry_data=_HAPPY_E300,
        created_by=user_id,
    )

    archived = await svc.archive_record(
        project_id=project_id, record_id=record.id, updated_by=user_id
    )
    assert archived is not None
    assert archived.status == "archived"


@pytest.mark.asyncio
async def test_compute_dashboard_counts_correctly(db):
    project_id, user_id = await _setup_project(db)
    svc = LiveRecordsService(db)

    a1 = await svc.create_record(
        project_id=project_id, register_type="E-300",
        entry_data=dict(_HAPPY_E300, codigo_activo="A-1"), created_by=user_id,
    )
    await svc.create_record(
        project_id=project_id, register_type="E-300",
        entry_data=dict(_HAPPY_E300, codigo_activo="A-2"), created_by=user_id,
    )
    a3 = await svc.create_record(
        project_id=project_id, register_type="E-300",
        entry_data=dict(_HAPPY_E300, codigo_activo="A-3"), created_by=user_id,
    )
    await svc.create_record(
        project_id=project_id, register_type="E-305",
        entry_data=_HAPPY_E305, created_by=user_id,
    )
    await svc.archive_record(
        project_id=project_id, record_id=a3.id, updated_by=user_id
    )

    blocks, total_active = await svc.compute_dashboard(project_id=project_id)
    assert len(blocks) == 26
    e300 = next(b for b in blocks if b.register_type == "E-300")
    e305 = next(b for b in blocks if b.register_type == "E-305")
    assert e300.active_count == 2
    assert e300.archived_count == 1
    assert e305.active_count == 1
    assert e305.archived_count == 0
    assert total_active == 3


@pytest.mark.asyncio
async def test_export_csv_includes_header_and_bom(db):
    project_id, user_id = await _setup_project(db)
    svc = LiveRecordsService(db)
    await svc.create_record(
        project_id=project_id, register_type="E-300",
        entry_data=_HAPPY_E300, created_by=user_id,
    )

    data = await svc.export_csv(project_id=project_id, register_type="E-300")
    assert data.startswith("﻿".encode("utf-8"))
    body = data.decode("utf-8")
    assert "codigo_activo" in body.splitlines()[0]
    assert "ACT-001" in body


@pytest.mark.asyncio
async def test_export_xlsx_renders_workbook(db):
    project_id, user_id = await _setup_project(db)
    svc = LiveRecordsService(db)
    await svc.create_record(
        project_id=project_id, register_type="E-300",
        entry_data=_HAPPY_E300, created_by=user_id,
    )

    data = await svc.export_xlsx(project_id=project_id, register_type="E-300")
    assert len(data) > 100
    wb = load_workbook(io.BytesIO(data))
    assert wb.active.title == "E-300"
    assert wb.active.cell(row=4, column=1).value == "id"


@pytest.mark.asyncio
async def test_rls_isolation_cross_project(db):
    """Insert in project_A · switching project_id MUST hide rows."""
    from backend.tests.conftest import _admin_setup  # type: ignore

    project_a, user_id = await _setup_project(db)

    svc = LiveRecordsService(db)
    rec_a = await svc.create_record(
        project_id=project_a, register_type="E-300",
        entry_data=dict(_HAPPY_E300, codigo_activo="PROJ-A"),
        created_by=user_id,
    )
    assert rec_a.id is not None

    project_b = uuid.uuid4()
    client_b = uuid.uuid4()
    cif_b = f"B{uuid.uuid4().hex[:8].upper()}"
    async with _admin_setup(db):
        await db.execute(text(
            "INSERT INTO clients (id, nombre, cif, created_at) "
            "VALUES (:id, 'Other Client', :cif, now())"
        ), {"id": str(client_b), "cif": cif_b})
        await db.execute(text(
            "INSERT INTO projects (id, client_id, nombre, created_at) "
            "VALUES (:id, :cid, 'Other Project', now())"
        ), {"id": str(project_b), "cid": str(client_b)})

    await db.execute(
        text("SELECT set_config('app.current_project_id', :pid, true)"),
        {"pid": str(project_b)},
    )
    rows_b, total_b = await svc.list_records(
        project_id=project_b, register_type="E-300"
    )
    assert total_b == 0
    assert rows_b == []

    await db.execute(
        text("SELECT set_config('app.current_project_id', :pid, true)"),
        {"pid": str(project_a)},
    )
    rows_a, total_a = await svc.list_records(
        project_id=project_a, register_type="E-300"
    )
    assert total_a == 1
    assert rows_a[0].entry_data["codigo_activo"] == "PROJ-A"
