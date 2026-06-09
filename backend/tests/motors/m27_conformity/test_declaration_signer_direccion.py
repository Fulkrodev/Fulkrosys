"""#2 Ola 7 · la Declaración de Conformidad 809 la firma la Dirección (sponsor).

CCN-STIC 809 Anexo A: el firmante es la Dirección/órgano superior (asume la
responsabilidad sobre la seguridad), NO el RSeg. El mecanismo de firma (canvas
Ed25519) y el distintivo NO se tocan · solo cambia QUIÉN firma y el bloque 7.
"""
from __future__ import annotations

import uuid

import pytest
from docx import Document
from sqlalchemy import text

from backend.app.motors.m27_conformity.distintivo_generator import (
    DistintivoContext,
    build_distintivo_context,
    generate_declaration_docx,
    generate_distintivo_svg,
)
from backend.tests.conftest import _admin_setup, setup_test_project


pytestmark = [pytest.mark.asyncio, pytest.mark.requires_db]


def _docx_text(bio) -> str:
    doc = Document(bio)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _ctx(**kw) -> DistintivoContext:
    base = dict(
        project_id=uuid.uuid4(),
        cert_id=uuid.uuid4(),
        client_name="ACME SL",
        client_cif="B12345678",
        client_domicilio="",
        system_name="Sistema X",
        system_category="BASICA",
        today="2026-06-04",
        expiry_date="2028-06-04",
        public_badge_url="/api/v1/conformity/badge/x/badge.svg",
        services_summary="",
        information_summary="",
        assets_essential_count=0,
        dda_total=0,
        dda_aplicables=0,
        dda_con_refuerzos=0,
        dda_no_aplica=0,
        conformes_count=0,
        no_conformes_count=0,
        pct_conformidad=0.0,
        rseg_name="Tecnico RSeg",
        rseg_email="rseg@acme.es",
        sponsor_name="Ana Dirección",
        sponsor_email="direccion@acme.es",
    )
    base.update(kw)
    return DistintivoContext(**base)


def test_declaration_docx_firma_la_direccion_no_rseg():
    txt = _docx_text(generate_declaration_docx(_ctx()))
    assert "Firma de la Dirección" in txt, "el bloque 7 debe ser de la Dirección"
    assert "Ana Dirección" in txt, "el sponsor aparece como firmante"
    assert "Firma del Responsable de Seguridad" not in txt, (
        "ya NO firma el RSeg (deprecado el bloque RSeg)"
    )
    assert "Tecnico RSeg" not in txt, "el RSeg no es el firmante del bloque 7"
    assert "CCN-STIC 809 Anexo A" in txt, "referencia normativa presente"


def test_distintivo_svg_still_generated():
    # Regresión (preocupación Marcos): el distintivo SVG sobrevive al cambio de
    # firmante · es entregable que el cliente quiere · sección aparte.
    svg = generate_distintivo_svg(_ctx())
    assert svg.lstrip().startswith("<"), "SVG válido generado"
    assert "ACME SL" in svg, "el distintivo lleva el nombre del cliente"


async def test_build_context_resolves_sponsor(db):
    client_id, project_id = await setup_test_project(db)
    async with _admin_setup(db):
        await db.execute(text(
            "INSERT INTO client_contacts "
            "(id, client_id, full_name, email, role_title, role_category, "
            " is_active, created_at) "
            "VALUES (gen_random_uuid(), :cid, 'Ana Dirección', 'dir@acme.es', "
            "'Directora General', 'sponsor', true, now())"
        ), {"cid": client_id})

    await db.execute(
        text("SELECT set_config('app.current_project_id', :p, true)"),
        {"p": project_id},
    )
    await db.execute(
        text("SELECT set_config('app.current_client_id', :c, true)"),
        {"c": client_id},
    )
    ctx = await build_distintivo_context(db, uuid.UUID(project_id))
    assert ctx.sponsor_name == "Ana Dirección", "resuelve el contacto sponsor (M30)"
    assert ctx.sponsor_email == "dir@acme.es"
