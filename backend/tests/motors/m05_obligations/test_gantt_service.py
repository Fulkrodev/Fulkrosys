"""Tests for Motor 5 -- Gantt Service (async DB tests).

Uses the same setup pattern as test_instantiation.py: create a project
via ``setup_test_project``, set tenant context, instantiate obligations,
then build the Gantt plan.
"""
from __future__ import annotations

import io
import uuid
from datetime import date

from openpyxl import load_workbook

from backend.app.database import set_tenant_context
from backend.app.motors.m05_obligations.gantt_service import (
    build_gantt_for_project,
    export_gantt_to_xlsx_bytes,
)
from backend.app.motors.m05_obligations.instantiation_service import (
    instantiate_obligations_for_multiple_gaps,
)
from backend.app.motors.m05_obligations.instantiation_types import (
    ClientContext,
    GapInput,
    ProjectContext,
)
from backend.tests.conftest import setup_test_project


async def _setup_and_instantiate(db):
    """Create project, set tenant, instantiate org.1 obligations.

    Returns (project_id, outcomes).
    """
    client_id, project_id = await setup_test_project(db)
    await set_tenant_context(
        db,
        client_id=uuid.UUID(client_id),
        project_id=uuid.UUID(project_id),
    )
    ctx = ProjectContext(
        project_id=uuid.UUID(project_id),
        nombre_proyecto="Gantt Test Project",
        categoria_ens="MEDIA",
        cliente=ClientContext(razon_social="Gantt Corp S.L.", sector="tech"),
    )
    gaps = [
        GapInput(gap_id=uuid.uuid4(), measure_code="org.1"),
    ]
    outcomes = await instantiate_obligations_for_multiple_gaps(db, gaps, ctx)
    await db.flush()
    return uuid.UUID(project_id), outcomes


class TestBuildGanttForProject:

    async def test_gantt_for_empty_project(self, db):
        """Empty project returns a plan with zero tasks."""
        client_id, project_id = await setup_test_project(db)
        await set_tenant_context(
            db,
            client_id=uuid.UUID(client_id),
            project_id=uuid.UUID(project_id),
        )
        plan = await build_gantt_for_project(
            db,
            project_id=uuid.UUID(project_id),
            fecha_kickoff=date(2025, 1, 6),
        )
        assert plan.tareas == []
        assert plan.duracion_total_dias_laborables == 0

    async def test_gantt_for_project_with_obligations(self, db):
        """Instantiate org.1, then build Gantt -- tasks should match."""
        project_id, outcomes = await _setup_and_instantiate(db)
        total_created = sum(len(o.obligations_created_ids) for o in outcomes)
        assert total_created > 0

        plan = await build_gantt_for_project(
            db,
            project_id=project_id,
            fecha_kickoff=date(2025, 1, 6),
            dedicacion_horas_semana=8.0,
        )
        assert len(plan.tareas) == total_created
        assert plan.duracion_total_dias_laborables > 0
        assert plan.total_esfuerzo_horas() > 0


class TestExportXlsx:

    async def test_export_gantt_to_xlsx_bytes(self, db):
        """XLSX export produces valid bytes that openpyxl can open."""
        project_id, _ = await _setup_and_instantiate(db)
        plan = await build_gantt_for_project(
            db,
            project_id=project_id,
            fecha_kickoff=date(2025, 1, 6),
        )
        xlsx_bytes = export_gantt_to_xlsx_bytes(plan)

        # PK magic bytes for ZIP (XLSX is a ZIP)
        assert xlsx_bytes[:2] == b"PK"

        # Must open without error
        wb = load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb.active
        assert ws.title == "Gantt Plan"
        # Header row "FULKRO - Plan de Implantacion ENS" in A1
        assert "FULKRO" in str(ws["A1"].value)
