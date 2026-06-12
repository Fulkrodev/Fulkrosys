"""SIMULACIÓN HIPERREALISTA · Cliente ENS BÁSICA (autodeclaración CCN-STIC 809).

Recorre el ciclo de implantación ENS completo de un cliente BÁSICA, lado admin +
cliente, ejercitando el código real (servicios + RLS + BD), e inspecciona cada
entregable que ve el auditor. Vuelca los outputs a out/sim_basica/ y ASEGURA
invariantes (cualquier defecto en el output rompe el test). NO usa navegador
(evita el bucle JWT del front · API/servicio in-process).

Flujo BÁSICA: alta cliente+proyecto -> categorización DICAT (todo BAJO -> BÁSICA)
-> DdA/SoA (52 medidas aplicables) + freeze -> entregables (E-040 Informe Final,
E-160 Manual SGSI, E-170 Plan Director) -> conformidad (ruta declaración ->
CONFORMANT -> distintivo CCN-STIC 809) -> cierre. (BÁSICA: sin auditor externo.)
"""
from __future__ import annotations

import io
import uuid
from pathlib import Path

import pytest
from sqlalchemy import text as sa_text

from backend.app.database import set_tenant_context
from backend.tests.conftest import _admin_setup, setup_test_project

_OUT = Path("/home/usuario/fulkro/out/sim_basica")
_CAT = "BASICA"


def _docx_text_from_path(path: str) -> str:
    from docx import Document as Docx
    d = Docx(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


def _docx_text_from_bio(bio: io.BytesIO) -> str:
    from docx import Document as Docx
    d = Docx(bio)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


def _dump(name: str, text: str) -> None:
    _OUT.mkdir(parents=True, exist_ok=True)
    (_OUT / name).write_text(text, encoding="utf-8")


async def _seed(db, *, cat: str):
    """Alta cliente + proyecto realista + sistema/servicios/info/sedes/roles."""
    client_id, project_id = await setup_test_project(db)
    cid, pid = uuid.UUID(client_id), uuid.UUID(project_id)
    sid = uuid.uuid4()
    # Valoraciones DICAT que producen la categoría objetivo.
    dic = {"D": "BAJO", "I": "BAJO", "C": "BAJO", "A": "BAJO", "T": "BAJO"}
    if cat == "MEDIA":
        dic["D"] = "MEDIO"
    elif cat == "ALTA":
        dic["D"] = "ALTO"
        dic["C"] = "ALTO"
    async with _admin_setup(db):
        await db.execute(sa_text(
            "UPDATE clients SET nombre='Ayuntamiento de Villaverde del Río', "
            "cif='P4109500A', domicilio_fiscal='Plaza de España 1, 41359 Villaverde del Río, Sevilla' "
            "WHERE id=:c"), {"c": client_id})
        await db.execute(sa_text(
            "UPDATE projects SET categoria_objetivo=:cat, nombre='Sede Electrónica Municipal', "
            "fecha_kickoff=current_date - interval '60 days', "
            "fecha_objetivo_certificacion=current_date + interval '30 days' "
            "WHERE id=:p"), {"cat": cat, "p": project_id})
        await db.execute(sa_text(
            "INSERT INTO systems (id, project_id, nombre, created_at) "
            "VALUES (:s,:p,'Sede Electrónica Municipal',now())"),
            {"s": str(sid), "p": project_id})
        await db.execute(sa_text(
            "INSERT INTO services (id, system_id, nombre, tipo, "
            "valoracion_d, valoracion_i, valoracion_c, valoracion_a, valoracion_t, created_at) "
            "VALUES (gen_random_uuid(), :s, 'Tramitación electrónica de expedientes', 'finalista', "
            ":d,:i,:c,:a,:t, now())"),
            {"s": str(sid), "d": dic["D"], "i": dic["I"], "c": dic["C"], "a": dic["A"], "t": dic["T"]})
        await db.execute(sa_text(
            "INSERT INTO information_types (id, system_id, nombre, "
            "valoracion_d, valoracion_i, valoracion_c, valoracion_a, valoracion_t, created_at) "
            "VALUES (gen_random_uuid(), :s, 'Datos de procedimientos administrativos', "
            ":d,:i,:c,:a,:t, now())"),
            {"s": str(sid), "d": dic["D"], "i": dic["I"], "c": dic["C"], "a": dic["A"], "t": dic["T"]})
        await db.execute(sa_text(
            "INSERT INTO system_sites (id, system_id, nombre, tipo, direccion, pais, created_at) "
            "VALUES (gen_random_uuid(), :s, 'Casa Consistorial', 'sede_fisica', "
            "'Plaza de España 1, Villaverde del Río', 'España', now())"), {"s": str(sid)})
        for full_name, role_cat, role_title in [
            ("Antonia Alcalde Ruiz", "sponsor", "Alcaldesa-Presidenta"),
            ("Beatriz Seguridad López", "responsable_seguridad", "Responsable de Seguridad"),
            ("Carlos Sistemas Pérez", "responsable_sistema", "Responsable del Sistema"),
            ("Diego Información Gómez", "responsable_informacion", "Responsable de la Información"),
            ("Elena Servicio Díaz", "responsable_servicio", "Responsable del Servicio"),
        ]:
            await db.execute(sa_text(
                "INSERT INTO client_contacts (id, client_id, full_name, email, "
                "role_category, role_title, is_active, created_at) "
                "VALUES (gen_random_uuid(), :c, :n, :e, :rc, :rt, true, now())"),
                {"c": client_id, "n": full_name, "e": role_cat + "@villaverde.es",
                 "rc": role_cat, "rt": role_title})
    await set_tenant_context(db, client_id=cid, project_id=pid)
    return cid, pid, sid


@pytest.mark.asyncio
async def test_sim_basica_full_lifecycle(db, monkeypatch):
    monkeypatch.setenv("FULKRO_SKIP_WORKFLOW_GATES", "1")  # los gates se validan aparte
    report: list[str] = ["# SIMULACIÓN BÁSICA · Ayuntamiento de Villaverde del Río\n"]

    def log(msg: str) -> None:
        report.append(msg)

    cid, pid, sid = await _seed(db, cat=_CAT)
    log("## 1. Alta cliente + proyecto\nAyuntamiento BÁSICA · Sede Electrónica Municipal · 5 roles ENS.\n")

    # 2. Categorización DICAT
    from backend.app.motors.m01_categorization.service import CategorizationService
    cat_result = await CategorizationService(db).compute_for_system(sid)
    categoria = cat_result.category if hasattr(cat_result, "category") else cat_result.categoria
    log(f"## 2. Categorización DICAT\nCategoría resultante: **{categoria}** (esperado {_CAT}).\n")
    assert str(categoria).upper().endswith(_CAT), f"categoría {categoria} != {_CAT}"
    # Persistir categorización firmada (acta E-012 aprobada).
    async with _admin_setup(db):
        await db.execute(sa_text(
            "INSERT INTO categorizations (id, system_id, categoria_resultante, "
            "aprobado_por, fecha_acta, created_at) "
            "VALUES (gen_random_uuid(), :s, :cat, 'Comité de Seguridad', current_date, now())"),
            {"s": str(sid), "cat": _CAT})

    # 3. DdA / SoA
    from backend.app.motors.m03_dda.enums import CategoriaSistema
    from backend.app.motors.m03_dda.service import DdaService
    dda_svc = DdaService(db)
    await dda_svc.generate_dda(pid, CategoriaSistema(_CAT), responsable="Beatriz Seguridad López",
                              enforce_gates=False)
    counts = (await db.execute(sa_text(
        "SELECT count(*) FILTER (WHERE aplicabilidad='aplica'), count(*) "
        "FROM dda_entries WHERE project_id=:p AND deleted_at IS NULL"), {"p": str(pid)})).first()
    aplicables, total = int(counts[0]), int(counts[1])
    log(f"## 3. Declaración de Aplicabilidad (SoA)\nMedidas aplicables: **{aplicables}** de {total} (Anexo II).\n")
    assert total == 73, f"DdA debe tener 73 entries, tiene {total}"
    assert 45 <= aplicables <= 60, f"BÁSICA aplicables fuera de rango: {aplicables}"
    # marcar como implantadas + congelar
    await db.execute(sa_text(
        "UPDATE dda_entries SET estado_implementacion='implantada', "
        "aprobado_por='Beatriz Seguridad López', fecha_aprobacion=current_date "
        "WHERE project_id=:p AND aplicabilidad='aplica'"), {"p": str(pid)})
    await db.execute(sa_text(
        "UPDATE dda_entries SET aprobado_por='Beatriz Seguridad López', fecha_aprobacion=current_date "
        "WHERE project_id=:p"), {"p": str(pid)})
    await db.flush()

    # 4. Entregables clave
    from backend.app.motors.m06_document_factory.service import DocumentFactoryService
    svc = DocumentFactoryService(db)
    await svc.load_template_metadata_from_catalog()

    from backend.app.motors.m06_document_factory.informe_final_generator import (
        build_informe_final_context,
    )
    e040_ctx = await build_informe_final_context(db, pid)
    e040 = await svc.generate_document(project_id=pid, template_codigo="E-040",
                                       context=e040_ctx, generate_pdf=False, sign=False,
                                       generated_by="sim")
    e040_text = _docx_text_from_path(e040["docx_path"])
    _dump("E-040_informe_final.txt", e040_text)
    log("## 4. Entregables\n### E-040 Informe Final de Adecuación (incluye SoA)\n"
        f"Render OK · {len(e040_text)} chars · cliente: "
        f"{'Villaverde' in e040_text} · sin fugas Jinja: {'{{' not in e040_text}\n")
    assert "Villaverde" in e040_text
    assert "{{" not in e040_text and "{%" not in e040_text
    assert "Declaración de Aplicabilidad" in e040_text or "Anexo II" in e040_text
    cumpl_global = e040_ctx["informe"]["cumplimiento_global"]
    log(f"Cumplimiento global del SoA: **{cumpl_global}%** (medidas implantadas/aplicables).\n")
    assert cumpl_global > 0, "cumplimiento global 0% — estado_implementacion no contabilizado"

    # Rectores E-160 + E-170
    from backend.app.motors.m06_document_factory.rectores_generator import (
        build_rectores_context, generate_manual_sgsi_docx, generate_plan_director_docx,
    )
    rctx = await build_rectores_context(db, pid)
    e160_text = _docx_text_from_bio(generate_manual_sgsi_docx(rctx))
    e170_text = _docx_text_from_bio(generate_plan_director_docx(rctx))
    _dump("E-160_manual_sgsi.txt", e160_text)
    _dump("E-170_plan_director.txt", e170_text)
    log("### E-160 Manual SGSI + E-170 Plan Director\n"
        f"Firmante RSEG presente E-160: {'Beatriz Seguridad López' in e160_text} · "
        f"E-170: {'Beatriz Seguridad López' in e170_text}\n")
    assert "Beatriz Seguridad López" in e160_text
    assert "Beatriz Seguridad López" in e170_text

    # 5. Conformidad BÁSICA: ruta declaración -> CONFORMANT -> distintivo CCN-STIC 809
    from backend.app.models.conformity_lifecycle import ConformityRouteRow
    route = ConformityRouteRow(project_id=pid, route_type="declaracion_basica", status="CONFORMANT")
    db.add(route)
    await db.flush()
    from backend.app.motors.m27_conformity.distintivo_persistence import (
        attach_distintivo_on_registered, read_distintivo_bytes,
    )
    dist = await attach_distintivo_on_registered(db, pid, generated_by="sim")
    dist_bytes, dist_name = await read_distintivo_bytes(db, pid)
    dist_text = _docx_text_from_bio(io.BytesIO(dist_bytes))
    _dump("E-049_distintivo.txt", dist_text)
    log("## 5. Conformidad BÁSICA (autodeclaración CCN-STIC 809)\n"
        f"Ruta -> {dist['route_state']} · distintivo emitido · fichero {dist_name}\n")
    assert dist["route_state"] == "REGISTERED"
    assert "CCN-STIC 809" in dist_text
    assert "autoevaluación" in dist_text  # BÁSICA = autoevaluación
    assert "Villaverde" in dist_text

    # 6. Dossier
    log("## 6. Dossier de auditoría\n(pendiente de inspección · ver out/sim_basica/)\n")

    _dump("REPORT.md", "\n".join(report))
    # Sanity final
    assert (_OUT / "E-040_informe_final.txt").exists()
