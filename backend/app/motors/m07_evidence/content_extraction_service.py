"""Extracción de contenido de los documentos que sube el cliente (lectura IA).

Cuando el cliente sube documentación dentro de una tarea, el sistema LEE el
contenido (no sólo lo almacena): así la evidencia es buscable, el clasificador
ENS puede sugerir la medida/familia a partir del CONTENIDO (no sólo del nombre
de fichero), y el copiloto/auditor pueden razonar sobre lo que el cliente aportó.

Formatos soportados (deps ya presentes · pdfplumber/python-docx/openpyxl):
- PDF con capa de texto → pdfplumber.
- Word .docx → python-docx (párrafos + tablas).
- Excel .xlsx/.xlsm → openpyxl (celdas no vacías).
- CSV / TXT / MD / JSON / LOG / HTML → decodificación directa.
- Imágenes (png/jpg/tiff/…) y PDF ESCANEADO (sin capa de texto) → OCR
  best-effort vía pytesseract (+ pdf2image para PDF). Estos NO son deps duros:
  si el binario tesseract / las libs no están, degrada con gracia
  (``ocr_available=False``) en vez de fallar — el resto de formatos sigue OK.

Política anti-alucinación: si un formato no se puede leer, devuelve texto vacío
con ``method``/``error`` explícitos. NUNCA inventa contenido. El texto se cap-ea
a ``MAX_CHARS`` para no inflar la BD ni el prompt del clasificador.
"""
from __future__ import annotations

import csv
import io
import logging
from typing import Any

logger = logging.getLogger(__name__)

MAX_CHARS = 20_000


def _result(text: str, method: str, *, ocr_used: bool = False,
            ocr_available: bool = True, error: str | None = None) -> dict[str, Any]:
    clean = (text or "").strip()
    if len(clean) > MAX_CHARS:
        clean = clean[:MAX_CHARS] + "\n…[truncado]"
    return {
        "text": clean,
        "chars": len(clean),
        "method": method,
        "ocr_used": ocr_used,
        "ocr_available": ocr_available,
        "error": error,
    }


def _decode(file_bytes: bytes) -> str:
    for enc in ("utf-8", "latin-1"):
        try:
            return file_bytes.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return file_bytes.decode("utf-8", errors="replace")


def _extract_pdf(file_bytes: bytes) -> dict[str, Any]:
    import pdfplumber  # dep presente

    parts: list[str] = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            parts.append(page.extract_text() or "")
    text = "\n".join(parts).strip()
    if text:
        return _result(text, "pdf-text")
    # PDF sin capa de texto (escaneado) → OCR best-effort
    return _ocr_pdf(file_bytes)


def _extract_docx(file_bytes: bytes) -> dict[str, Any]:
    import docx  # python-docx (dep presente)

    d = docx.Document(io.BytesIO(file_bytes))
    parts = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return _result("\n".join(parts), "docx")


def _extract_xlsx(file_bytes: bytes) -> dict[str, Any]:
    import openpyxl  # dep presente

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"# Hoja: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v not in (None, "")]
            if vals:
                parts.append(" | ".join(vals))
    wb.close()
    return _result("\n".join(parts), "xlsx")


def _ocr_image(file_bytes: bytes) -> dict[str, Any]:
    try:
        import pytesseract  # NO es dep dura · lazy
        from PIL import Image
    except ImportError:
        return _result("", "image-ocr-unavailable", ocr_available=False)
    try:
        img = Image.open(io.BytesIO(file_bytes))
        text = pytesseract.image_to_string(img, lang="spa+eng")
        return _result(text, "image-ocr", ocr_used=True)
    except Exception as e:  # binario tesseract ausente, etc.
        logger.warning("OCR imagen fallo: %s", e)
        return _result("", "image-ocr-failed", ocr_available=False, error=str(e))


def _ocr_pdf(file_bytes: bytes) -> dict[str, Any]:
    try:
        import pytesseract  # lazy
        from pdf2image import convert_from_bytes
    except ImportError:
        return _result("", "pdf-scanned-ocr-unavailable", ocr_available=False)
    try:
        images = convert_from_bytes(file_bytes, dpi=200, fmt="png")
        parts = [pytesseract.image_to_string(im, lang="spa+eng") for im in images[:30]]
        return _result("\n".join(parts), "pdf-ocr", ocr_used=True)
    except Exception as e:
        logger.warning("OCR PDF escaneado fallo: %s", e)
        return _result("", "pdf-ocr-failed", ocr_available=False, error=str(e))


def extract_text(file_bytes: bytes, mime_type: str | None, filename: str) -> dict[str, Any]:
    """Extrae el texto de un documento subido. Devuelve dict (nunca lanza).

    Keys: text, chars, method, ocr_used, ocr_available, error.
    """
    if not file_bytes:
        return _result("", "empty")
    ext = filename.rsplit(".", 1)[-1].lower() if "." in (filename or "") else ""
    mime = (mime_type or "").lower()
    try:
        if "pdf" in mime or ext == "pdf":
            return _extract_pdf(file_bytes)
        if "wordprocessing" in mime or ext == "docx":
            return _extract_docx(file_bytes)
        if "spreadsheet" in mime or ext in ("xlsx", "xlsm"):
            return _extract_xlsx(file_bytes)
        if ext == "csv" or "csv" in mime:
            return _result(_decode(file_bytes), "csv")
        if mime.startswith("text/") or ext in (
            "txt", "md", "json", "log", "html", "htm", "xml", "yaml", "yml",
        ):
            return _result(_decode(file_bytes), "text")
        if mime.startswith("image/") or ext in (
            "png", "jpg", "jpeg", "tif", "tiff", "bmp", "webp", "gif",
        ):
            return _ocr_image(file_bytes)
        # desconocido: intento decodificar como texto; si es binario, vacío
        text = _decode(file_bytes)
        printable = sum(c.isprintable() or c.isspace() for c in text[:2000])
        if text and printable / max(len(text[:2000]), 1) > 0.85:
            return _result(text, "fallback-decode")
        return _result("", "unsupported-binary")
    except Exception as e:  # pragma: no cover — defensivo
        logger.exception("extract_text fallo · file=%s mime=%s", filename, mime)
        return _result("", "error", error=str(e))
