"""RoPA Service — Record of Processing Activities (Art. 30 GDPR).

Manages the ``fulkro_ropa_treatments`` registry: list, single fetch,
field updates, mark-as-reviewed, and Excel export in AEPD-compatible
layout (one row per treatment with the columns required by the
Spanish supervisory authority audit guide).

The service is admin-only. Reads return all rows including processor
sub-rows; the Trust Center public API (atom 9.bis.5) reads a derived
``sub_processors_count`` from this table.
"""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from backend.app.models.ropa_treatments import FulkroRoPATreatment


_EXCEL_COLUMNS: list[tuple[str, str, int]] = [
    ("treatment_code", "Código", 12),
    ("treatment_name", "Nombre del tratamiento", 38),
    ("role", "Rol FULKRO", 12),
    ("purpose", "Finalidad (Art. 30.1.b)", 60),
    ("legal_basis", "Base jurídica (Art. 6)", 28),
    ("data_categories", "Categorías de datos (Art. 30.1.c)", 40),
    (
        "data_subjects_categories",
        "Categorías de interesados (Art. 30.1.c)",
        32,
    ),
    ("recipients", "Destinatarios (Art. 30.1.d)", 36),
    ("transfers_outside_eu", "Transferencias fuera UE (Art. 30.1.e)", 16),
    ("transfer_safeguards", "Garantías de transferencia", 48),
    ("retention_period", "Plazo de conservación (Art. 30.1.f)", 32),
    ("security_measures", "Medidas técnicas y organizativas (Art. 32)", 60),
    ("processor_name", "Sub-encargado (si aplica)", 28),
    ("is_sub_processor", "¿Implica sub-encargado?", 14),
    ("dpa_signed", "DPA firmado", 12),
    ("controller_dpo", "DPO responsable", 40),
    ("last_reviewed_at", "Última revisión", 20),
]


class RoPAService:
    """Service for managing the Article 30 RoPA registry."""

    def __init__(self, db: AsyncSession):
        self.db = db

    async def list_treatments(self) -> list[FulkroRoPATreatment]:
        """Return all treatments ordered by treatment_code."""
        rows = (
            await self.db.execute(
                select(FulkroRoPATreatment).order_by(
                    FulkroRoPATreatment.treatment_code
                )
            )
        ).scalars().all()
        return list(rows)

    async def get_treatment(self, code: str) -> FulkroRoPATreatment | None:
        return (
            await self.db.execute(
                select(FulkroRoPATreatment).where(
                    FulkroRoPATreatment.treatment_code == code
                )
            )
        ).scalar_one_or_none()

    async def update_treatment(
        self, code: str, fields: dict[str, Any]
    ) -> FulkroRoPATreatment:
        """Patch updatable fields on a treatment, bumping last_reviewed_at."""
        row = await self.get_treatment(code)
        if row is None:
            raise KeyError(f"Treatment {code} not found")

        allowed = {
            "treatment_name",
            "role",
            "purpose",
            "legal_basis",
            "data_categories",
            "data_subjects_categories",
            "recipients",
            "transfers_outside_eu",
            "transfer_safeguards",
            "retention_period",
            "security_measures",
            "processor_name",
            "is_sub_processor",
            "dpa_signed",
            "dpa_expires_at",
            "controller_dpo",
        }
        for key, value in fields.items():
            if key not in allowed:
                raise ValueError(f"Field '{key}' is not updatable")
            setattr(row, key, value)
        row.last_reviewed_at = datetime.now(timezone.utc)
        await self.db.flush()
        return row

    async def mark_reviewed(self, code: str) -> FulkroRoPATreatment:
        row = await self.get_treatment(code)
        if row is None:
            raise KeyError(f"Treatment {code} not found")
        row.last_reviewed_at = datetime.now(timezone.utc)
        await self.db.flush()
        return row

    async def count_treatments(self) -> int:
        return (
            (await self.db.execute(select(func.count(FulkroRoPATreatment.id)))).scalar()
            or 0
        )

    async def count_sub_processors(self) -> int:
        return (
            (
                await self.db.execute(
                    select(func.count(FulkroRoPATreatment.id)).where(
                        FulkroRoPATreatment.is_sub_processor.is_(True)
                    )
                )
            ).scalar()
            or 0
        )

    async def export_aepd_excel(self) -> BytesIO:
        """Render the full RoPA into an XLSX file in AEPD-friendly format.

        Layout:
        - Sheet 1 (RoPA): one row per treatment, every Art. 30 column
        - Sheet 2 (Metadatos): responsable / DPO / domicilio / fecha export
        """
        treatments = await self.list_treatments()

        # Identidad fiscal FULKRO (FUENTE ÚNICA · punto #44).
        from backend.app.core.fiscal_identity import get_fiscal_identity
        fi = await get_fiscal_identity(self.db)

        wb = Workbook()
        ws = wb.active
        ws.title = "RoPA"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="1F2937")
        wrap = Alignment(wrap_text=True, vertical="top")

        for col_idx, (_, label, width) in enumerate(_EXCEL_COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(
                wrap_text=True, vertical="center", horizontal="center"
            )
            ws.column_dimensions[cell.column_letter].width = width
        ws.row_dimensions[1].height = 32

        for row_idx, t in enumerate(treatments, start=2):
            for col_idx, (attr, _, _) in enumerate(_EXCEL_COLUMNS, start=1):
                value = getattr(t, attr)
                if isinstance(value, list):
                    value = "\n".join(f"• {item}" for item in value)
                elif isinstance(value, bool):
                    value = "Sí" if value else "No"
                elif isinstance(value, datetime):
                    value = value.strftime("%Y-%m-%d %H:%M")
                cell = ws.cell(row=row_idx, column=col_idx, value=value or "")
                cell.alignment = wrap

        # Metadatos sheet
        meta = wb.create_sheet("Metadatos")
        meta_rows: list[tuple[str, str]] = [
            (
                "Responsable del tratamiento",
                f"{fi.nombre_fiscal or 'Marcos Mata García'} "
                f"({fi.nombre_comercial or 'FULKRO'})",
            ),
            ("CIF / NIF", fi.nif or "[pendiente · configurar en /admin/settings/fiscal]"),
            ("Domicilio", fi.domicilio_completo or "Madrid, España"),
            ("DPO", f"{fi.nombre_fiscal or 'Marcos Mata García'} · interim"),
            ("Email DPO", "dpo@fulkro.es"),
            ("Actividad principal", "Consultoría servicios IT · ENS"),
            ("Fecha exportación", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
            ("Base regulatoria", "Art. 30 GDPR + LOPDGDD 3/2018"),
            ("Total tratamientos", str(len(treatments))),
        ]
        for idx, (label, value) in enumerate(meta_rows, start=1):
            label_cell = meta.cell(row=idx, column=1, value=label)
            label_cell.font = Font(bold=True)
            meta.cell(row=idx, column=2, value=value)
        meta.column_dimensions["A"].width = 32
        meta.column_dimensions["B"].width = 50

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf
