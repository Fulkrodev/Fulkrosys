"""Motor 5 -- Gantt Service -- async DB bridge + XLSX export.

Connects the pure ``gantt_planner`` algorithm to the database and
provides an XLSX export via openpyxl.
"""
from __future__ import annotations

import io
import uuid
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from backend.app.models.ens import Obligation
from backend.app.motors.m05_obligations.gantt_planner import build_gantt_plan
from backend.app.motors.m05_obligations.gantt_types import GanttPlan


async def build_gantt_for_project(
    session: AsyncSession,
    project_id: uuid.UUID,
    fecha_kickoff: date,
    dedicacion_horas_semana: float = 8.0,
) -> GanttPlan:
    """Query obligations from DB and build a Gantt plan."""
    stmt = (
        select(Obligation)
        .where(
            Obligation.project_id == project_id,
            Obligation.deleted_at.is_(None),
        )
        .order_by(Obligation.created_at)
    )
    result = await session.execute(stmt)
    rows = result.scalars().all()

    obligation_dicts = [
        {
            "obligation_id": row.id,
            "template_id": row.template_id or str(row.id),
            "measure_code": row.measure_code or "",
            "titulo": row.titulo or row.descripcion[:80],
            "modo_ejecucion": row.modo_ejecucion,
            "estado": row.estado,
            "esfuerzo_estimado": row.esfuerzo_estimado or 0.0,
            "dependencias_template_ids": (
                row.dependencias_template_ids
                if isinstance(row.dependencias_template_ids, list)
                else []
            ),
        }
        for row in rows
    ]

    return build_gantt_plan(
        project_id=project_id,
        obligations=obligation_dicts,
        fecha_kickoff=fecha_kickoff,
        dedicacion_horas_semana=dedicacion_horas_semana,
    )


def export_gantt_to_xlsx_bytes(plan: GanttPlan) -> bytes:
    """Create an XLSX workbook from a GanttPlan and return raw bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Gantt Plan"

    # ── Header metadata rows ──────────────────────────────────────────
    header_font = Font(bold=True, size=12)
    ws.append(["FULKRO - Plan de Implantacion ENS"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Proyecto", str(plan.project_id)])
    ws.append(["Fecha kickoff", plan.fecha_kickoff.isoformat()])
    ws.append(["Fecha fin estimada", plan.fecha_fin_estimada.isoformat()])
    ws.append(["Duracion total (dias lab.)", plan.duracion_total_dias_laborables])
    ws.append(["Dedicacion (h/semana)", plan.dedicacion_cliente_horas_semana])
    ws.append(["Esfuerzo total (horas)", plan.total_esfuerzo_horas()])
    ws.append([])  # blank separator

    # ── Table headers ─────────────────────────────────────────────────
    columns = [
        "Template ID",
        "Medida",
        "Titulo",
        "Modo Ejecucion",
        "Estado",
        "Esfuerzo (h)",
        "Fecha Inicio",
        "Fecha Fin",
        "Duracion (dias lab.)",
        "Predecesores",
    ]
    ws.append(columns)
    header_row = ws.max_row
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font_white = Font(bold=True, color="FFFFFF", size=11)
    for col_idx, _ in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font_white
        cell.alignment = Alignment(horizontal="center")

    # ── Data rows ─────────────────────────────────────────────────────
    for t in plan.tareas:
        ws.append([
            t.template_id,
            t.measure_code,
            t.titulo,
            t.modo_ejecucion or "",
            t.estado or "",
            t.esfuerzo_horas,
            t.fecha_inicio.isoformat(),
            t.fecha_fin.isoformat(),
            t.duracion_dias_laborables,
            ", ".join(t.predecesores_template_ids) if t.predecesores_template_ids else "",
        ])

    # ── Auto-width columns ────────────────────────────────────────────
    for col_cells in ws.columns:
        max_len = 0
        col_letter = col_cells[0].column_letter
        for cell in col_cells:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 3, 50)

    # Serialise to bytes
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
