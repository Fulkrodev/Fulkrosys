"""Planning service (M17).

Genera planes desde catalogo WBS + CPM (forward+backward pass) + replan.
"""
from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timedelta, timezone
from typing import Optional

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from backend.app.models.planning import (
    ChangeRequest,
    ProjectPlan,
    WbsTask,
)

from .effort_estimator import (
    VALID_CATEGORIAS,
    VALID_COMPLEXITY,
    VALID_SIZES,
    estimate_duration_weeks,
    estimate_effort,
)
from .wbs_catalog import get_tasks_for_categoria


class PlanningError(ValueError):
    pass


VALID_TASK_STATUS = {
    "por_hacer", "en_curso", "bloqueada",
    "en_revision", "hecha", "descartada",
}


DEFAULT_MEETING_PLAN = {
    "weekly": {"day": "friday", "time": "16:00", "duration_min": 30},
    "monthly": {"week": 1, "day": "tuesday", "duration_min": 90},
    "quarterly": {"month_offset": 0, "week": 1, "duration_min": 60},
}


# =============== Plan generation ===============

async def generate_plan(
    db: AsyncSession,
    project_id: uuid.UUID,
    categoria: str,
    start_date: date,
    client_size: str = "mediana",
    complexity: str = "media",
    marcos_weekly_hours: float = 20.0,
    client_weekly_hours: float = 8.0,
) -> ProjectPlan:
    cat = (categoria or "").strip().upper()
    if cat not in VALID_CATEGORIAS:
        raise PlanningError(
            f"Categoria '{categoria}' invalida. Validas: {sorted(VALID_CATEGORIAS)}"
        )
    if client_size not in VALID_SIZES:
        raise PlanningError(
            f"client_size '{client_size}' invalido. Validos: {sorted(VALID_SIZES)}"
        )
    if complexity not in VALID_COMPLEXITY:
        raise PlanningError(
            f"complexity '{complexity}' invalido. Validos: {sorted(VALID_COMPLEXITY)}"
        )

    # Check: evitar duplicados por proyecto
    r = await db.execute(
        select(ProjectPlan).where(
            ProjectPlan.project_id == project_id,
            ProjectPlan.deleted_at.is_(None),
        )
    )
    existing = r.scalar_one_or_none()
    if existing is not None:
        raise PlanningError(
            f"Proyecto {project_id} ya tiene plan. Borra el anterior o usa replan."
        )

    plan = ProjectPlan(
        project_id=project_id,
        version=1,
        categoria=cat,
        start_date=start_date,
        estado="draft",
        marcos_weekly_capacity_hours=marcos_weekly_hours,
        client_weekly_capacity_hours=client_weekly_hours,
        total_duration_weeks=estimate_duration_weeks(cat),
        meeting_plan=dict(DEFAULT_MEETING_PLAN),
    )
    db.add(plan)
    await db.flush()

    # Instantiate tasks
    templates = get_tasks_for_categoria(cat)
    tasks_by_code: dict[str, WbsTask] = {}
    for t in templates:
        eff_m = estimate_effort(t.effort_marcos_h, cat, client_size, complexity)
        eff_p = estimate_effort(t.effort_platform_h, cat, client_size, complexity)
        task = WbsTask(
            project_plan_id=plan.id,
            project_id=project_id,
            task_code=t.code,
            task_name=t.name,
            phase=str(t.phase),
            duration_days=t.duration_days,
            effort_marcos_hours=eff_m,
            effort_platform_hours=eff_p,
            responsible=t.responsible,
            dependencies=list(t.dependencies),
            deliverable_e_code=t.deliverable,
            status="por_hacer",
            progress_pct=0,
            is_critical_path=False,
        )
        db.add(task)
        await db.flush()
        tasks_by_code[t.code] = task

    # Calcular fechas (forward pass por dependencias)
    _forward_pass_dates(tasks_by_code, start_date)

    # CPM: forward + backward + slack
    _calculate_critical_path(tasks_by_code)

    # Totales + end_date estimado
    total_marcos = sum(t.effort_marcos_hours or 0 for t in tasks_by_code.values())
    total_platform = sum(
        t.effort_platform_hours or 0 for t in tasks_by_code.values()
    )
    max_end = max(
        (t.end_date for t in tasks_by_code.values() if t.end_date),
        default=start_date,
    )
    plan.total_effort_marcos_hours = round(total_marcos, 1)
    plan.total_effort_platform_hours = round(total_platform, 1)
    plan.end_date_estimated = max_end
    plan.critical_path_tasks = [
        t.task_code for t in tasks_by_code.values() if t.is_critical_path
    ]
    # Ruta critica longitud en semanas (suma duraciones criticas)
    cp_days = sum(
        (t.duration_days or 0)
        for t in tasks_by_code.values() if t.is_critical_path
    )
    plan.critical_path_length_weeks = max(1, cp_days // 7)

    # Mermaid
    plan.mermaid_gantt = _build_mermaid_gantt(
        tasks_by_code, plan_name=f"Plan ENS {cat}",
    )
    plan.estado = "active"
    await db.flush()
    return plan


async def replan(
    db: AsyncSession,
    project_id: uuid.UUID,
    new_categoria: str,
    client_size: str = "mediana",
    complexity: str = "media",
) -> ProjectPlan:
    """#5 cabo N2 · regenera el plan a una nueva categoría preservando
    ``start_date`` + capacidades del plan vigente. Soft-delete del plan anterior
    + sus tareas (traza · NO borrado físico) y genera uno nuevo con la categoría
    nueva. Lo invoca el orquestador del suelo SOLO en N2 (borradores · sin
    contrato firmado); nunca toca planes de proyectos con firma protegida.
    """
    r = await db.execute(
        select(ProjectPlan).where(
            ProjectPlan.project_id == project_id,
            ProjectPlan.deleted_at.is_(None),
        )
    )
    old = r.scalar_one_or_none()
    if old is None:
        raise PlanningError(
            f"Proyecto {project_id} no tiene plan que regenerar."
        )
    start_date = old.start_date or date.today()
    m_h = old.marcos_weekly_capacity_hours or 20.0
    c_h = old.client_weekly_capacity_hours or 8.0

    now = datetime.now(timezone.utc)
    old.deleted_at = now
    tasks_r = await db.execute(
        select(WbsTask).where(
            WbsTask.project_plan_id == old.id,
            WbsTask.deleted_at.is_(None),
        )
    )
    for t in tasks_r.scalars().all():
        t.deleted_at = now
    await db.flush()

    return await generate_plan(
        db,
        project_id,
        new_categoria,
        start_date,
        client_size=client_size,
        complexity=complexity,
        marcos_weekly_hours=m_h,
        client_weekly_hours=c_h,
    )


def _forward_pass_dates(
    tasks_by_code: dict[str, WbsTask], start_date: date,
) -> None:
    """Asigna start_date/end_date respetando dependencias via orden topologico."""
    # Topological sort
    order: list[str] = []
    visited: set[str] = set()
    visiting: set[str] = set()

    def _visit(code: str) -> None:
        if code in visited:
            return
        if code in visiting:
            return  # ciclo — lo ignoramos
        visiting.add(code)
        task = tasks_by_code.get(code)
        if task is not None:
            for dep in (task.dependencies or []):
                if dep in tasks_by_code:
                    _visit(dep)
        visiting.discard(code)
        visited.add(code)
        order.append(code)

    for code in tasks_by_code:
        _visit(code)

    for code in order:
        task = tasks_by_code[code]
        deps = task.dependencies or []
        if not deps:
            task.start_date = start_date
        else:
            dep_ends: list[date] = []
            for dep_code in deps:
                dep = tasks_by_code.get(dep_code)
                if dep and dep.end_date:
                    dep_ends.append(dep.end_date)
            task.start_date = (
                max(dep_ends) + timedelta(days=1) if dep_ends else start_date
            )
        dur = task.duration_days or 0
        task.end_date = task.start_date + timedelta(days=max(0, dur - 1))


def _calculate_critical_path(tasks_by_code: dict[str, WbsTask]) -> None:
    """CPM estandar:
       1. Forward pass: ES, EF
       2. Backward pass: LS, LF
       3. Slack = LS - ES
       4. Critical path: slack == 0
    """
    if not tasks_by_code:
        return

    # Topological order
    order: list[str] = []
    visited: set[str] = set()
    visiting: set[str] = set()

    def _visit(code: str) -> None:
        if code in visited:
            return
        if code in visiting:
            return
        visiting.add(code)
        task = tasks_by_code.get(code)
        if task is not None:
            for dep in (task.dependencies or []):
                if dep in tasks_by_code:
                    _visit(dep)
        visiting.discard(code)
        visited.add(code)
        order.append(code)

    for code in tasks_by_code:
        _visit(code)

    # Forward pass (ES = max EF de preds, EF = ES + dur)
    ES: dict[str, int] = {}
    EF: dict[str, int] = {}
    for code in order:
        task = tasks_by_code[code]
        deps = task.dependencies or []
        es = 0
        for d in deps:
            if d in EF:
                es = max(es, EF[d])
        ES[code] = es
        EF[code] = es + (task.duration_days or 0)

    project_end = max(EF.values(), default=0)

    # Sucesores para backward pass
    successors: dict[str, list[str]] = {c: [] for c in tasks_by_code}
    for c, t in tasks_by_code.items():
        for d in (t.dependencies or []):
            if d in successors:
                successors[d].append(c)

    # Backward pass
    LS: dict[str, int] = {}
    LF: dict[str, int] = {}
    for code in reversed(order):
        task = tasks_by_code[code]
        succ = successors.get(code, [])
        lf = project_end if not succ else min(LS[s] for s in succ if s in LS)
        LF[code] = lf
        LS[code] = lf - (task.duration_days or 0)

    # Slack + critical path
    for code, task in tasks_by_code.items():
        slack = LS[code] - ES[code]
        task.slack_days = slack
        task.is_critical_path = (slack == 0 and (task.duration_days or 0) > 0)


def _build_mermaid_gantt(
    tasks_by_code: dict[str, WbsTask], plan_name: str,
) -> str:
    """Genera Mermaid gantt desde tareas."""
    lines = [
        "gantt",
        f"    title {plan_name}",
        "    dateFormat YYYY-MM-DD",
    ]
    # Agrupar por fase
    by_phase: dict[str, list[WbsTask]] = {}
    for t in tasks_by_code.values():
        by_phase.setdefault(t.phase or "0", []).append(t)
    for phase in sorted(by_phase.keys()):
        lines.append(f"    section Fase {phase}")
        for t in sorted(by_phase[phase], key=lambda x: x.task_code):
            if not t.start_date or not t.end_date:
                continue
            dur = (t.end_date - t.start_date).days + 1
            crit = "crit, " if t.is_critical_path else ""
            safe_name = (t.task_name or "").replace(":", " -").replace("\n", " ")
            lines.append(
                f"    {safe_name} :{crit}{t.task_code}, "
                f"{t.start_date.isoformat()}, {dur}d"
            )
    return "\n".join(lines)


# =============== Seguimiento ===============

async def update_task_status(
    db: AsyncSession,
    task_id: uuid.UUID,
    status: str,
    progress_pct: Optional[int] = None,
    blocker: Optional[str] = None,
) -> WbsTask:
    if status not in VALID_TASK_STATUS:
        raise PlanningError(
            f"status '{status}' invalido. Validos: {sorted(VALID_TASK_STATUS)}"
        )
    r = await db.execute(
        select(WbsTask).where(
            WbsTask.id == task_id, WbsTask.deleted_at.is_(None),
        )
    )
    task = r.scalar_one_or_none()
    if task is None:
        raise PlanningError(f"WbsTask {task_id} no encontrado")
    task.status = status
    if progress_pct is not None:
        task.progress_pct = max(0, min(100, int(progress_pct)))
    if status == "hecha":
        task.progress_pct = 100
    if status == "bloqueada" and blocker:
        task.blocker_description = blocker
    await db.flush()
    return task


async def set_baseline(
    db: AsyncSession, plan_id: uuid.UUID,
) -> ProjectPlan:
    r = await db.execute(
        select(ProjectPlan).where(
            ProjectPlan.id == plan_id, ProjectPlan.deleted_at.is_(None),
        )
    )
    plan = r.scalar_one_or_none()
    if plan is None:
        raise PlanningError(f"ProjectPlan {plan_id} no encontrado")
    if plan.baseline_date is not None:
        raise PlanningError("Baseline ya congelada")

    # Snapshot de tareas
    r = await db.execute(
        select(WbsTask).where(WbsTask.project_plan_id == plan_id)
    )
    tasks = list(r.scalars().all())
    snapshot = {
        "plan_id": str(plan_id),
        "created_at": datetime.now(timezone.utc).isoformat(),
        "tasks": [
            {
                "task_code": t.task_code,
                "start_date": t.start_date.isoformat() if t.start_date else None,
                "end_date": t.end_date.isoformat() if t.end_date else None,
                "duration_days": t.duration_days,
                "is_critical_path": t.is_critical_path,
            }
            for t in tasks
        ],
    }
    plan.baseline_snapshot = snapshot
    plan.baseline_date = datetime.now(timezone.utc)
    for t in tasks:
        t.baseline_start_date = t.start_date
        t.baseline_end_date = t.end_date
    await db.flush()
    return plan


async def detect_delays(
    db: AsyncSession, plan_id: uuid.UUID,
) -> list[dict]:
    r = await db.execute(
        select(WbsTask).where(
            WbsTask.project_plan_id == plan_id,
            WbsTask.deleted_at.is_(None),
            WbsTask.baseline_end_date.isnot(None),
        )
    )
    tasks = list(r.scalars().all())
    delays: list[dict] = []
    for t in tasks:
        if t.end_date is None or t.baseline_end_date is None:
            continue
        delay = (t.end_date - t.baseline_end_date).days
        if delay > 0:
            delays.append({
                "task_code": t.task_code,
                "task_name": t.task_name,
                "baseline_end": t.baseline_end_date.isoformat(),
                "current_end": t.end_date.isoformat(),
                "delay_days": delay,
                "is_critical": bool(t.is_critical_path),
            })
    delays.sort(key=lambda x: (-x["delay_days"], x["task_code"]))
    return delays


async def get_plan_progress(
    db: AsyncSession, plan_id: uuid.UUID,
) -> dict:
    r = await db.execute(
        select(WbsTask).where(
            WbsTask.project_plan_id == plan_id,
            WbsTask.deleted_at.is_(None),
        )
    )
    tasks = list(r.scalars().all())
    total = len(tasks)
    by_status: dict[str, int] = {}
    total_progress = 0
    for t in tasks:
        s = t.status or "por_hacer"
        by_status[s] = by_status.get(s, 0) + 1
        total_progress += (t.progress_pct or 0)
    avg_progress = round(total_progress / total, 1) if total else 0.0
    done = by_status.get("hecha", 0)
    pct_done = round(100.0 * done / total, 1) if total else 0.0
    return {
        "plan_id": str(plan_id),
        "total_tasks": total,
        "by_status": by_status,
        "tasks_done": done,
        "pct_done": pct_done,
        "avg_progress_pct": avg_progress,
    }


# =============== Change Requests ===============

async def create_change_request(
    db: AsyncSession,
    project_id: uuid.UUID,
    plan_id: uuid.UUID,
    titulo: str,
    descripcion: str,
    impacto_plazo_dias: Optional[int] = None,
    impacto_esfuerzo_horas: Optional[float] = None,
    impacto_presupuesto_eur: Optional[float] = None,
    solicitado_por: str = "marcos",
) -> ChangeRequest:
    # FIX P1-8: código CR serializado (advisory lock compartido m17+m19 · evita
    # CR-NNN duplicados bajo concurrencia · unicidad en BD como red de seguridad).
    from backend.app.core.sequences import next_change_request_code
    code = await next_change_request_code(db, project_id)
    cr = ChangeRequest(
        project_id=project_id,
        plan_id=plan_id,
        code=code,
        titulo=titulo,
        descripcion=descripcion,
        impacto_plazo_dias=impacto_plazo_dias,
        impacto_esfuerzo_horas=impacto_esfuerzo_horas,
        impacto_presupuesto_eur=impacto_presupuesto_eur,
        estado="propuesto",
        solicitado_por=solicitado_por,
    )
    db.add(cr)
    await db.flush()
    return cr


async def approve_change_request(
    db: AsyncSession, cr_id: uuid.UUID, aprobado_por: str,
) -> ChangeRequest:
    r = await db.execute(
        select(ChangeRequest).where(
            ChangeRequest.id == cr_id, ChangeRequest.deleted_at.is_(None),
        )
    )
    cr = r.scalar_one_or_none()
    if cr is None:
        raise PlanningError(f"ChangeRequest {cr_id} no encontrado")
    if cr.estado != "propuesto":
        raise PlanningError(
            f"CR ya en estado '{cr.estado}'. Solo 'propuesto' es aprobable."
        )
    cr.estado = "aprobado"
    cr.aprobado_por = aprobado_por
    cr.aprobado_at = datetime.now(timezone.utc)

    # Aplicar impacto al plan
    if cr.impacto_plazo_dias:
        r2 = await db.execute(
            select(ProjectPlan).where(ProjectPlan.id == cr.plan_id)
        )
        plan = r2.scalar_one_or_none()
        if plan is not None and plan.end_date_estimated is not None:
            plan.end_date_estimated = (
                plan.end_date_estimated + timedelta(days=cr.impacto_plazo_dias)
            )
    if cr.impacto_esfuerzo_horas:
        r2 = await db.execute(
            select(ProjectPlan).where(ProjectPlan.id == cr.plan_id)
        )
        plan = r2.scalar_one_or_none()
        if plan is not None:
            plan.total_effort_platform_hours = (
                (plan.total_effort_platform_hours or 0)
                + float(cr.impacto_esfuerzo_horas)
            )
    await db.flush()
    return cr


async def reject_change_request(
    db: AsyncSession, cr_id: uuid.UUID, rechazado_por: str,
) -> ChangeRequest:
    r = await db.execute(
        select(ChangeRequest).where(
            ChangeRequest.id == cr_id, ChangeRequest.deleted_at.is_(None),
        )
    )
    cr = r.scalar_one_or_none()
    if cr is None:
        raise PlanningError(f"ChangeRequest {cr_id} no encontrado")
    cr.estado = "rechazado"
    cr.aprobado_por = rechazado_por
    cr.aprobado_at = datetime.now(timezone.utc)
    await db.flush()
    return cr


# =============== Queries ===============

async def get_plan(
    db: AsyncSession, project_id: uuid.UUID,
) -> Optional[ProjectPlan]:
    r = await db.execute(
        select(ProjectPlan).where(
            ProjectPlan.project_id == project_id,
            ProjectPlan.deleted_at.is_(None),
        )
    )
    return r.scalar_one_or_none()


async def get_plan_by_id(
    db: AsyncSession, plan_id: uuid.UUID,
) -> Optional[ProjectPlan]:
    r = await db.execute(
        select(ProjectPlan).where(
            ProjectPlan.id == plan_id, ProjectPlan.deleted_at.is_(None),
        )
    )
    return r.scalar_one_or_none()


async def get_tasks(
    db: AsyncSession,
    plan_id: uuid.UUID,
    phase: Optional[str] = None,
    status: Optional[str] = None,
    is_critical: Optional[bool] = None,
) -> list[WbsTask]:
    stmt = select(WbsTask).where(
        WbsTask.project_plan_id == plan_id,
        WbsTask.deleted_at.is_(None),
    )
    if phase is not None:
        stmt = stmt.where(WbsTask.phase == str(phase))
    if status:
        stmt = stmt.where(WbsTask.status == status)
    if is_critical is not None:
        stmt = stmt.where(WbsTask.is_critical_path.is_(is_critical))
    r = await db.execute(stmt.order_by(WbsTask.task_code.asc()))
    return list(r.scalars().all())


async def get_task(
    db: AsyncSession, task_id: uuid.UUID,
) -> Optional[WbsTask]:
    r = await db.execute(
        select(WbsTask).where(
            WbsTask.id == task_id, WbsTask.deleted_at.is_(None),
        )
    )
    return r.scalar_one_or_none()


async def list_change_requests(
    db: AsyncSession, project_id: uuid.UUID,
) -> list[ChangeRequest]:
    r = await db.execute(
        select(ChangeRequest)
        .where(
            ChangeRequest.project_id == project_id,
            ChangeRequest.deleted_at.is_(None),
        )
        .order_by(ChangeRequest.created_at.desc())
    )
    return list(r.scalars().all())


async def get_change_request(
    db: AsyncSession, cr_id: uuid.UUID,
) -> Optional[ChangeRequest]:
    r = await db.execute(
        select(ChangeRequest).where(
            ChangeRequest.id == cr_id, ChangeRequest.deleted_at.is_(None),
        )
    )
    return r.scalar_one_or_none()


# =============== Exports ===============

def export_gantt_xlsx(tasks: list[WbsTask], plan: ProjectPlan) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Gantt"
    headers = [
        "WBS", "Tarea", "Fase", "Inicio", "Fin", "Dias",
        "Esfuerzo Marcos (h)", "Esfuerzo Plataforma (h)",
        "Responsable", "Estado", "Progreso %", "Ruta Crítica",
    ]
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    crit_fill = PatternFill("solid", fgColor="FFC7CE")
    done_fill = PatternFill("solid", fgColor="C6EFCE")
    for row_idx, t in enumerate(
        sorted(tasks, key=lambda x: x.task_code), start=2,
    ):
        values = [
            t.task_code, t.task_name, t.phase or "",
            t.start_date.isoformat() if t.start_date else "",
            t.end_date.isoformat() if t.end_date else "",
            t.duration_days or 0,
            t.effort_marcos_hours or 0,
            t.effort_platform_hours or 0,
            t.responsible or "",
            t.status or "",
            t.progress_pct or 0,
            "Sí" if t.is_critical_path else "",
        ]
        for col_idx, v in enumerate(values, start=1):
            c = ws.cell(row=row_idx, column=col_idx, value=v)
            c.alignment = Alignment(vertical="center", wrap_text=True)
        if t.is_critical_path:
            ws.cell(row=row_idx, column=12).fill = crit_fill
        if (t.status or "") == "hecha":
            ws.cell(row=row_idx, column=10).fill = done_fill

    widths = [12, 45, 8, 12, 12, 7, 10, 10, 14, 14, 10, 14]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# =============== Dicts ===============

def plan_to_dict(p: ProjectPlan) -> dict:
    return {
        "id": str(p.id),
        "project_id": str(p.project_id),
        "version": p.version,
        "categoria": p.categoria,
        "start_date": p.start_date.isoformat() if p.start_date else None,
        "end_date_estimated": (
            p.end_date_estimated.isoformat() if p.end_date_estimated else None
        ),
        "end_date_actual": (
            p.end_date_actual.isoformat() if p.end_date_actual else None
        ),
        "total_effort_marcos_hours": p.total_effort_marcos_hours,
        "total_effort_platform_hours": p.total_effort_platform_hours,
        "total_duration_weeks": p.total_duration_weeks,
        "critical_path_length_weeks": p.critical_path_length_weeks,
        "critical_path_tasks": p.critical_path_tasks or [],
        "marcos_weekly_capacity_hours": p.marcos_weekly_capacity_hours,
        "client_weekly_capacity_hours": p.client_weekly_capacity_hours,
        "estado": p.estado,
        "baseline_date": p.baseline_date.isoformat() if p.baseline_date else None,
        "meeting_plan": p.meeting_plan or {},
        "mermaid_gantt": p.mermaid_gantt,
        "created_at": p.created_at.isoformat() if p.created_at else None,
    }


def task_to_dict(t: WbsTask) -> dict:
    return {
        "id": str(t.id),
        "project_plan_id": str(t.project_plan_id),
        "project_id": str(t.project_id) if t.project_id else None,
        "task_code": t.task_code,
        "task_name": t.task_name,
        "phase": t.phase,
        "start_date": t.start_date.isoformat() if t.start_date else None,
        "end_date": t.end_date.isoformat() if t.end_date else None,
        "duration_days": t.duration_days,
        "effort_marcos_hours": t.effort_marcos_hours,
        "effort_platform_hours": t.effort_platform_hours,
        "responsible": t.responsible,
        "dependencies": t.dependencies or [],
        "deliverable_e_code": t.deliverable_e_code,
        "status": t.status,
        "progress_pct": t.progress_pct,
        "is_critical_path": t.is_critical_path,
        "slack_days": t.slack_days,
        "blocker_description": t.blocker_description,
        "baseline_start_date": (
            t.baseline_start_date.isoformat() if t.baseline_start_date else None
        ),
        "baseline_end_date": (
            t.baseline_end_date.isoformat() if t.baseline_end_date else None
        ),
    }


def cr_to_dict(cr: ChangeRequest) -> dict:
    return {
        "id": str(cr.id),
        "project_id": str(cr.project_id),
        "plan_id": str(cr.plan_id),
        "code": cr.code,
        "titulo": cr.titulo,
        "descripcion": cr.descripcion,
        "impacto_plazo_dias": cr.impacto_plazo_dias,
        "impacto_esfuerzo_horas": cr.impacto_esfuerzo_horas,
        "impacto_presupuesto_eur": cr.impacto_presupuesto_eur,
        "estado": cr.estado,
        "solicitado_por": cr.solicitado_por,
        "aprobado_por": cr.aprobado_por,
        "aprobado_at": cr.aprobado_at.isoformat() if cr.aprobado_at else None,
        "created_at": cr.created_at.isoformat() if cr.created_at else None,
    }
