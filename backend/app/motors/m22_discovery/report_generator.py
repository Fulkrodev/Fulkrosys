"""Technical Report Generator (M22-C).

Consolida los datos de los 8 modulos de discovery en un informe tecnico
unificado (JSON + DOCX). Incluye resumen ejecutivo con stats globales,
alertas criticas destacadas y recomendaciones priorizadas deterministas.
"""
from __future__ import annotations

import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from backend.app.models.discovery import (
    DiscoveryAlert,
    VulnerabilityFinding,
)

from . import (
    alerts_service,
    asset_discovery,
    config_discovery,
    continuity_service,
    data_discovery,
    dataflow_service,
    identity_discovery,
    log_assessment,
    vuln_discovery,
)


REPO_ROOT = Path(__file__).resolve().parents[4]
REPORTS_DIR = REPO_ROOT / "var" / "reports"
TEMPLATE_PATH = Path(__file__).parent / "templates" / "informe_tecnico_template.docx"


class ReportError(ValueError):
    pass


def _ensure_reports_dir(project_id: uuid.UUID) -> Path:
    d = REPORTS_DIR / str(project_id) / "m22_technical"
    d.mkdir(parents=True, exist_ok=True)
    return d


async def generate(
    session: AsyncSession, project_id: uuid.UUID,
) -> dict:
    """Consolida datos de los 8 modulos en un dict de informe tecnico."""
    # Summaries
    assets_sum = await asset_discovery.assets_summary(session, project_id)
    ident_sum = await identity_discovery.identity_summary(session, project_id)
    config_sum = await config_discovery.configurations_summary(session, project_id)
    vuln_sum = await vuln_discovery.vulnerabilities_summary(session, project_id)
    data_sum = await data_discovery.data_summary(session, project_id)
    alert_sum = await alerts_service.alerts_summary(session, project_id)

    logging_ass = await log_assessment.get_latest(session, project_id)
    continuity_ass = await continuity_service.get_latest(session, project_id)
    dfds = await dataflow_service.list_dfds(session, project_id)

    # Vulns top + alerts criticas
    r = await session.execute(
        select(VulnerabilityFinding)
        .where(
            VulnerabilityFinding.project_id == project_id,
            VulnerabilityFinding.cvss_severity.in_(["critica", "alta"]),
            VulnerabilityFinding.deleted_at.is_(None),
        )
        .order_by(VulnerabilityFinding.cvss_score.desc().nulls_last())
        .limit(10)
    )
    top_vulns = [
        {
            "cve_id": v.cve_id, "titulo": v.titulo,
            "cvss_score": v.cvss_score, "severity": v.cvss_severity,
            "asset": v.asset_afectado,
        }
        for v in r.scalars().all()
    ]

    r = await session.execute(
        select(DiscoveryAlert)
        .where(
            DiscoveryAlert.project_id == project_id,
            DiscoveryAlert.severidad.in_(["critica", "alta"]),
        )
        .order_by(DiscoveryAlert.created_at.desc())
        .limit(10)
    )
    top_alerts = [
        {
            "severidad": a.severidad, "codigo": a.codigo,
            "titulo": a.titulo, "modulo": a.modulo,
        }
        for a in r.scalars().all()
    ]

    summary_ejecutivo = _build_executive_summary(
        assets_sum, ident_sum, config_sum, vuln_sum, data_sum,
        alert_sum, logging_ass, continuity_ass,
    )

    recomendaciones = _build_recommendations(
        top_alerts, top_vulns, config_sum, ident_sum,
        logging_ass, continuity_ass,
    )

    report = {
        "meta": {
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "project_id": str(project_id),
        },
        "resumen_ejecutivo": summary_ejecutivo,
        "assets": assets_sum,
        "identities": ident_sum,
        "configurations": config_sum,
        "vulnerabilities": {
            **vuln_sum,
            "top_findings": top_vulns,
        },
        "data_stores": data_sum,
        "logging": (
            log_assessment.to_dict(logging_ass) if logging_ass else None
        ),
        "dataflows": [dataflow_service.to_dict(d) for d in dfds],
        "continuity": (
            continuity_service.to_dict(continuity_ass) if continuity_ass else None
        ),
        "alerts": {
            **alert_sum,
            "top": top_alerts,
        },
        "recomendaciones": recomendaciones,
    }
    return report


def _build_executive_summary(
    assets_sum: dict, ident_sum: dict, config_sum: dict, vuln_sum: dict,
    data_sum: dict, alert_sum: dict, logging_ass, continuity_ass,
) -> dict:
    return {
        "total_assets": assets_sum.get("total", 0),
        "total_identities": ident_sum.get("total", 0),
        "mfa_coverage_pct": ident_sum.get("mfa_coverage", {}).get("percentage", 0),
        "total_configurations": config_sum.get("total", 0),
        "configs_fail": config_sum.get("by_estado", {}).get("fail", 0),
        "total_vulns": vuln_sum.get("total", 0),
        "vulns_criticas": vuln_sum.get("by_severity", {}).get("critica", 0),
        "vulns_altas": vuln_sum.get("by_severity", {}).get("alta", 0),
        "total_data_stores": data_sum.get("total", 0),
        "stores_con_datos_personales": data_sum.get("con_datos_personales", 0),
        "total_alerts": alert_sum.get("total", 0),
        "alerts_criticas": alert_sum.get("by_severidad", {}).get("critica", 0),
        "alerts_altas": alert_sum.get("by_severidad", {}).get("alta", 0),
        "nivel_madurez_logging": (
            logging_ass.nivel_madurez_logging if logging_ass else "L0"
        ),
        "nivel_madurez_continuidad": (
            continuity_ass.nivel_madurez_continuidad if continuity_ass else "L0"
        ),
    }


def _build_recommendations(
    top_alerts: list[dict], top_vulns: list[dict], config_sum: dict,
    ident_sum: dict, logging_ass, continuity_ass,
) -> list[dict]:
    recs: list[dict] = []

    # Alertas criticas -> urgente
    for a in top_alerts:
        if a["severidad"] == "critica":
            recs.append({
                "prioridad": "urgente",
                "area": a["modulo"],
                "titulo": a["titulo"],
                "origen": f"alert:{a['codigo']}",
            })

    # Vulns CVSS >= 9 -> urgente
    for v in top_vulns:
        if v.get("cvss_score") and v["cvss_score"] >= 9.0:
            recs.append({
                "prioridad": "urgente",
                "area": "vulnerabilities",
                "titulo": f"Parchear {v.get('cve_id') or v['titulo']} (CVSS {v['cvss_score']})",
                "origen": f"vuln:{v.get('cve_id') or v['titulo']}",
            })

    # Config fail con gap critica -> alta
    fail_count = config_sum.get("by_estado", {}).get("fail", 0)
    gap_crit = config_sum.get("by_gap_severidad", {}).get("critica", 0)
    if gap_crit > 0:
        recs.append({
            "prioridad": "alta",
            "area": "configurations",
            "titulo": (
                f"Corregir {gap_crit} configuraciones con gap critico "
                f"(total fails: {fail_count})"
            ),
            "origen": "config:gap_critica",
        })

    # MFA coverage < 80% -> alta
    mfa_pct = ident_sum.get("mfa_coverage", {}).get("percentage", 0)
    if mfa_pct < 80:
        recs.append({
            "prioridad": "alta",
            "area": "identities",
            "titulo": f"Elevar cobertura MFA del {mfa_pct}% al >=95%",
            "origen": "ident:mfa_coverage",
        })

    # Sin DRP -> alta
    if continuity_ass and not continuity_ass.tiene_drp:
        recs.append({
            "prioridad": "alta",
            "area": "continuity",
            "titulo": "Disenar, documentar y probar DRP (op.cont.1-2)",
            "origen": "cont:no_drp",
        })

    # Logging L0/L1 -> alta
    if logging_ass and logging_ass.nivel_madurez_logging in {"L0", "L1"}:
        recs.append({
            "prioridad": "alta",
            "area": "logs",
            "titulo": (
                f"Implantar/madurar SIEM: nivel actual "
                f"{logging_ass.nivel_madurez_logging} -> objetivo L3 (op.mon/op.exp.8)"
            ),
            "origen": "logs:madurez",
        })

    # Deduplicar por origen
    seen: set[str] = set()
    deduped: list[dict] = []
    prio_order = {"urgente": 0, "alta": 1, "media": 2, "baja": 3}
    for r in recs:
        if r["origen"] in seen:
            continue
        seen.add(r["origen"])
        deduped.append(r)
    deduped.sort(key=lambda x: prio_order.get(x["prioridad"], 9))
    return deduped


def _build_docx_context(report: dict) -> dict:
    resumen = report.get("resumen_ejecutivo", {}) or {}
    return {
        "fecha_generacion": (report.get("meta", {}) or {}).get("generated_at", ""),
        "project_id": (report.get("meta", {}) or {}).get("project_id", ""),
        "resumen": resumen,
        "assets": report.get("assets", {}),
        "identities": report.get("identities", {}),
        "configurations": report.get("configurations", {}),
        "vulnerabilities": report.get("vulnerabilities", {}),
        "data_stores": report.get("data_stores", {}),
        "logging": report.get("logging"),
        "continuity": report.get("continuity"),
        "dataflows": report.get("dataflows", []),
        "alerts": report.get("alerts", {}),
        "recomendaciones": report.get("recomendaciones", []),
    }


def _generate_provisional_template(path: Path) -> None:
    """Crea un template DOCX provisional con python-docx.

    Sesión 3B-2B.4 Phase 4 · header con logo cliente InlineImage + footer
    branding.footer_text (graceful fallback si NULL).
    """
    from docx import Document
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    # Phase 4 · header cliente branding
    doc.add_paragraph("{% if branding.has_logo %}{{ branding.logo_image }}{% else %}{{ branding.client_name|upper }}{% endif %}")

    doc.add_heading("Informe Tecnico de Discovery - Motor 22", level=0)
    doc.add_paragraph("Fecha: {{ fecha_generacion }}")
    doc.add_paragraph("Proyecto: {{ project_id }}")

    doc.add_heading("1. Resumen Ejecutivo", level=1)
    doc.add_paragraph("Assets: {{ resumen.total_assets }}")
    doc.add_paragraph("Identidades: {{ resumen.total_identities }} (MFA {{ resumen.mfa_coverage_pct }}%)")
    doc.add_paragraph("Configuraciones: {{ resumen.total_configurations }} ({{ resumen.configs_fail }} fail)")
    doc.add_paragraph("Vulnerabilidades: {{ resumen.total_vulns }} ({{ resumen.vulns_criticas }} criticas, {{ resumen.vulns_altas }} altas)")
    doc.add_paragraph("Data stores: {{ resumen.total_data_stores }} ({{ resumen.stores_con_datos_personales }} con datos personales)")
    doc.add_paragraph("Alertas: {{ resumen.total_alerts }} ({{ resumen.alerts_criticas }} criticas, {{ resumen.alerts_altas }} altas)")
    doc.add_paragraph("Madurez logging: {{ resumen.nivel_madurez_logging }} / Continuidad: {{ resumen.nivel_madurez_continuidad }}")

    doc.add_heading("2. Inventario de Assets", level=1)
    doc.add_paragraph("Total: {{ assets.total }}")
    doc.add_paragraph("Por tipo MAGERIT:")
    doc.add_paragraph("{% for k, v in (assets.by_tipo_magerit or {}).items() %}- {{ k }}: {{ v }}\n{% endfor %}")

    doc.add_heading("3. Identidades", level=1)
    doc.add_paragraph("Total: {{ identities.total }}, privilegiadas: {{ identities.privilegiadas }}, sin MFA: {{ identities.sin_mfa }}")

    doc.add_heading("4. Configuraciones", level=1)
    doc.add_paragraph("{% for k, v in (configurations.by_estado or {}).items() %}{{ k }}: {{ v }}\n{% endfor %}")

    doc.add_heading("5. Vulnerabilidades", level=1)
    doc.add_paragraph("Total: {{ vulnerabilities.total }}")
    doc.add_paragraph("Top hallazgos:")
    doc.add_paragraph("{% for v in vulnerabilities.top_findings or [] %}* {{ v.cve_id or v.titulo }} (CVSS {{ v.cvss_score }}) - {{ v.severity }}\n{% endfor %}")

    doc.add_heading("6. Almacenes de Datos", level=1)
    doc.add_paragraph("Total: {{ data_stores.total }}, con datos personales: {{ data_stores.con_datos_personales }}")
    doc.add_paragraph("Volumen total: {{ data_stores.volumen_total_gb }} GB")

    doc.add_heading("7. Logging y Monitorizacion", level=1)
    doc.add_paragraph("{% if logging %}Madurez: {{ logging.nivel_madurez_logging }}. SIEM: {{ logging.siem_producto or 'ninguno' }}. op.exp.8 cumple: {{ logging.cumple_op_exp_8 }}.{% else %}Sin assessment registrado.{% endif %}")

    doc.add_heading("8. Flujos de Datos", level=1)
    doc.add_paragraph("{% for d in dataflows %}- {{ d.nombre }} ({{ d.tipo }}) - clasif max: {{ d.clasificacion_max_datos }}\n{% endfor %}")

    doc.add_heading("9. Continuidad", level=1)
    doc.add_paragraph("{% if continuity %}Madurez: {{ continuity.nivel_madurez_continuidad }}. DRP: {{ continuity.tiene_drp }}. SPOFs: {{ continuity.spofs_detectados|length }}.{% else %}Sin assessment registrado.{% endif %}")

    doc.add_heading("10. Alertas", level=1)
    doc.add_paragraph("Total: {{ alerts.total }}")
    doc.add_paragraph("{% for k, v in (alerts.by_severidad or {}).items() %}{{ k }}: {{ v }}\n{% endfor %}")

    doc.add_heading("11. Recomendaciones Priorizadas", level=1)
    doc.add_paragraph("{% for r in recomendaciones %}[{{ r.prioridad|upper }}] {{ r.titulo }} ({{ r.area }})\n{% endfor %}")

    # Phase 4 · footer cliente branding
    doc.add_paragraph("{{ branding.footer_text }}")

    doc.save(str(path))


async def generate_docx(
    session: AsyncSession, project_id: uuid.UUID,
) -> Optional[str]:
    """Genera DOCX del informe tecnico. Devuelve path del fichero o None.

    Sesión 3B-2B.4 Phase 4 · inyecta ClientBranding (logo + colors + footer)
    desde DB · cliente piloto recibe DOCX con SU branding (audit Phase 0 D3).
    """
    try:
        from docxtpl import DocxTemplate, InlineImage
        from docx.shared import Mm
    except ImportError:
        return None

    report = await generate(session, project_id)

    if not TEMPLATE_PATH.exists():
        _generate_provisional_template(TEMPLATE_PATH)

    # Phase 4 · fetch cliente branding ANTES render.
    from backend.app.core.branding import build_branding_pdf_context

    branding = await build_branding_pdf_context(session, project_id)

    doc = DocxTemplate(str(TEMPLATE_PATH))
    context = _build_docx_context(report)
    context["branding"] = branding.template_dict()
    if branding.logo_path and branding.logo_path.exists():
        context["branding"]["logo_image"] = InlineImage(
            doc, str(branding.logo_path), height=Mm(12),
        )
    doc.render(context)

    out_dir = _ensure_reports_dir(project_id)
    ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S")
    out_path = out_dir / f"informe_tecnico_{ts}.docx"
    doc.save(str(out_path))

    # Cleanup temp logo file
    if branding.logo_path and branding.logo_path.exists():
        try:
            branding.logo_path.unlink()
        except OSError:
            pass

    return str(out_path)
