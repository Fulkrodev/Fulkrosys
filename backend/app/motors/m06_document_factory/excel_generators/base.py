"""Helpers comunes para los 16 Excel generators · SAN-C.MB-10.1.

Patrón canónico: cada generator importa estas helpers para mantener
estilo visual consistente (cabeceras + freeze + auto-width + colores
por categoría) y look profesional ENAC-ready.
"""
from __future__ import annotations

from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from backend.app.fulkro_identity import FULKRO_FOOTER_TEXT


# Paleta corporativa FULKRO
HEADER_FILL = PatternFill(
    start_color="1F2937",
    end_color="1F2937",
    fill_type="solid",
)
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
SUBHEADER_FILL = PatternFill(
    start_color="E5E7EB",
    end_color="E5E7EB",
    fill_type="solid",
)
SUBHEADER_FONT = Font(name="Calibri", bold=True, size=10)
BODY_FONT = Font(name="Calibri", size=10)
THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)
WRAP_CENTER = Alignment(wrap_text=True, vertical="center")
WRAP_LEFT = Alignment(wrap_text=True, vertical="center", horizontal="left")


def new_workbook(sheet_title: str) -> Workbook:
    """Workbook listo con la sheet activa renombrada."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]  # Excel limit 31 chars
    # F-14-01 (Ejecutable 8 Pasada 16): identidad Fulkro centralizada en el pie
    # de impresión · cubre las plantillas Excel (no pasan por _inject_brand de
    # DOCX) en UN solo sitio (todos los generators usan new_workbook).
    ws.oddFooter.center.text = FULKRO_FOOTER_TEXT
    return wb


def write_headers(ws: Worksheet, headers: Sequence[str], row: int = 1) -> None:
    """Escribe fila cabecera estilo corporativo."""
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP_CENTER
        cell.border = THIN_BORDER
    ws.row_dimensions[row].height = 28
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def write_data_rows(
    ws: Worksheet,
    rows: Iterable[Sequence[object]],
    start_row: int = 2,
) -> int:
    """Escribe filas data con estilo body. Returns next row index."""
    next_row = start_row
    for row_data in rows:
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=next_row, column=col_idx, value=value)
            cell.font = BODY_FONT
            cell.alignment = WRAP_LEFT
            cell.border = THIN_BORDER
        next_row += 1
    return next_row


def write_metadata_block(
    ws: Worksheet,
    title: str,
    project_name: str,
    today: str,
    template_code: str,
    start_row: int = 1,
) -> int:
    """Bloque metadata estándar al inicio del workbook."""
    ws.cell(row=start_row, column=1, value=title).font = Font(
        name="Calibri", bold=True, size=14
    )
    ws.merge_cells(
        start_row=start_row, start_column=1, end_row=start_row, end_column=6
    )

    meta_row = start_row + 2
    pairs = [
        ("Proyecto:", project_name),
        ("Fecha emisión:", today),
        ("Plantilla:", template_code),
        ("Norma referencia:", "RD 311/2022 + CCN-STIC"),
    ]
    for offset, (label, value) in enumerate(pairs):
        ws.cell(row=meta_row + offset, column=1, value=label).font = Font(
            name="Calibri", bold=True, size=10
        )
        ws.cell(row=meta_row + offset, column=2, value=value).font = BODY_FONT

    return meta_row + len(pairs) + 1


def auto_size_columns(ws: Worksheet, max_width: int = 60) -> None:
    """Auto-width approx (openpyxl no expone medición real · usa max char count)."""
    from openpyxl.utils import get_column_letter

    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 10
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, values_only=True):
            for cell in row:
                if cell is None:
                    continue
                length = len(str(cell))
                if length > max_len:
                    max_len = length
        ws.column_dimensions[letter].width = min(max_len + 2, max_width)
