"""
Motor 2 — MAGERIT v3: XLSX Exporters for individual deliverables.

7 XLSX generators (E-020 to E-026) producing deterministic, standalone
spreadsheets for each MAGERIT analysis deliverable. Used by REST endpoints
for individual document downloads.

All generators:
- Accept SQLAlchemy AsyncSession + analysis_id
- Query the DB directly (no dependency on service layer for reads)
- Return bytes (XLSX content) ready for HTTP Response
- Sort output deterministically (by code/name/dimension) for test reproducibility
"""
import io
import uuid
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from backend.app.motors.m02_magerit.models import (
    MageritAnalysis,
    MageritAsset,
    MageritAssetDependency,
    MageritThreatAssessment,
    MageritSafeguardDeployment,
    MageritRiskCalculation,
    MageritTreatmentPlan,
)


# ================================================================
# SHARED STYLES
# ================================================================

_HEADER_FONT = Font(bold=True, size=11)
_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_HEADER_FONT_WHITE = Font(bold=True, size=11, color="FFFFFF")
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _apply_header_style(ws, headers: list[str]):
    """Write styled header row and set column widths."""
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _HEADER_FONT_WHITE
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = _THIN_BORDER
    # Auto-width estimation
    for col_idx, header in enumerate(headers, start=1):
        ws.column_dimensions[chr(64 + col_idx) if col_idx <= 26 else "A"].width = max(len(header) + 4, 14)


def _set_column_widths(ws, widths: list[int]):
    """Set explicit column widths."""
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _wb_to_bytes(wb: Workbook) -> bytes:
    """Serialize workbook to bytes."""
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ================================================================
# E-020: Inventario de activos
# ================================================================

async def export_asset_inventory(
    db: AsyncSession,
    analysis_id: uuid.UUID,
) -> bytes:
    """E-020: Inventario de activos con valoración DICAT.

    Columns: Código | Nombre | Tipo | Propietario | D | I | C | A | T |
             Acum_D | Acum_I | Acum_C | Acum_A | Acum_T
    """
    assets = (await db.execute(
        select(MageritAsset).where(
            MageritAsset.analysis_id == analysis_id,
            MageritAsset.deleted_at.is_(None),
        ).order_by(MageritAsset.code)
    )).scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "E-020 Inventario Activos"

    headers = [
        "Codigo", "Nombre", "Tipo Activo", "Propietario",
        "Val_D", "Val_I", "Val_C", "Val_A", "Val_T",
        "Acum_D", "Acum_I", "Acum_C", "Acum_A", "Acum_T",
    ]
    _apply_header_style(ws, headers)
    _set_column_widths(ws, [15, 30, 12, 20, 8, 8, 8, 8, 8, 10, 10, 10, 10, 10])

    for row_idx, a in enumerate(assets, start=2):
        ws.cell(row=row_idx, column=1, value=a.code)
        ws.cell(row=row_idx, column=2, value=a.name)
        ws.cell(row=row_idx, column=3, value=a.asset_type_code)
        ws.cell(row=row_idx, column=4, value=a.owner or "")
        ws.cell(row=row_idx, column=5, value=a.value_d or 0)
        ws.cell(row=row_idx, column=6, value=a.value_i or 0)
        ws.cell(row=row_idx, column=7, value=a.value_c or 0)
        ws.cell(row=row_idx, column=8, value=a.value_a or 0)
        ws.cell(row=row_idx, column=9, value=a.value_t or 0)
        ws.cell(row=row_idx, column=10, value=float(a.accumulated_d) if a.accumulated_d else "")
        ws.cell(row=row_idx, column=11, value=float(a.accumulated_i) if a.accumulated_i else "")
        ws.cell(row=row_idx, column=12, value=float(a.accumulated_c) if a.accumulated_c else "")
        ws.cell(row=row_idx, column=13, value=float(a.accumulated_a) if a.accumulated_a else "")
        ws.cell(row=row_idx, column=14, value=float(a.accumulated_t) if a.accumulated_t else "")

    return _wb_to_bytes(wb)


# ================================================================
# E-021: Mapa de dependencias
# ================================================================

async def export_dependency_map(
    db: AsyncSession,
    analysis_id: uuid.UUID,
) -> bytes:
    """E-021: Mapa de dependencias entre activos.

    Columns: Activo Superior (code) | Activo Inferior (code) | Grado | Razon
    """
    deps = (await db.execute(
        select(MageritAssetDependency).where(
            MageritAssetDependency.analysis_id == analysis_id
        )
    )).scalars().all()

    # Build asset_id → code map
    assets = (await db.execute(
        select(MageritAsset).where(
            MageritAsset.analysis_id == analysis_id,
            MageritAsset.deleted_at.is_(None),
        )
    )).scalars().all()
    id_to_code = {str(a.id): a.code for a in assets}

    wb = Workbook()
    ws = wb.active
    ws.title = "E-021 Dependencias"

    headers = ["Activo Superior", "Activo Inferior", "Grado Dependencia", "Razon"]
    _apply_header_style(ws, headers)
    _set_column_widths(ws, [25, 25, 20, 40])

    # Sort deterministically by (superior_code, inferior_code)
    sorted_deps = sorted(
        deps,
        key=lambda d: (id_to_code.get(str(d.superior_asset_id), ""), id_to_code.get(str(d.inferior_asset_id), "")),
    )

    for row_idx, d in enumerate(sorted_deps, start=2):
        ws.cell(row=row_idx, column=1, value=id_to_code.get(str(d.superior_asset_id), str(d.superior_asset_id)))
        ws.cell(row=row_idx, column=2, value=id_to_code.get(str(d.inferior_asset_id), str(d.inferior_asset_id)))
        ws.cell(row=row_idx, column=3, value=d.dependency_degree)
        ws.cell(row=row_idx, column=4, value=d.reason or "")

    return _wb_to_bytes(wb)


# ================================================================
# E-022: Valoración de amenazas
# ================================================================

async def export_threat_assessment(
    db: AsyncSession,
    analysis_id: uuid.UUID,
) -> bytes:
    """E-022: Valoración de amenazas por activo.

    Columns: Activo | Amenaza | Probabilidad | Deg_D | Deg_I | Deg_C | Deg_A | Deg_T
    """
    threats = (await db.execute(
        select(MageritThreatAssessment).where(
            MageritThreatAssessment.analysis_id == analysis_id
        )
    )).scalars().all()

    assets = (await db.execute(
        select(MageritAsset).where(
            MageritAsset.analysis_id == analysis_id,
            MageritAsset.deleted_at.is_(None),
        )
    )).scalars().all()
    id_to_code = {str(a.id): a.code for a in assets}

    wb = Workbook()
    ws = wb.active
    ws.title = "E-022 Amenazas"

    headers = ["Activo", "Amenaza", "Probabilidad", "Deg_D", "Deg_I", "Deg_C", "Deg_A", "Deg_T"]
    _apply_header_style(ws, headers)
    _set_column_widths(ws, [20, 12, 14, 8, 8, 8, 8, 8])

    sorted_threats = sorted(
        threats,
        key=lambda t: (id_to_code.get(str(t.asset_id), ""), t.threat_code),
    )

    for row_idx, t in enumerate(sorted_threats, start=2):
        ws.cell(row=row_idx, column=1, value=id_to_code.get(str(t.asset_id), str(t.asset_id)))
        ws.cell(row=row_idx, column=2, value=t.threat_code)
        ws.cell(row=row_idx, column=3, value=t.probability)
        ws.cell(row=row_idx, column=4, value=t.degradation_d or 0)
        ws.cell(row=row_idx, column=5, value=t.degradation_i or 0)
        ws.cell(row=row_idx, column=6, value=t.degradation_c or 0)
        ws.cell(row=row_idx, column=7, value=t.degradation_a or 0)
        ws.cell(row=row_idx, column=8, value=t.degradation_t or 0)

    return _wb_to_bytes(wb)


# ================================================================
# E-023: Salvaguardas desplegadas
# ================================================================

async def export_safeguard_deployment(
    db: AsyncSession,
    analysis_id: uuid.UUID,
) -> bytes:
    """E-023: Salvaguardas desplegadas con eficacia.

    Columns: Salvaguarda | Estado | Eficacia (%) | Tipo Efecto | Responsable | Notas
    """
    safeguards = (await db.execute(
        select(MageritSafeguardDeployment).where(
            MageritSafeguardDeployment.analysis_id == analysis_id
        )
    )).scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "E-023 Salvaguardas"

    headers = ["Salvaguarda", "Estado", "Eficacia (%)", "Tipo Efecto", "Responsable", "Notas"]
    _apply_header_style(ws, headers)
    _set_column_widths(ws, [20, 14, 14, 16, 20, 40])

    sorted_sg = sorted(safeguards, key=lambda s: s.safeguard_code)

    for row_idx, s in enumerate(sorted_sg, start=2):
        ws.cell(row=row_idx, column=1, value=s.safeguard_code)
        ws.cell(row=row_idx, column=2, value=s.status)
        ws.cell(row=row_idx, column=3, value=s.efficacy)
        ws.cell(row=row_idx, column=4, value=s.effect_type)
        ws.cell(row=row_idx, column=5, value=s.responsible or "")
        ws.cell(row=row_idx, column=6, value=s.notes or "")

    return _wb_to_bytes(wb)


# ================================================================
# E-024: Cálculos de riesgo
# ================================================================

async def export_risk_calculations(
    db: AsyncSession,
    analysis_id: uuid.UUID,
) -> bytes:
    """E-024: Resultados de cálculo de riesgo por (activo, amenaza, dimensión).

    Columns: Activo | Amenaza | Dim | Imp.Intríns | R.Intrins.Acum | R.Intrins.Reperc |
             Imp.Efectivo | R.Efectivo | R.Residual | Nivel
    """
    calcs = (await db.execute(
        select(MageritRiskCalculation).where(
            MageritRiskCalculation.analysis_id == analysis_id
        )
    )).scalars().all()

    assets = (await db.execute(
        select(MageritAsset).where(
            MageritAsset.analysis_id == analysis_id,
            MageritAsset.deleted_at.is_(None),
        )
    )).scalars().all()
    id_to_code = {str(a.id): a.code for a in assets}

    wb = Workbook()
    ws = wb.active
    ws.title = "E-024 Calculos Riesgo"

    headers = [
        "Activo", "Amenaza", "Dim",
        "Imp.Intrinseco", "R.Intrins.Acum", "R.Intrins.Reperc",
        "Imp.Efectivo", "R.Efectivo", "R.Residual", "Nivel",
    ]
    _apply_header_style(ws, headers)
    _set_column_widths(ws, [20, 12, 6, 14, 14, 14, 14, 12, 12, 8])

    sorted_calcs = sorted(
        calcs,
        key=lambda c: (id_to_code.get(str(c.asset_id), ""), c.threat_code, c.dimension),
    )

    for row_idx, c in enumerate(sorted_calcs, start=2):
        ws.cell(row=row_idx, column=1, value=id_to_code.get(str(c.asset_id), str(c.asset_id)))
        ws.cell(row=row_idx, column=2, value=c.threat_code)
        ws.cell(row=row_idx, column=3, value=c.dimension)
        ws.cell(row=row_idx, column=4, value=c.impact_intrinsic)
        ws.cell(row=row_idx, column=5, value=c.risk_intrinsic_accumulated)
        ws.cell(row=row_idx, column=6, value=c.risk_intrinsic_repercuted)
        ws.cell(row=row_idx, column=7, value=c.impact_effective)
        ws.cell(row=row_idx, column=8, value=c.risk_effective)
        ws.cell(row=row_idx, column=9, value=c.risk_residual)
        ws.cell(row=row_idx, column=10, value=c.risk_level or "")

    return _wb_to_bytes(wb)


# ================================================================
# E-025: Plan de tratamiento
# ================================================================

async def export_treatment_plan(
    db: AsyncSession,
    analysis_id: uuid.UUID,
) -> bytes:
    """E-025: Plan de tratamiento de riesgos.

    Columns: Activo | Amenaza | Dim | Riesgo Actual | Tratamiento |
             Descripcion | Salvaguardas Propuestas | Riesgo Objetivo |
             Responsable | Plazo | Estado
    """
    plans = (await db.execute(
        select(MageritTreatmentPlan).where(
            MageritTreatmentPlan.analysis_id == analysis_id,
            MageritTreatmentPlan.deleted_at.is_(None),
        )
    )).scalars().all()

    assets = (await db.execute(
        select(MageritAsset).where(
            MageritAsset.analysis_id == analysis_id,
            MageritAsset.deleted_at.is_(None),
        )
    )).scalars().all()
    id_to_code = {str(a.id): a.code for a in assets}

    wb = Workbook()
    ws = wb.active
    ws.title = "E-025 Plan Tratamiento"

    headers = [
        "Activo", "Amenaza", "Dim", "Riesgo Actual", "Tratamiento",
        "Descripcion", "Salvaguardas Propuestas", "Riesgo Objetivo",
        "Responsable", "Plazo", "Estado",
    ]
    _apply_header_style(ws, headers)
    _set_column_widths(ws, [20, 12, 6, 14, 14, 30, 25, 14, 20, 14, 12])

    sorted_plans = sorted(
        plans,
        key=lambda p: (id_to_code.get(str(p.asset_id), ""), p.threat_code, p.dimension),
    )

    for row_idx, p in enumerate(sorted_plans, start=2):
        ws.cell(row=row_idx, column=1, value=id_to_code.get(str(p.asset_id), str(p.asset_id)))
        ws.cell(row=row_idx, column=2, value=p.threat_code)
        ws.cell(row=row_idx, column=3, value=p.dimension)
        ws.cell(row=row_idx, column=4, value=p.current_risk_level)
        ws.cell(row=row_idx, column=5, value=p.treatment)
        ws.cell(row=row_idx, column=6, value=p.action_description or "")
        ws.cell(row=row_idx, column=7, value=", ".join(p.proposed_safeguards) if p.proposed_safeguards else "")
        ws.cell(row=row_idx, column=8, value=p.target_risk_level or "")
        ws.cell(row=row_idx, column=9, value=p.responsible or "")
        ws.cell(row=row_idx, column=10, value=p.deadline.isoformat() if p.deadline else "")
        ws.cell(row=row_idx, column=11, value=p.status)

    return _wb_to_bytes(wb)


# ================================================================
# E-026: Resumen ejecutivo
# ================================================================

async def export_executive_summary(
    db: AsyncSession,
    analysis_id: uuid.UUID,
) -> bytes:
    """E-026: Resumen ejecutivo del análisis MAGERIT.

    Multi-sheet workbook:
    - Hoja 1: Datos generales del análisis
    - Hoja 2: Distribución de riesgo por nivel
    - Hoja 3: Top 10 riesgos más altos
    """
    from backend.app.motors.m02_magerit.service import LEVEL_TO_INDEX

    analysis = await db.get(MageritAnalysis, analysis_id)

    assets = (await db.execute(
        select(MageritAsset).where(
            MageritAsset.analysis_id == analysis_id,
            MageritAsset.deleted_at.is_(None),
        )
    )).scalars().all()

    calcs = (await db.execute(
        select(MageritRiskCalculation).where(
            MageritRiskCalculation.analysis_id == analysis_id
        )
    )).scalars().all()

    threats = (await db.execute(
        select(MageritThreatAssessment).where(
            MageritThreatAssessment.analysis_id == analysis_id
        )
    )).scalars().all()

    safeguards = (await db.execute(
        select(MageritSafeguardDeployment).where(
            MageritSafeguardDeployment.analysis_id == analysis_id
        )
    )).scalars().all()

    plans = (await db.execute(
        select(MageritTreatmentPlan).where(
            MageritTreatmentPlan.analysis_id == analysis_id,
            MageritTreatmentPlan.deleted_at.is_(None),
        )
    )).scalars().all()

    id_to_code = {str(a.id): a.code for a in assets}

    wb = Workbook()

    # --- Sheet 1: Datos generales ---
    ws1 = wb.active
    ws1.title = "Resumen General"

    ws1.cell(row=1, column=1, value="Campo").font = _HEADER_FONT
    ws1.cell(row=1, column=2, value="Valor").font = _HEADER_FONT
    _set_column_widths(ws1, [30, 50])

    risk_levels = [c.risk_level for c in calcs if c.risk_level]
    max_risk = max(risk_levels, key=lambda x: LEVEL_TO_INDEX.get(x, 0)) if risk_levels else "N/A"

    summary_data = [
        ("Nombre del analisis", analysis.name),
        ("Modo de calculo", analysis.calculation_mode),
        ("Version metodologia", analysis.methodology_version),
        ("Estado", analysis.status),
        ("Fecha exportacion", date.today().isoformat()),
        ("Total activos", len(assets)),
        ("Total amenazas valoradas", len(threats)),
        ("Total salvaguardas", len(safeguards)),
        ("Total calculos de riesgo", len(calcs)),
        ("Riesgo maximo", max_risk),
        ("Acciones de tratamiento", len(plans)),
        ("Congelado", "Si" if analysis.snapshot_frozen_at else "No"),
    ]
    for row_idx, (campo, valor) in enumerate(summary_data, start=2):
        ws1.cell(row=row_idx, column=1, value=campo)
        ws1.cell(row=row_idx, column=2, value=str(valor))

    # --- Sheet 2: Distribución de riesgo ---
    ws2 = wb.create_sheet("Distribucion Riesgo")
    headers2 = ["Nivel", "Cantidad", "Porcentaje"]
    _apply_header_style(ws2, headers2)
    _set_column_widths(ws2, [12, 12, 14])

    total = len(risk_levels) if risk_levels else 1
    for row_idx, level in enumerate(["MB", "B", "M", "A", "MA"], start=2):
        count = sum(1 for r in risk_levels if r == level)
        ws2.cell(row=row_idx, column=1, value=level)
        ws2.cell(row=row_idx, column=2, value=count)
        ws2.cell(row=row_idx, column=3, value=f"{count/total*100:.1f}%")

    # --- Sheet 3: Top 10 riesgos ---
    ws3 = wb.create_sheet("Top 10 Riesgos")
    headers3 = ["Activo", "Amenaza", "Dimension", "R.Efectivo", "Nivel"]
    _apply_header_style(ws3, headers3)
    _set_column_widths(ws3, [20, 12, 10, 14, 8])

    # Sort by risk level descending, then effective risk descending
    sorted_calcs = sorted(
        calcs,
        key=lambda c: (
            -LEVEL_TO_INDEX.get(c.risk_level or "MB", 0),
            -(c.risk_effective or 0),
            id_to_code.get(str(c.asset_id), ""),
            c.threat_code,
            c.dimension,
        ),
    )

    for row_idx, c in enumerate(sorted_calcs[:10], start=2):
        ws3.cell(row=row_idx, column=1, value=id_to_code.get(str(c.asset_id), str(c.asset_id)))
        ws3.cell(row=row_idx, column=2, value=c.threat_code)
        ws3.cell(row=row_idx, column=3, value=c.dimension)
        ws3.cell(row=row_idx, column=4, value=c.risk_effective)
        ws3.cell(row=row_idx, column=5, value=c.risk_level or "")

    return _wb_to_bytes(wb)
