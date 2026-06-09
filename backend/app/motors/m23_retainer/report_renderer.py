"""Retainer report DOCX renderer · MB-7.bis atom 7.bis.4.

Builds quarterly and annual DOCX reports from `retainer_quarterly_reports`
data using python-docx (programmatic builder · no external .docx template
required).

Q5.B+D cement: 2 templates retainer-specific (quarterly + annual) + UI
drill-down endpoint. Reuses M06 docxtpl convention but here we build
from scratch since the report is rich tabular + can vary per period.
"""
from __future__ import annotations

import io
import uuid
from datetime import date
from typing import Optional

from docx import Document
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession



def _add_kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    """Helper · build a 2-column key-value table from list of pairs."""
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Light Grid Accent 1"
    for i, (k, v) in enumerate(rows):
        cells = table.rows[i].cells
        cells[0].text = k
        cells[1].text = v


def _fmt_rag(rag: Optional[str]) -> str:
    if not rag:
        return "—"
    return {"green": "🟢 Verde", "amber": "🟡 Ámbar", "red": "🔴 Rojo"}.get(
        rag, rag,
    )


async def _fetch_quarterly_report_row(
    db: AsyncSession,
    retainer_id: uuid.UUID,
    period_start: str,
) -> Optional[dict]:
    # asyncpg requires a real `date` object for DATE columns.
    ps_date = date.fromisoformat(period_start)
    row = (await db.execute(
        text(
            "SELECT id, period_start, period_end, activities_completed, "
            "activities_pending, activities_overdue, incidents_detected, "
            "normativa_changes_relevant, vulns_critical, rag_overall, "
            "summary_jsonb "
            "FROM retainer_quarterly_reports "
            "WHERE retainer_contract_id = :rid "
            "AND period_start = :ps "
            "AND period_type = 'quarterly' "
            "ORDER BY created_at DESC LIMIT 1"
        ),
        {"rid": str(retainer_id), "ps": ps_date},
    )).mappings().first()
    return dict(row) if row else None


async def _fetch_annual_aggregate(
    db: AsyncSession,
    retainer_id: uuid.UUID,
    year: int,
) -> dict:
    """Aggregate 4 quarterly rows of a year to produce annual summary."""
    rows = (await db.execute(
        text(
            "SELECT period_start, period_end, activities_completed, "
            "activities_pending, activities_overdue, incidents_detected, "
            "normativa_changes_relevant, vulns_critical, rag_overall "
            "FROM retainer_quarterly_reports "
            "WHERE retainer_contract_id = :rid "
            "AND period_type = 'quarterly' "
            "AND EXTRACT(YEAR FROM period_start) = :y "
            "ORDER BY period_start"
        ),
        {"rid": str(retainer_id), "y": year},
    )).mappings().all()
    quarters = [dict(r) for r in rows]
    return {
        "year": year,
        "quarters": quarters,
        "activities_completed": sum(q["activities_completed"] or 0 for q in quarters),
        "activities_pending": sum(q["activities_pending"] or 0 for q in quarters),
        "activities_overdue": sum(q["activities_overdue"] or 0 for q in quarters),
        "incidents_detected": sum(q["incidents_detected"] or 0 for q in quarters),
        "normativa_changes_relevant": sum(
            q["normativa_changes_relevant"] or 0 for q in quarters
        ),
        "vulns_critical": sum(q["vulns_critical"] or 0 for q in quarters),
        # Worst RAG across quarters (red > amber > green)
        "rag_overall": _worst_rag([q.get("rag_overall") for q in quarters]),
    }


def _worst_rag(rags: list[Optional[str]]) -> Optional[str]:
    order = {"red": 3, "amber": 2, "green": 1}
    candidates = [r for r in rags if r]
    if not candidates:
        return None
    return max(candidates, key=lambda r: order.get(r, 0))


async def _fetch_retainer_with_client(
    db: AsyncSession, retainer_id: uuid.UUID,
) -> Optional[dict]:
    row = (await db.execute(
        text(
            "SELECT rc.id, rc.perfil, rc.precio_mensual, rc.estado, "
            "rc.inicio, rc.fin, c.nombre AS client_name, c.cif "
            "FROM retainer_contracts rc "
            "JOIN clients c ON c.id = rc.client_id "
            "WHERE rc.id = :rid"
        ),
        {"rid": str(retainer_id)},
    )).mappings().first()
    return dict(row) if row else None


async def render_quarterly_report_docx(
    db: AsyncSession,
    retainer_id: uuid.UUID,
    period_start: str,  # ISO date YYYY-MM-DD
) -> bytes:
    """Build quarterly DOCX bytes for a retainer + period."""
    retainer = await _fetch_retainer_with_client(db, retainer_id)
    if not retainer:
        raise ValueError(f"Retainer {retainer_id} not found")

    report = await _fetch_quarterly_report_row(db, retainer_id, period_start)
    doc = Document()

    # Header
    h = doc.add_heading("Informe Trimestral Retainer", level=0)
    p = doc.add_paragraph()
    run = p.add_run(f"{retainer['client_name']} · {period_start}")
    run.bold = True

    # Cliente / contrato
    doc.add_heading("Datos del contrato", level=1)
    _add_kv_table(doc, [
        ("Cliente", retainer["client_name"]),
        ("CIF", retainer.get("cif") or "—"),
        ("Perfil", retainer["perfil"]),
        ("Precio mensual", f"{retainer['precio_mensual']} €"),
        ("Estado", retainer["estado"]),
        ("Inicio", str(retainer["inicio"])),
        ("Fin", str(retainer.get("fin") or "—")),
    ])

    if report:
        doc.add_heading("Resumen del trimestre", level=1)
        _add_kv_table(doc, [
            ("Periodo", f"{report['period_start']} → {report['period_end']}"),
            ("Actividades completadas", str(report["activities_completed"])),
            ("Actividades pendientes", str(report["activities_pending"])),
            ("Actividades vencidas", str(report["activities_overdue"])),
            ("Incidentes detectados", str(report["incidents_detected"])),
            ("Cambios normativos relevantes", str(report["normativa_changes_relevant"])),
            ("Vulnerabilidades críticas", str(report["vulns_critical"])),
            ("RAG global", _fmt_rag(report["rag_overall"])),
        ])
        if report.get("summary_jsonb"):
            doc.add_heading("Resumen ejecutivo", level=1)
            for key, value in (report["summary_jsonb"] or {}).items():
                doc.add_paragraph(
                    f"{key.replace('_', ' ').title()}: {value}",
                    style="List Bullet",
                )
    else:
        doc.add_paragraph(
            "(Sin datos cuantitativos disponibles para este periodo · "
            "informe parcial.)",
        )

    doc.add_heading("Conclusiones", level=1)
    doc.add_paragraph(
        "Este informe consolida la actividad del trimestre según los "
        "indicadores agregados del retainer (M23). Para detalle por "
        "obligación, evidencia o incidente consultar el portal cliente.",
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def render_annual_report_docx(
    db: AsyncSession,
    retainer_id: uuid.UUID,
    year: int,
) -> bytes:
    """Build annual DOCX bytes aggregating 4 quarters of `year`."""
    retainer = await _fetch_retainer_with_client(db, retainer_id)
    if not retainer:
        raise ValueError(f"Retainer {retainer_id} not found")

    agg = await _fetch_annual_aggregate(db, retainer_id, year)
    doc = Document()

    doc.add_heading(f"Informe Anual Retainer · {year}", level=0)
    p = doc.add_paragraph()
    run = p.add_run(f"{retainer['client_name']}")
    run.bold = True

    doc.add_heading("Datos del contrato", level=1)
    _add_kv_table(doc, [
        ("Cliente", retainer["client_name"]),
        ("CIF", retainer.get("cif") or "—"),
        ("Perfil", retainer["perfil"]),
        ("Precio mensual", f"{retainer['precio_mensual']} €"),
        ("Estado", retainer["estado"]),
    ])

    doc.add_heading("Agregado anual", level=1)
    _add_kv_table(doc, [
        ("Año", str(year)),
        ("Trimestres consolidados", str(len(agg["quarters"]))),
        ("Actividades completadas (total)", str(agg["activities_completed"])),
        ("Actividades pendientes (total)", str(agg["activities_pending"])),
        ("Actividades vencidas (total)", str(agg["activities_overdue"])),
        ("Incidentes detectados (total)", str(agg["incidents_detected"])),
        ("Cambios normativos relevantes (total)", str(agg["normativa_changes_relevant"])),
        ("Vulnerabilidades críticas (total)", str(agg["vulns_critical"])),
        ("RAG anual (peor trimestre)", _fmt_rag(agg["rag_overall"])),
    ])

    if agg["quarters"]:
        doc.add_heading("Desglose por trimestre", level=1)
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Periodo"
        hdr[1].text = "Completadas"
        hdr[2].text = "Vencidas"
        hdr[3].text = "Incidentes"
        hdr[4].text = "RAG"
        for q in agg["quarters"]:
            row = table.add_row().cells
            row[0].text = f"{q['period_start']} → {q['period_end']}"
            row[1].text = str(q["activities_completed"] or 0)
            row[2].text = str(q["activities_overdue"] or 0)
            row[3].text = str(q["incidents_detected"] or 0)
            row[4].text = _fmt_rag(q["rag_overall"])
    else:
        doc.add_paragraph(
            "(Sin trimestres consolidados para este año.)",
        )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
