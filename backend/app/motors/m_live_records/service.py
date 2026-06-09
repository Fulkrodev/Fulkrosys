"""Service layer m_live_records · sub-lote 1.C.B fase 3a.

CRUD generico sobre tabla ``live_records`` con discriminator register_type.
Validacion entry_data via Pydantic schemas (``validate_entry_data``). RLS
enforced project-scoped: caller debe SET app.current_project_id antes de
invocar metodos del service (helper ``backend.app.database.set_tenant_context``).

Operaciones expuestas:
- list_records (paginated · filtros register_type + status)
- get_record
- create_record (validacion entry_data)
- update_record (partial)
- archive_record (soft archive · status='archived')
- compute_dashboard (counts active/archived per register_type)
- export_csv / export_xlsx
"""
from __future__ import annotations

import csv
import uuid
from io import BytesIO, StringIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import and_, func, select, update
from sqlalchemy.ext.asyncio import AsyncSession

from backend.app.models.live_record import LiveRecord
from backend.app.motors.m_live_records.schemas import (
    ENTRY_SCHEMAS,
    REGISTER_TYPE_BLOQUES,
    REGISTER_TYPE_LABELS,
    LiveRecordsDashboardBlock,
    validate_entry_data,
)


class LiveRecordsService:
    """Service registros vivos E-300..E-325 (project-scoped RLS)."""

    def __init__(self, db: AsyncSession) -> None:
        self.db = db

    async def list_records(
        self,
        *,
        project_id: uuid.UUID,
        register_type: str | None = None,
        status: str = "active",
        limit: int = 50,
        offset: int = 0,
    ) -> tuple[list[LiveRecord], int]:
        """List records with pagination. Returns (rows, total_count)."""
        if register_type is not None and register_type not in ENTRY_SCHEMAS:
            raise ValueError(f"Unknown register_type {register_type!r}")

        conditions = [LiveRecord.project_id == project_id]
        if register_type is not None:
            conditions.append(LiveRecord.register_type == register_type)
        if status is not None:
            conditions.append(LiveRecord.status == status)

        where_clause = and_(*conditions)

        total_q = select(func.count()).select_from(LiveRecord).where(where_clause)
        total = (await self.db.execute(total_q)).scalar_one()

        rows_q = (
            select(LiveRecord)
            .where(where_clause)
            .order_by(LiveRecord.created_at.desc())
            .limit(limit)
            .offset(offset)
        )
        rows = (await self.db.execute(rows_q)).scalars().all()
        return list(rows), int(total)

    async def get_record(
        self, *, project_id: uuid.UUID, record_id: uuid.UUID
    ) -> LiveRecord | None:
        q = select(LiveRecord).where(
            and_(
                LiveRecord.id == record_id,
                LiveRecord.project_id == project_id,
            )
        )
        return (await self.db.execute(q)).scalar_one_or_none()

    async def create_record(
        self,
        *,
        project_id: uuid.UUID,
        register_type: str,
        entry_data: dict[str, Any],
        created_by: uuid.UUID,
    ) -> LiveRecord:
        """Create a live record. Validates entry_data against schema."""
        validated = validate_entry_data(register_type, entry_data)
        record = LiveRecord(
            project_id=project_id,
            register_type=register_type,
            entry_data=validated,
            status="active",
            created_by=created_by,
            updated_by=created_by,
        )
        self.db.add(record)
        await self.db.flush()
        await self.db.refresh(record)
        return record

    async def update_record(
        self,
        *,
        project_id: uuid.UUID,
        record_id: uuid.UUID,
        updated_by: uuid.UUID,
        entry_data: dict[str, Any] | None = None,
        status: str | None = None,
    ) -> LiveRecord | None:
        """Update entry_data and/or status. Validates entry_data if provided."""
        record = await self.get_record(project_id=project_id, record_id=record_id)
        if record is None:
            return None

        if entry_data is not None:
            record.entry_data = validate_entry_data(record.register_type, entry_data)
        if status is not None:
            if status not in ("active", "archived"):
                raise ValueError(f"Invalid status {status!r}")
            record.status = status

        record.updated_by = updated_by
        await self.db.flush()
        await self.db.refresh(record)
        return record

    async def archive_record(
        self,
        *,
        project_id: uuid.UUID,
        record_id: uuid.UUID,
        updated_by: uuid.UUID,
    ) -> LiveRecord | None:
        """Soft archive (status='archived'). Returns updated record or None."""
        return await self.update_record(
            project_id=project_id,
            record_id=record_id,
            updated_by=updated_by,
            status="archived",
        )

    async def compute_dashboard(
        self, *, project_id: uuid.UUID
    ) -> tuple[list[LiveRecordsDashboardBlock], int]:
        """Compute count per (register_type, status). Returns (blocks, total_active)."""
        q = (
            select(
                LiveRecord.register_type,
                LiveRecord.status,
                func.count().label("cnt"),
            )
            .where(LiveRecord.project_id == project_id)
            .group_by(LiveRecord.register_type, LiveRecord.status)
        )
        rows = (await self.db.execute(q)).all()

        counts: dict[str, dict[str, int]] = {
            rt: {"active": 0, "archived": 0} for rt in ENTRY_SCHEMAS.keys()
        }
        for register_type, status, cnt in rows:
            if register_type in counts and status in counts[register_type]:
                counts[register_type][status] = int(cnt)

        blocks: list[LiveRecordsDashboardBlock] = []
        total_active = 0
        for register_type in sorted(ENTRY_SCHEMAS.keys()):
            active = counts[register_type]["active"]
            archived = counts[register_type]["archived"]
            total_active += active
            blocks.append(
                LiveRecordsDashboardBlock(
                    register_type=register_type,
                    label=REGISTER_TYPE_LABELS[register_type],
                    bloque=REGISTER_TYPE_BLOQUES[register_type],
                    active_count=active,
                    archived_count=archived,
                )
            )
        return blocks, total_active

    async def export_csv(
        self, *, project_id: uuid.UUID, register_type: str
    ) -> bytes:
        """Export all records of register_type as CSV (utf-8 with BOM)."""
        if register_type not in ENTRY_SCHEMAS:
            raise ValueError(f"Unknown register_type {register_type!r}")

        rows, _ = await self.list_records(
            project_id=project_id,
            register_type=register_type,
            status=None,  # include archived too
            limit=10_000,
            offset=0,
        )

        schema = ENTRY_SCHEMAS[register_type]
        field_names = list(schema.model_fields.keys())

        buf = StringIO()
        buf.write("﻿")  # BOM for Excel utf-8 detection
        writer = csv.writer(buf, dialect="excel")
        writer.writerow(["id", "status", "created_at", "updated_at", *field_names])
        for r in rows:
            entry: dict[str, Any] = r.entry_data or {}
            row_values: list[Any] = [
                str(r.id),
                r.status,
                r.created_at.isoformat(),
                r.updated_at.isoformat(),
            ]
            for fn in field_names:
                v = entry.get(fn)
                if isinstance(v, (list, dict)):
                    row_values.append(repr(v))
                else:
                    row_values.append(v if v is not None else "")
            writer.writerow(row_values)

        return buf.getvalue().encode("utf-8")

    async def export_xlsx(
        self, *, project_id: uuid.UUID, register_type: str
    ) -> bytes:
        """Export all records of register_type as XLSX (openpyxl)."""
        if register_type not in ENTRY_SCHEMAS:
            raise ValueError(f"Unknown register_type {register_type!r}")

        rows, _ = await self.list_records(
            project_id=project_id,
            register_type=register_type,
            status=None,
            limit=10_000,
            offset=0,
        )

        schema = ENTRY_SCHEMAS[register_type]
        field_names = list(schema.model_fields.keys())
        label = REGISTER_TYPE_LABELS[register_type]

        wb = Workbook()
        ws = wb.active
        ws.title = register_type

        ws.cell(row=1, column=1, value=f"{register_type} · {label}").font = Font(
            bold=True, size=14
        )
        ws.cell(row=2, column=1, value=f"project_id: {project_id}").font = Font(
            italic=True, size=10
        )

        headers = ["id", "status", "created_at", "updated_at", *field_names]
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=col_idx, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                start_color="2C3E50", end_color="2C3E50", fill_type="solid"
            )
            cell.alignment = Alignment(horizontal="center")

        for r_idx, r in enumerate(rows, start=5):
            entry: dict[str, Any] = r.entry_data or {}
            ws.cell(row=r_idx, column=1, value=str(r.id))
            ws.cell(row=r_idx, column=2, value=r.status)
            ws.cell(row=r_idx, column=3, value=r.created_at.isoformat())
            ws.cell(row=r_idx, column=4, value=r.updated_at.isoformat())
            for offset, fn in enumerate(field_names, start=5):
                v = entry.get(fn)
                if isinstance(v, (list, dict)):
                    ws.cell(row=r_idx, column=offset, value=repr(v))
                else:
                    ws.cell(row=r_idx, column=offset, value=v)

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()
