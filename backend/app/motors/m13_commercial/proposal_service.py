"""Genera y gestiona propuestas comerciales (M13).

Estados: draft → sent → under_review → negotiating → won / lost.
"""
from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timedelta, timezone
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from backend.app.models.commercial import Proposal
from backend.app.motors.m17_planning.effort_estimator import (
    estimate_duration_weeks,
    estimate_effort,
)

from .pricing_service import PricingService


VALID_ESTADOS = {
    "draft", "sent", "under_review", "negotiating", "won", "lost",
    "superseded",  # MB-19.3 · revisión obsoleta · ADR-041
}


class ProposalError(Exception):
    pass


class ProposalService:
    """Servicio de propuestas comerciales."""

    def __init__(self):
        self.pricing = PricingService()

    async def generate_proposal(
        self,
        db: AsyncSession,
        lead_id: uuid.UUID,
        pricing_model_id: str,
        categoria: str,
        project_id: uuid.UUID | None = None,
        empleados: int = 0,
        sistemas: int = 0,
        ubicaciones: int = 1,
        sector_regulado: bool = False,
        cpds: int = 0,
        client_size: str = "mediana",
        complexity: str = "media",
        importe_override: float | None = None,
        duracion_semanas_override: int | None = None,
        notas_marcos: str | None = None,
        alcance: dict | None = None,
        validez_dias: int = 30,
    ) -> Proposal:
        price = self.pricing.calculate_price(
            model_id=pricing_model_id,
            empleados=empleados,
            sistemas=sistemas,
            ubicaciones=ubicaciones,
            sector_regulado=sector_regulado,
            cpds=cpds,
        )

        if importe_override is not None:
            importe_total = float(importe_override)
            hitos = [
                {
                    "nombre": h["nombre"],
                    "pct": h["pct"],
                    "importe": round(importe_total * h["pct"] / 100, 2),
                }
                for h in price["hitos"]
            ]
        else:
            importe_total = price["total"]
            hitos = price["hitos"]

        duracion = duracion_semanas_override or estimate_duration_weeks(categoria)
        # Esfuerzo aproximado: base 150h para MEDIA, escalado por categoría + tamaño + complejidad
        effort_h = estimate_effort(150, categoria, client_size=client_size, complexity=complexity)

        proposal = Proposal(
            lead_id=lead_id,
            project_id=project_id,
            pricing_model_id=pricing_model_id,
            version=1,
            categoria_objetivo=(categoria or "").upper(),
            alcance=alcance or {
                "empleados": empleados,
                "sistemas": sistemas,
                "ubicaciones": ubicaciones,
                "sector_regulado": sector_regulado,
                "cpds": cpds,
                "client_size": client_size,
                "complexity": complexity,
            },
            duracion_semanas=duracion,
            effort_marcos_horas=effort_h,
            importe_total=importe_total,
            importe_desglose={
                "base": price["base"],
                "extras": price["extras"],
                "iva_percent": price["iva_percent"],
                "iva_importe": price["iva_importe"],
                "total_con_iva": round(importe_total * (1 + price["iva_percent"] / 100), 2),
            },
            hitos_pago={"hitos": hitos},
            validez_hasta=date.today() + timedelta(days=validez_dias),
            estado="draft",
            notas_marcos=notas_marcos,
        )
        db.add(proposal)
        await db.flush()
        return proposal

    # #5 cabo N2 · modelo de pricing legacy por categoría (proyecto · NO
    # retainer). SUPERSEDED por BUG5: ``regenerate_for_categoria`` ya NO usa este
    # mapping (cotizaba MEDIA 22.000 vía PRICING_CATALOG['media_hitos']); ahora
    # delega en ``generate_proposal_apendice_m`` (BASE_PRICES_CANONICAL · MEDIA 10.700).
    # Conservado solo como referencia histórica.
    _PROJECT_PRICING_MODEL = {
        "BASICA": "basica_fijo",
        "MEDIA": "media_hitos",
        "ALTA": "alta_fases_exito",
    }

    async def regenerate_for_categoria(
        self,
        db: AsyncSession,
        project_id: uuid.UUID,
        new_categoria: str,
    ) -> Proposal | None:
        """#5 cabo N2 · regenera la propuesta vigente a una NUEVA categoría
        (category-aware · corrige el ``create_version`` que copiaba la categoría
        vieja). Marca ``superseded`` la anterior y genera una versión nueva con
        la categoría nueva + pricing recomputado con el modelo de la nueva
        categoría. Devuelve ``None`` si no hay propuesta que regenerar.
        """
        cat = (new_categoria or "").strip().upper()
        res = await db.execute(
            select(Proposal)
            .where(
                Proposal.project_id == project_id,
                Proposal.superseded.is_(False),
                Proposal.deleted_at.is_(None),
            )
            .order_by(Proposal.version.desc())
        )
        actual = res.scalars().first()
        if actual is None:
            return None

        alc = actual.alcance if isinstance(actual.alcance, dict) else {}

        actual.superseded = True
        await db.flush()

        # BUG5 fix · la elevación de categoría debe cotizar por la MISMA vía
        # canónica que la propuesta inicial (``generate_proposal_apendice_m`` →
        # ``PricingCalculator`` → BASE_PRICES_CANONICAL: MEDIA 10.700), NO por el legacy
        # ``PRICING_CATALOG['media_hitos']`` (22.000). El alcance viejo guardaba
        # ``sistemas``/``ubicaciones`` (vía legacy); ``sistemas_en_alcance``/
        # ``sedes`` lo respeta si ya viene del Apendice M.
        sistemas_en_alcance = int(
            alc.get("sistemas_en_alcance", alc.get("sistemas", 1)) or 1
        )
        sedes = int(alc.get("sedes", alc.get("ubicaciones", 1)) or 1)
        madurez_pct = alc.get("madurez_pct")
        nueva = await self.generate_proposal_apendice_m(
            db,
            lead_id=actual.lead_id,
            categoria=cat,
            sector=alc.get("sector"),
            sistemas_en_alcance=sistemas_en_alcance,
            sedes=sedes,
            madurez_pct=madurez_pct,
            project_id=project_id,
            notas_marcos=(
                f"Regenerada por elevación de categoría a {cat} (suelo AAPP · #5)."
            ),
            narrative_mode="deterministic",
        )
        nueva.version = (actual.version or 1) + 1
        await db.flush()
        return nueva

    async def get_proposal(self, db: AsyncSession, proposal_id: uuid.UUID) -> Proposal | None:
        res = await db.execute(select(Proposal).where(Proposal.id == proposal_id))
        return res.scalar_one_or_none()

    async def list_proposals(
        self,
        db: AsyncSession,
        project_id: uuid.UUID | None = None,
        lead_id: uuid.UUID | None = None,
    ) -> list[Proposal]:
        stmt = select(Proposal)
        if project_id:
            stmt = stmt.where(Proposal.project_id == project_id)
        if lead_id:
            stmt = stmt.where(Proposal.lead_id == lead_id)
        stmt = stmt.order_by(Proposal.created_at.desc())
        res = await db.execute(stmt)
        return list(res.scalars().all())

    async def update_status(
        self, db: AsyncSession, proposal_id: uuid.UUID, estado: str
    ) -> Proposal:
        if estado not in VALID_ESTADOS:
            raise ProposalError(
                f"Estado inválido: {estado}. Válidos: {sorted(VALID_ESTADOS)}"
            )
        p = await self.get_proposal(db, proposal_id)
        if not p:
            raise ProposalError(f"Proposal {proposal_id} no encontrada")
        p.estado = estado
        await db.flush()
        return p

    # ══════════════════════════════════════════════════════════════════
    # Paso 6 — propuestas con pricing oficial Apendice M v2.2
    # ══════════════════════════════════════════════════════════════════

    async def generate_proposal_apendice_m(
        self,
        db: AsyncSession,
        *,
        lead_id: uuid.UUID,
        categoria: str,
        cliente: Any = None,
        sector: str | None = None,
        sistemas_en_alcance: int = 1,
        sedes: int = 1,
        madurez_pct: int | None = None,
        dias_hasta_plazo: int | None = None,
        project_id: uuid.UUID | None = None,
        notas_marcos: str | None = None,
        validez_dias: int = 30,
        narrative_mode: str = "deterministic",
    ) -> Proposal:
        """Genera una Proposal alineada con el Apendice M v2.2 oficial.

        Usa ``backend.app.core.pricing.PricingCalculator`` para aplicar:
        - Precios base según BASE_PRICES_CANONICAL · fuente única pricing_config
          (BASICA 3.200 / MEDIA 10.700 / ALTA 22.800).
        - Extras MEDIA (sector_regulado, multi_ubicacion, madurez_l0_l1,
          sistemas_adicionales).
        - Hitos oficiales 30/40/30, 26/21/21/21/11 o 7 hitos ALTA.
        - Recargo urgencia +30% si plazo <6 semanas.
        - Garantias por categoria.
        - Plazo pago 60d AAPP / 30d privado.
        """
        from backend.app.core.pricing import PricingCalculator
        calc = PricingCalculator()
        result = calc.calculate_implantacion(
            categoria,
            cliente=cliente,
            sector=sector,
            sistemas_en_alcance=sistemas_en_alcance,
            sedes=sedes,
            madurez_pct=madurez_pct,
            dias_hasta_plazo=dias_hasta_plazo,
        )

        importe_total = float(result.total)

        importe_desglose = {
            "apendice_m_version": "v2.2",
            "base": float(result.base),
            "extras": [
                {
                    "code": e.code,
                    "description": e.description,
                    "amount": float(e.amount),
                }
                for e in result.extras
            ],
            "urgency_surcharge": float(result.urgency_surcharge),
            "urgent": result.urgency_surcharge > 0,
            "total": importe_total,
            "iva_percent": 21.0,
            "iva_importe": round(importe_total * 0.21, 2),
            "total_con_iva": round(importe_total * 1.21, 2),
            "is_aapp": result.is_aapp,
            "payment_days": result.payment_days,
            "breakdown_text": result.breakdown_text,
        }

        hitos_dict = {
            "hitos": [
                {
                    "code": h.code,
                    "pct": float(h.pct * 100),
                    "description": h.description,
                    "amount": float(h.amount),
                }
                for h in result.hitos
            ],
        }

        proposal = Proposal(
            lead_id=lead_id,
            project_id=project_id,
            pricing_model_id=f"apendice_m_{result.categoria.lower()}",
            version=1,
            categoria_objetivo=result.categoria,
            alcance={
                "sector": sector,
                "sistemas_en_alcance": sistemas_en_alcance,
                "sedes": sedes,
                "madurez_pct": madurez_pct,
                "dias_hasta_plazo": dias_hasta_plazo,
            },
            duracion_semanas=result.hours_range[1] // 5 + 2,
            effort_marcos_horas=result.hours_range[1],
            importe_total=importe_total,
            importe_desglose=importe_desglose,
            hitos_pago=hitos_dict,
            validez_hasta=date.today() + timedelta(days=validez_dias),
            estado="draft",
            notas_marcos=(
                (notas_marcos or "") + "\n\nGarantia: " + result.garantia
            ).strip(),
        )
        db.add(proposal)
        await db.flush()

        if narrative_mode == "llm" and project_id is not None:
            try:
                llm_draft = await self._generate_narrative_via_agent_19(
                    db,
                    project_id=project_id,
                    sistemas_en_alcance=sistemas_en_alcance,
                    sedes=sedes,
                    madurez_pct=madurez_pct,
                    dias_hasta_plazo=dias_hasta_plazo,
                    sector_override=sector,
                    categoria_override=categoria,
                )
                # Persistimos en importe_desglose bajo clave "llm_narrative" para
                # no alterar el contrato existente de los consumidores.
                desglose = dict(proposal.importe_desglose or {})
                desglose["llm_narrative"] = llm_draft
                proposal.importe_desglose = desglose
                await db.flush()
            except Exception as exc:
                # Fallback transparente: el agente ya tiene fallback interno, pero
                # si explota (p. ej. falta ctx), marcamos el modo y seguimos.
                desglose = dict(proposal.importe_desglose or {})
                desglose["llm_narrative_error"] = str(exc)[:500]
                proposal.importe_desglose = desglose
                await db.flush()

        return proposal

    async def _generate_narrative_via_agent_19(
        self,
        db: AsyncSession,
        *,
        project_id: uuid.UUID,
        sistemas_en_alcance: int,
        sedes: int,
        madurez_pct: int | None,
        dias_hasta_plazo: int | None,
        sector_override: str | None,
        categoria_override: str,
    ) -> dict[str, Any]:
        """Invoca el Agente 19 para producir narrativa LLM (Opus 4.7)."""
        # Import diferido para no romper import cycles ni forzar LiteLLM
        # en arranques que no usan LLM.
        from backend.app.agents.agent_19_propuestas import Agent19Proposals

        agent = Agent19Proposals()
        return await agent.generate_proposal(
            db,
            project_id=project_id,
            sistemas_en_alcance=sistemas_en_alcance,
            sedes=sedes,
            madurez_pct=madurez_pct,
            dias_hasta_plazo=dias_hasta_plazo,
            sector_override=sector_override,
            categoria_override=categoria_override,
        )

    async def send_proposal(self, db: AsyncSession, proposal_id: uuid.UUID) -> Proposal:
        p = await self.get_proposal(db, proposal_id)
        if not p:
            raise ProposalError(f"Proposal {proposal_id} no encontrada")
        if p.estado not in ("draft", "under_review"):
            raise ProposalError(
                f"Solo se envían propuestas en draft/under_review (actual: {p.estado})"
            )
        p.estado = "sent"
        p.enviado_at = datetime.now(timezone.utc)
        await db.flush()
        return p

    async def create_version(
        self,
        db: AsyncSession,
        proposal_id: uuid.UUID,
        ajustes: dict[str, Any],
    ) -> Proposal:
        """Crea nueva versión de propuesta aplicando ajustes de negociación."""
        original = await self.get_proposal(db, proposal_id)
        if not original:
            raise ProposalError(f"Proposal {proposal_id} no encontrada")

        # §4.4 audit-2026-06-15 · si cambia importe_total, RECALCULAR el desglose
        # (base/IVA/total_con_iva) y RE-PROPORCIONAR los hitos por sus % originales
        # (mirror generate_revision). Antes se copiaba importe_desglose y hitos_pago
        # VERBATIM con un importe_total nuevo → IVA y sumatorio de hitos incoherentes
        # impresos en el DOCX que ve el prospecto.
        nuevo_importe = ajustes.get("importe_total", original.importe_total)
        nuevo_desglose = original.importe_desglose
        nuevos_hitos = ajustes.get("hitos_pago", original.hitos_pago)
        if (
            "importe_total" in ajustes
            and original.importe_total
            and float(nuevo_importe or 0) != float(original.importe_total or 0)
        ):
            imp = float(nuevo_importe or 0)
            base_desglose = original.importe_desglose or {}
            iva_pct = base_desglose.get("iva_percent", 21)
            nuevo_desglose = dict(base_desglose)
            nuevo_desglose["base"] = imp
            nuevo_desglose["iva_importe"] = round(imp * iva_pct / 100, 2)
            nuevo_desglose["total_con_iva"] = round(imp * (1 + iva_pct / 100), 2)
            nuevo_desglose["importe_anterior"] = float(original.importe_total or 0)
            if "hitos_pago" not in ajustes:
                anterior_hitos = (original.hitos_pago or {}).get("hitos") or []
                total_original = float(original.importe_total or 0)
                if anterior_hitos and total_original > 0:
                    nuevos_hitos = {
                        "hitos": [
                            {
                                "nombre": h.get("nombre", ""),
                                "pct": h.get("pct", 0),
                                "importe": round(imp * h.get("pct", 0) / 100, 2),
                            }
                            for h in anterior_hitos
                        ]
                    }

        nueva = Proposal(
            lead_id=original.lead_id,
            project_id=original.project_id,
            pricing_model_id=original.pricing_model_id,
            version=original.version + 1,
            categoria_objetivo=original.categoria_objetivo,
            alcance=ajustes.get("alcance", original.alcance),
            duracion_semanas=ajustes.get("duracion_semanas", original.duracion_semanas),
            effort_marcos_horas=original.effort_marcos_horas,
            importe_total=nuevo_importe,
            importe_desglose=nuevo_desglose,
            hitos_pago=nuevos_hitos,
            validez_hasta=date.today() + timedelta(days=30),
            estado="draft",
            notas_marcos=ajustes.get("notas_marcos", original.notas_marcos),
        )
        db.add(nueva)
        # original pasa a "negotiating" si no lo está ya
        if original.estado not in ("negotiating", "lost", "won"):
            original.estado = "negotiating"
        await db.flush()
        return nueva

    # ══════════════════════════════════════════════════════════════════
    # SAN-D MB-19.3 · revisions tracking + accept/supersede (ADR-041)
    # ══════════════════════════════════════════════════════════════════

    async def generate_revision(
        self,
        db: AsyncSession,
        *,
        lead_id: uuid.UUID,
        feedback_cliente: str,
        cambios_desde_anterior: str | None = None,
        importe_override: float | None = None,
        notas_marcos: str | None = None,
        validez_dias: int = 30,
        ajustes_alcance: dict | None = None,
        agent_19_metadata: dict | None = None,
    ) -> Proposal:
        """Crea revisión nueva basada en propuesta active anterior + feedback cliente.

        Workflow MB-19 CRM:
        1. Localiza propuesta active del lead (superseded=FALSE · highest version).
        2. Marca anterior superseded=TRUE · estado='superseded'.
        3. Crea nueva proposal version+1 con feedback persistido.
        4. Importe override aplicable · si NO dado · copia anterior.
        5. agent_19_metadata persistido si LLM regenera narrative (opcional · MB-19+).

        Modo "manual" (default · MB-19.A): copia contenido propuesta anterior +
        feedback persistido. NO invoca Agent_19 (deferrable MB-19.A para evitar
        coupling LLM en flujo CRM crítico · DEC-MB19A-PROPOSAL-AI-REGENERATE).

        Args:
            lead_id: UUID Lead asociado.
            feedback_cliente: feedback recibido del cliente sobre propuesta anterior.
            cambios_desde_anterior: resumen diff aplicado en esta revisión.
            importe_override: importe nuevo (€ sin IVA) · None = mantener anterior.
            notas_marcos: notas internas Marcos.
            validez_dias: días validez nueva propuesta (default 30).
            ajustes_alcance: dict ajustes alcance (override anterior si dado).
            agent_19_metadata: metadata Agent_19 si LLM regeneration aplicada.

        Returns:
            Proposal nueva (version+1) · estado='draft' · superseded=FALSE.

        Raises:
            ProposalError: lead no tiene propuesta anterior active.
        """
        # Localizar propuesta active del lead
        anterior_stmt = (
            select(Proposal)
            .where(Proposal.lead_id == lead_id)
            .where(Proposal.superseded.is_(False))
            .order_by(Proposal.version.desc())
            .limit(1)
        )
        anterior = (await db.execute(anterior_stmt)).scalar_one_or_none()
        if not anterior:
            raise ProposalError(
                f"Lead {lead_id} no tiene propuesta active · invocar "
                f"generate_proposal_apendice_m primero"
            )

        # Marcar anterior superseded
        anterior.superseded = True
        if anterior.estado not in ("won", "lost"):
            anterior.estado = "superseded"

        # Recalcular hitos si importe_override (proportional re-distribution)
        nuevo_importe = anterior.importe_total
        nuevos_hitos = anterior.hitos_pago
        if importe_override is not None:
            nuevo_importe = float(importe_override)
            anterior_hitos = (anterior.hitos_pago or {}).get("hitos") or []
            if anterior_hitos and anterior.importe_total:
                # Re-proportion hitos manteniendo % originales
                total_original = float(anterior.importe_total)
                if total_original > 0:
                    nuevos_hitos = {
                        "hitos": [
                            {
                                "nombre": h.get("nombre", ""),
                                "pct": h.get("pct", 0),
                                "importe": round(
                                    nuevo_importe * h.get("pct", 0) / 100, 2,
                                ),
                            }
                            for h in anterior_hitos
                        ]
                    }

        # Recalcular IVA en desglose
        anterior_desglose = anterior.importe_desglose or {}
        iva_pct = anterior_desglose.get("iva_percent", 21)
        nuevo_desglose = dict(anterior_desglose)
        if importe_override is not None:
            nuevo_desglose["base"] = nuevo_importe
            nuevo_desglose["iva_importe"] = round(
                nuevo_importe * iva_pct / 100, 2,
            )
            nuevo_desglose["total_con_iva"] = round(
                nuevo_importe * (1 + iva_pct / 100), 2,
            )
            # Documentar override en desglose para audit
            nuevo_desglose["importe_override_aplicado"] = True
            nuevo_desglose["importe_anterior"] = float(anterior.importe_total or 0)

        nueva = Proposal(
            lead_id=lead_id,
            project_id=anterior.project_id,
            pricing_model_id=anterior.pricing_model_id,
            version=anterior.version + 1,
            categoria_objetivo=anterior.categoria_objetivo,
            alcance=ajustes_alcance or anterior.alcance,
            duracion_semanas=anterior.duracion_semanas,
            effort_marcos_horas=anterior.effort_marcos_horas,
            importe_total=nuevo_importe,
            importe_desglose=nuevo_desglose,
            hitos_pago=nuevos_hitos,
            validez_hasta=date.today() + timedelta(days=validez_dias),
            estado="draft",
            notas_marcos=notas_marcos or anterior.notas_marcos,
            # MB-19.3 extension fields
            feedback_cliente=feedback_cliente,
            cambios_desde_anterior=cambios_desde_anterior,
            superseded=False,
            agent_19_metadata=agent_19_metadata,
        )
        db.add(nueva)
        await db.flush()
        return nueva

    async def mark_accepted(
        self,
        db: AsyncSession,
        proposal_id: uuid.UUID,
    ) -> Proposal:
        """Marca propuesta como aceptada por cliente.

        Side effects:
        - estado → 'won'
        - fecha_aceptacion = NOW()
        - superseded permanece FALSE (la accepted es la canónica)

        Sirve trigger downstream:
        - ContractSigningFlow puede leer proposal.fecha_aceptacion para
          generar contract.
        - CommercialWorkflowService puede tomar accepted proposal para
          MilestoneFactory.
        """
        from datetime import datetime as _dt, timezone as _tz
        proposal = await self.get_proposal(db, proposal_id)
        if not proposal:
            raise ProposalError(f"Proposal {proposal_id} no encontrada")

        proposal.estado = "won"
        proposal.fecha_aceptacion = _dt.now(_tz.utc)
        await db.flush()
        return proposal

    async def list_revisions(
        self,
        db: AsyncSession,
        lead_id: uuid.UUID,
    ) -> list[Proposal]:
        """Lista todas revisiones de propuestas para un lead, ordenadas por version DESC."""
        stmt = (
            select(Proposal)
            .where(Proposal.lead_id == lead_id)
            .order_by(Proposal.version.desc())
        )
        result = await db.execute(stmt)
        return list(result.scalars().all())

    async def get_active_proposal(
        self,
        db: AsyncSession,
        lead_id: uuid.UUID,
    ) -> Proposal | None:
        """Retorna la propuesta active (superseded=FALSE) más reciente del lead.

        None si lead no tiene propuestas o todas están superseded.
        """
        stmt = (
            select(Proposal)
            .where(Proposal.lead_id == lead_id)
            .where(Proposal.superseded.is_(False))
            .order_by(Proposal.version.desc())
            .limit(1)
        )
        result = await db.execute(stmt)
        return result.scalar_one_or_none()

    async def generate_docx(
        self, db: AsyncSession, proposal_id: uuid.UUID
    ) -> bytes:
        """Genera DOCX de la propuesta programáticamente (sin template externo).

        Usa python-docx para crear documento comercial estándar.
        """
        from docx import Document

        p = await self.get_proposal(db, proposal_id)
        if not p:
            raise ProposalError(f"Proposal {proposal_id} no encontrada")

        doc = Document()
        doc.add_heading(f"Propuesta comercial P-{str(p.id)[:8]} v{p.version}", level=0)
        doc.add_paragraph(
            f"Categoría objetivo ENS: {p.categoria_objetivo or '-'}"
        )
        doc.add_paragraph(f"Modelo de pricing: {p.pricing_model_id or '-'}")
        doc.add_paragraph(f"Duración estimada: {p.duracion_semanas or '?'} semanas")
        doc.add_paragraph(f"Esfuerzo Marcos: {p.effort_marcos_horas or 0} h")

        doc.add_heading("Alcance", level=1)
        alcance = p.alcance or {}
        for k, v in alcance.items():
            doc.add_paragraph(f"• {k}: {v}")

        doc.add_heading("Importe", level=1)
        desglose = p.importe_desglose or {}
        doc.add_paragraph(f"Base: {desglose.get('base', 0):.2f} €")
        for extra in desglose.get("extras") or []:
            # §4.4 audit-2026-06-15 · soporta ambos esquemas de extras: el legacy
            # (concepto/importe) y el canónico Apéndice M v2.2 (code/description/
            # amount). Antes sólo leía concepto/importe → en propuestas generadas por
            # la vía apéndice-M cada extra salía "  + : 0.00 €" (línea en blanco).
            label = extra.get("concepto") or extra.get("description") or ""
            amount = extra.get("importe", extra.get("amount", 0)) or 0
            doc.add_paragraph(f"  + {label}: {float(amount):.2f} €")
        doc.add_paragraph(f"Total (sin IVA): {float(p.importe_total or 0):.2f} €")
        doc.add_paragraph(
            f"IVA {desglose.get('iva_percent', 21)}%: {desglose.get('iva_importe', 0):.2f} €"
        )
        doc.add_paragraph(
            f"Total con IVA: {desglose.get('total_con_iva', 0):.2f} €"
        )

        doc.add_heading("Hitos de pago", level=1)
        for hito in (p.hitos_pago or {}).get("hitos") or []:
            # §2.2 audit-2026-06-16 · soporta ambos esquemas de hitos: el legacy
            # (nombre/importe) y el canónico Apéndice M (code/description/amount).
            # Antes sólo leía nombre/importe → en propuestas apéndice-M cada hito
            # salía "• : 0.00 €" (etiqueta en blanco).
            label = hito.get("nombre") or hito.get("description") or hito.get("code") or ""
            amount = hito.get("importe", hito.get("amount", 0)) or 0
            doc.add_paragraph(
                f"• {label}: {hito.get('pct', 0)}% = {float(amount):.2f} €"
            )

        doc.add_paragraph(f"\nValidez hasta: {p.validez_hasta or '-'}")
        if p.notas_marcos:
            doc.add_heading("Notas", level=1)
            doc.add_paragraph(p.notas_marcos)

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
