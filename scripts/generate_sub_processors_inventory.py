"""Generate Sub_Processors_Inventory.xlsx (AEPD-compatible format).

Output: docs/compliance/06-Sub_Processors/Sub_Processors_Inventory.xlsx

Columns:
- Provider · Service · Location · GDPR mechanism · DPA renewal date ·
  Sub-sub-processors · Risk level · Notes

Mantained in Spanish formal for AEPD-facing inventory.
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROWS = [
    {
        "proveedor": "Hetzner Online GmbH",
        "servicio": "Infraestructura de alojamiento (IaaS)",
        "localizacion": "Falkenstein, Alemania (EEE)",
        "mecanismo_rgpd": "DPA · sin SCC (intra-EEE)",
        "fecha_renovacion_dpa": "2027-04-01",
        "sub_sub_procesadores": "Ninguno declarado",
        "nivel_riesgo": "Bajo (intra-EEE · proveedor consolidado)",
        "notas": "Bases de datos PostgreSQL · objetos MinIO · servidores backend/frontend",
    },
    {
        "proveedor": "Postmark (ActiveCampaign Inc.)",
        "servicio": "Correo electrónico transaccional",
        "localizacion": "Región UE Fráncfort (seleccionada)",
        "mecanismo_rgpd": "DPA + SCC 2021",
        "fecha_renovacion_dpa": "2027-03-15",
        "sub_sub_procesadores": "AWS Fráncfort (DPA heredado)",
        "nivel_riesgo": "Bajo (EU data region · SCC vigentes)",
        "notas": "5 plantillas MJML brand-coherent · webhook HMAC firmado",
    },
    {
        "proveedor": "Anthropic PBC",
        "servicio": "Inteligencia artificial generativa (Claude API)",
        "localizacion": "US primario · UE Fráncfort disponible",
        "mecanismo_rgpd": "DPA + SCC 2021 + TADPF",
        "fecha_renovacion_dpa": "2026-12-20",
        "sub_sub_procesadores": "AWS US-East · Google Cloud Platform US-Central",
        "nivel_riesgo": "Medio (transferencia extra-EEE mitigada · seudonimización)",
        "notas": "Datos seudonimizados antes envío · Plan contingencia AWS Bedrock Fráncfort si Schrems III",
    },
    {
        "proveedor": "360dialog GmbH",
        "servicio": "WhatsApp Business API (Business Solution Provider)",
        "localizacion": "Alemania (EEE)",
        "mecanismo_rgpd": "DPA · sin SCC (intra-EEE)",
        "fecha_renovacion_dpa": "2027-05-30",
        "sub_sub_procesadores": "WhatsApp/Meta (DPA heredado · 360dialog intermediario BSP)",
        "nivel_riesgo": "Medio (sub-encargado Meta · jurisdicción compleja)",
        "notas": "Templates aprobados · KYC Meta Business pendiente Marcos",
    },
    {
        "proveedor": "MinIO (autoalojado Hetzner)",
        "servicio": "Almacenamiento de objetos",
        "localizacion": "Hetzner Alemania (autoalojado)",
        "mecanismo_rgpd": "N/A · FULKRO controla 100% infraestructura",
        "fecha_renovacion_dpa": "N/A",
        "sub_sub_procesadores": "Ninguno",
        "nivel_riesgo": "Bajo (control total)",
        "notas": "AES-256 en reposo · URLs firmadas TTL limitado · cubos: cliente-data, compliance-reports, backups, evidence",
    },
]


HEADERS = [
    ("proveedor", "Proveedor"),
    ("servicio", "Servicio"),
    ("localizacion", "Localización"),
    ("mecanismo_rgpd", "Mecanismo RGPD"),
    ("fecha_renovacion_dpa", "Fecha renovación DPA"),
    ("sub_sub_procesadores", "Sub-encargados"),
    ("nivel_riesgo", "Nivel de riesgo"),
    ("notas", "Notas"),
]


def main() -> None:
    out_path = Path("docs/compliance/06-Sub_Processors/Sub_Processors_Inventory.xlsx")
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sub-procesadores"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrapped = Alignment(vertical="top", wrap_text=True)

    for col_idx, (_, label) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = centered

    for row_idx, row in enumerate(ROWS, start=2):
        for col_idx, (key, _) in enumerate(HEADERS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row[key])
            cell.alignment = wrapped

    widths = [28, 38, 30, 28, 18, 38, 36, 50]
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 30
    for row_idx in range(2, len(ROWS) + 2):
        ws.row_dimensions[row_idx].height = 90

    meta_ws = wb.create_sheet("Metadatos")
    meta_rows = [
        ("Código", "FULKRO-SUB-INV-001"),
        ("Versión", "1.0"),
        ("Fecha aprobación", date(2026, 5, 12).isoformat()),
        ("Próxima revisión", date(2027, 5, 12).isoformat()),
        ("Aprobado por", "Marcos Mata García"),
        ("Cargo aprobador", "Delegado de Protección de Datos (interim)"),
        ("Clasificación", "Interno · referenciable AEPD"),
        ("Referencia normativa", "RGPD Art. 28, Art. 30 · LOPDGDD Art. 31"),
        (
            "Periodicidad revisión",
            "Anual obligatoria + extraordinaria por incorporación o baja de sub-procesador",
        ),
    ]
    for row_idx, (key, value) in enumerate(meta_rows, start=1):
        ws_key = meta_ws.cell(row=row_idx, column=1, value=key)
        ws_key.font = Font(bold=True)
        meta_ws.cell(row=row_idx, column=2, value=value)

    meta_ws.column_dimensions["A"].width = 28
    meta_ws.column_dimensions["B"].width = 70

    wb.save(out_path)
    print(f"OK · escrito {out_path} con {len(ROWS)} sub-procesadores")


if __name__ == "__main__":
    main()
